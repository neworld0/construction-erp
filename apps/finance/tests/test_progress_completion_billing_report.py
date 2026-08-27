from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.test import Client

from apps.core.rbac.models import LegalEntity, LegalEntityAccessScope, Role, UserLegalEntityMembership, UserProfile
from apps.finance.models import BillingReportType
from apps.finance.services.billing_reports import build_billing_report_preview, create_billing_report, export_billing_report_excel
from apps.finance.services.progress_billing import issue_progress_billing, register_advance_payment
from apps.projects.models import BudgetCategory, BudgetItem, Project
from apps.cost.models import CostItem, CostItemCategory
from apps.schedule.models import DailyProgress, SchedulePlan, ScheduleTask


@pytest.mark.django_db
def test_progress_and_completion_report_snapshot_and_export(settings):
    settings.MIDDLEWARE = [m for m in settings.MIDDLEWARE if "two_factor_enforce" not in m]
    hq = get_user_model().objects.create_user(username="billing-report-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    UserLegalEntityMembership.objects.create(
        user=hq,
        legal_entity=LegalEntity.objects.get(code="ASAN"),
        access_scope=LegalEntityAccessScope.ENTITY_HQ,
    )
    project = Project.objects.create(code="REPORT-001", name="기성 보고서 현장", project_type="civil", contract_amount=Decimal("132000000"))
    plan = SchedulePlan.objects.create(project=project, version_no=1, name="기준선")
    task = ScheduleTask.objects.create(plan=plan, name="공정", weight_percent=Decimal("100"), start_date=date(2026, 8, 1), end_date=date(2026, 9, 13))
    labor_item = CostItem.objects.create(code="REPORT-LABOR", name="노무비", category=CostItemCategory.LABOR)
    material_item = CostItem.objects.create(code="REPORT-MATERIAL", name="재료비", category=CostItemCategory.MATERIAL)
    BudgetItem.objects.create(project=project, cost_item=labor_item, category=BudgetCategory.LABOR, name="노무비", planned_amount=25000000)
    BudgetItem.objects.create(project=project, cost_item=labor_item, category=BudgetCategory.LABOR, name="노무비", planned_amount=15000000)
    BudgetItem.objects.create(project=project, cost_item=material_item, category=BudgetCategory.MATERIAL, name="재료비", planned_amount=92000000)
    register_advance_payment(project=project, advance_rate_percent="40", received_date=date(2026, 8, 1), actor=hq)
    DailyProgress.objects.create(project=project, plan=plan, task=task, report_date=date(2026, 8, 31), progress_percent=Decimal("70.5"), status="approved", reporter=hq)
    issue_progress_billing(project=project, billing_date=date(2026, 8, 31), actor=hq)
    first = build_billing_report_preview(project, report_type=BillingReportType.PROGRESS, billing_date=date(2026, 8, 31))
    assert first["gross_claim_amount"] == Decimal("93060000.00")
    assert first["advance_deduction_amount"] == Decimal("37224000.00")
    assert first["net_claim_amount"] == Decimal("55836000.00")
    progress_report = create_billing_report(project=project, report_type=BillingReportType.PROGRESS, billing_date=date(2026, 8, 31), billing_round=1, actor=hq)

    DailyProgress.objects.create(project=project, plan=plan, task=task, report_date=date(2026, 9, 13), progress_percent=Decimal("100"), status="approved", reporter=hq)
    completion = build_billing_report_preview(project, report_type=BillingReportType.COMPLETION, billing_date=date(2026, 9, 13))
    assert completion["gross_claim_amount"] == Decimal("38940000.00")
    assert completion["advance_deduction_amount"] == Decimal("15576000.00")
    assert completion["net_claim_amount"] == Decimal("23364000.00")
    assert completion["advance_balance_after"] == Decimal("0.00")
    completion_report = create_billing_report(project=project, report_type=BillingReportType.COMPLETION, billing_date=date(2026, 9, 13), billing_round=1, actor=hq)

    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(export_billing_report_excel(progress_report)))
    assert workbook.sheetnames == ["갑지", "원가계산서", "내역서", "선급금정산", "엔지니어작성사항", "증빙체크리스트"]
    assert "A1:B1" in {str(merged) for merged in workbook["갑지"].merged_cells.ranges}
    assert workbook["내역서"].page_setup.orientation == "landscape"
    assert workbook["내역서"]["A3"].fill.fill_type == "solid"
    assert workbook["갑지"]["B13"].number_format == "0.0%"
    assert workbook["내역서"]["A4"].value == "노무비"
    assert workbook["내역서"]["A5"].value == "재료비"
    assert workbook["내역서"]["A6"].value is None
    assert sum(workbook["내역서"].cell(row, 4).value for row in (4, 5)) == 93060000
    assert sum(workbook["내역서"].cell(row, 5).value for row in (4, 5)) == 93060000
    assert sum(workbook["내역서"].cell(row, 6).value for row in (4, 5)) == 38940000
    assert completion_report.remaining_advance_balance == Decimal("0.00")
    client = Client()
    client.force_login(hq)
    response = client.get(f"/app/hq/billing/reports/{progress_report.id}/")
    assert response.status_code == 200
    assert "엔지니어 작성 항목" in response.content.decode("utf-8")
