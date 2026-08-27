from datetime import date
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.labor.models import IncomeTaxTableRow, IncomeTaxTableVersion


class Command(BaseCommand):
    help = "Import an official Korean monthly income-tax withholding table workbook."

    def add_arguments(self, parser):
        parser.add_argument("file_path")
        parser.add_argument("--effective-from", required=True)
        parser.add_argument("--actor-username", required=True)

    def handle(self, *args, **options):
        try:
            effective_from = date.fromisoformat(options["effective_from"])
        except ValueError as exc:
            raise CommandError("--effective-from must be YYYY-MM-DD") from exc
        actor = get_user_model().objects.filter(username=options["actor_username"]).first()
        if actor is None:
            raise CommandError("actor user not found")
        file_path = options["file_path"]
        if file_path == "AUTO":
            matches = [path for path in (Path.home() / "Downloads").glob("*.xlsx") if path.name.endswith("2026.03.01.xlsx")]
            if len(matches) != 1:
                raise CommandError("AUTO could not uniquely find the 2026.03.01 workbook in Downloads")
            file_path = str(matches[0])
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.worksheets[-1]
        rows = []
        for values in sheet.iter_rows(min_row=5, min_col=1, max_col=13, values_only=True):
            lower, upper = values[0], values[1]
            if not isinstance(lower, (int, float)) or not isinstance(upper, (int, float)):
                continue
            for family_count, raw_tax in enumerate(values[2:13], 1):
                tax = 0 if raw_tax in (None, "-") else int(raw_tax)
                rows.append(IncomeTaxTableRow(monthly_pay_from=int(lower * 1000), monthly_pay_to=int(upper * 1000), dependent_count=family_count, income_tax=tax))
        if not rows:
            raise CommandError("No tax-table rows found. Expected official workbook layout.")
        with transaction.atomic():
            IncomeTaxTableVersion.objects.filter(is_active=True).update(is_active=False)
            version, _ = IncomeTaxTableVersion.objects.update_or_create(effective_from=effective_from, defaults={"source_name": Path(file_path).name, "is_active": True, "created_by": actor})
            version.rows.all().delete()
            for row in rows:
                row.version = version
            IncomeTaxTableRow.objects.bulk_create(rows, batch_size=1000)
        self.stdout.write(self.style.SUCCESS(f"Imported {len(rows)} rows into table version {version.id}."))
