from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


PAY_FIELDS = (
    ("base_pay", "기본급"), ("meal_allowance_pay", "식대(비과세)"), ("fuel_allowance_pay", "주유대(비과세)"),
    ("site_allowance_pay", "현장수당"), ("overtime_pay", "연장수당"), ("bonus_pay", "상여"),
    ("income_tax", "소득세"), ("local_income_tax", "지방소득세"), ("national_pension", "국민연금"),
    ("health_insurance", "건강보험"), ("long_term_care", "장기요양"), ("employment_insurance", "고용보험"),
    ("other_deduction", "기타공제"),
)
PAYMENT_FIELDS = PAY_FIELDS[:6]
DEDUCTION_FIELDS = PAY_FIELDS[6:]


def _name(slip):
    return slip.employee.user.get_full_name() or slip.employee.user.username


def _apply_sheet_style(sheet, end_col, title):
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = sheet.cell(1, 1, title)
    cell.font = Font(name="맑은 고딕", size=16, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1E3A5F")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28
    for column in range(1, end_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    return border


def payroll_workbook(run):
    book = Workbook()
    sheet = book.active
    sheet.title = "급여대장"
    headers = ["사번", "성명", "부서"] + [label for _, label in PAY_FIELDS] + ["총지급", "공제합계", "실지급"]
    border = _apply_sheet_style(sheet, len(headers), f"{run.period_year}년 {run.period_month:02d}월 본사 급여대장")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.cell(2, 1, f"상태: {run.get_status_display()}  |  지급일자: {run.payment_date:%Y년 %m월 %d일}  |  확정·지급된 급여 기준 출력")
    sheet.cell(2, 1).alignment = Alignment(horizontal="left")
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(4, col, header)
        cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_no, slip in enumerate(run.slips.select_related("employee__user"), 5):
        values = [slip.employee.employee_no, _name(slip), slip.employee.department] + [getattr(slip, field) for field, _ in PAY_FIELDS] + [slip.gross_pay, slip.total_deduction, slip.net_pay]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_no, col, value)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if col <= 3 else "right", vertical="center")
            if col > 3:
                cell.number_format = '#,##0;[Red]-#,##0'
    total_row = 5 + run.slips.count()
    sheet.cell(total_row, 1, "합계")
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    for col in range(1, len(headers) + 1):
        cell = sheet.cell(total_row, col)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor="DBEAFE")
        cell.font = Font(name="맑은 고딕", bold=True)
        if col >= 4:
            cell.value = f"=SUM({get_column_letter(col)}5:{get_column_letter(col)}{total_row - 1})"
            cell.number_format = '#,##0;[Red]-#,##0'
    sheet.freeze_panes = "D5"
    stream = BytesIO(); book.save(stream); stream.seek(0)
    return stream


def payslip_workbook(run, slip):
    book = Workbook()
    sheet = book.active
    sheet.title = "급여명세서"
    border = _apply_sheet_style(sheet, 4, f"{run.period_year}년 {run.period_month:02d}월 급여명세서")
    details = [("사번", slip.employee.employee_no), ("성명", _name(slip)), ("부서", slip.employee.department), ("상태", run.get_status_display()), ("지급일자", run.payment_date.strftime("%Y년 %m월 %d일"))]
    for row, (label, value) in enumerate(details, 3):
        sheet.cell(row, 1, label); sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4); sheet.cell(row, 2, value)
        for col in range(1, 5): sheet.cell(row, col).border = border
    sheet.cell(9, 1, "지급 항목"); sheet.cell(9, 2, "금액"); sheet.cell(9, 3, "공제 항목"); sheet.cell(9, 4, "금액")
    for col in range(1, 5):
        sheet.cell(9, col).fill = PatternFill("solid", fgColor="2563EB"); sheet.cell(9, col).font = Font(name="맑은 고딕", bold=True, color="FFFFFF"); sheet.cell(9, col).border = border
    earnings = [(label, getattr(slip, field)) for field, label in PAYMENT_FIELDS]
    deductions = [(label, getattr(slip, field)) for field, label in DEDUCTION_FIELDS]
    for index in range(max(len(earnings), len(deductions))):
        row = 10 + index
        left = earnings[index] if index < len(earnings) else ("", "")
        right = deductions[index] if index < len(deductions) else ("", "")
        for col, value in enumerate([left[0], left[1], right[0], right[1]], 1):
            cell = sheet.cell(row, col, value); cell.border = border; cell.number_format = '#,##0;[Red]-#,##0' if col in (2, 4) else "General"
    for row, label, value in [(17, "총지급", slip.gross_pay), (18, "공제합계", slip.total_deduction), (19, "실지급", slip.net_pay)]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2); sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        sheet.cell(row, 1, label); sheet.cell(row, 3, value)
        for col in range(1, 5): sheet.cell(row, col).border = border; sheet.cell(row, col).fill = PatternFill("solid", fgColor="DBEAFE")
        sheet.cell(row, 3).number_format = '#,##0;[Red]-#,##0'
    stream = BytesIO(); book.save(stream); stream.seek(0)
    return stream


def _pdf_document(landscape_mode=False):
    stream = BytesIO()
    page_size = landscape(A4) if landscape_mode else A4
    document = SimpleDocTemplate(stream, pagesize=page_size, leftMargin=12*mm, rightMargin=12*mm, topMargin=12*mm, bottomMargin=12*mm)
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="KoreanTitle", parent=styles["Title"], fontName="HYSMyeongJo-Medium", fontSize=17, leading=24, alignment=1))
    styles.add(ParagraphStyle(name="Korean", parent=styles["Normal"], fontName="HYSMyeongJo-Medium", fontSize=8, leading=12))
    styles.add(ParagraphStyle(name="KoreanTable", parent=styles["Normal"], fontName="HYSMyeongJo-Medium", fontSize=7, leading=9, wordWrap="CJK"))
    # Amounts must stay on one line in the compact landscape register.  The
    # CJK wrapping used for Korean labels would otherwise split 5,300,000
    # across multiple lines.
    styles.add(ParagraphStyle(name="KoreanTableNumber", parent=styles["Normal"], fontName="HYSMyeongJo-Medium", fontSize=5.5, leading=7, alignment=2, splitLongWords=False))
    styles.add(ParagraphStyle(name="KoreanTableHeader", parent=styles["KoreanTable"], textColor=colors.white, alignment=1, leading=10))
    return stream, document, styles


def payroll_pdf(run):
    stream, document, styles = _pdf_document(landscape_mode=True)
    # Keep the PDF register aligned with the screen and Excel register.  In
    # particular, 지방소득세 is a statutory deduction and must never be
    # omitted from a document that presents the total deduction.
    headers = ["사번", "성명"] + [label for _, label in PAY_FIELDS] + ["총지급", "공제합계", "실지급"]
    # ReportLab plain strings do not wrap.  A long employee number therefore
    # painted over the name cell in the downloaded register.  Use paragraphs
    # in every cell and add valid break opportunities to the employee number.
    data = [[Paragraph(escape(header), styles["KoreanTableHeader"]) for header in headers]]
    for slip in run.slips.select_related("employee__user"):
        employee_no = escape(slip.employee.employee_no).replace("-", "-\u200b")
        values = (
            *[getattr(slip, field) for field, _ in PAY_FIELDS],
            slip.gross_pay,
            slip.total_deduction,
            slip.net_pay,
        )
        data.append([
            Paragraph(employee_no, styles["KoreanTable"]),
            Paragraph(escape(_name(slip)), styles["KoreanTable"]),
            *[Paragraph(f"{value:,.0f}", styles["KoreanTableNumber"]) for value in values],
        ])
    story = [Paragraph(f"{run.period_year}년 {run.period_month:02d}월 본사 급여대장", styles["KoreanTitle"]), Spacer(1, 5*mm), Paragraph(f"상태: {run.get_status_display()} / 지급일자: {run.payment_date:%Y년 %m월 %d일} / 확정·지급된 급여 기준 출력", styles["Korean"]), Spacer(1, 4*mm)]
    # The landscape A4 printable width is 273mm.  The compact amount columns
    # retain all statutory deductions while avoiding text overlap.
    table = Table(data, repeatRows=1, colWidths=[25*mm, 18*mm] + [12*mm] * len(PAY_FIELDS) + [16*mm] * 3)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1E3A5F")), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#94A3B8")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")])]))
    story.append(table); document.build(story); stream.seek(0); return stream


def payslip_pdf(run, slip):
    stream, document, styles = _pdf_document()
    story = [Paragraph(f"{run.period_year}년 {run.period_month:02d}월 급여명세서", styles["KoreanTitle"]), Spacer(1, 5*mm)]
    profile = [["사번", slip.employee.employee_no], ["성명", _name(slip)], ["부서", slip.employee.department], ["상태", run.get_status_display()], ["지급일자", run.payment_date.strftime("%Y년 %m월 %d일")]]
    pay = [["지급 항목", "금액", "공제 항목", "금액"]]
    earnings = [(label, getattr(slip, field)) for field, label in PAYMENT_FIELDS]
    deductions = [(label, getattr(slip, field)) for field, label in DEDUCTION_FIELDS]
    for index in range(max(len(earnings), len(deductions))):
        left = earnings[index] if index < len(earnings) else ("", "")
        right = deductions[index] if index < len(deductions) else ("", "")
        pay.append([left[0], f"{left[1]:,.0f}" if left[0] else "", right[0], f"{right[1]:,.0f}" if right[0] else ""])
    pay += [["총지급", f"{slip.gross_pay:,.0f}", "공제합계", f"{slip.total_deduction:,.0f}"], ["실지급", f"{slip.net_pay:,.0f}", "", ""]]
    for data, widths in ((profile, [35*mm, 130*mm]), (pay, [42*mm, 42*mm, 42*mm, 42*mm])):
        table = Table(data, colWidths=widths)
        table.setStyle(TableStyle([("FONTNAME", (0,0), (-1,-1), "HYSMyeongJo-Medium"), ("FONTSIZE", (0,0), (-1,-1), 9), ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#94A3B8")), ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#DBEAFE")), ("ALIGN", (1,0), (-1,-1), "RIGHT"), ("VALIGN", (0,0), (-1,-1), "MIDDLE")]))
        story.append(table); story.append(Spacer(1, 5*mm))
    document.build(story); stream.seek(0); return stream
