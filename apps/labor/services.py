import logging
import hashlib
from datetime import date as date_type, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import re
from calendar import monthrange

from django.core.files.base import ContentFile
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Q, Sum
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.services.logger import log_action
from apps.closing.guards import guard_write
from apps.closing.services import is_month_closed
from apps.core.rbac.models import Role
from apps.core.rbac.permissions import require_project_access, require_role
from apps.cost.models import CostActual, CostItem, CostItemCategory
from apps.projects.models import Project

from .models import (
    ElectronicCardImportBatch,
    ElectronicCardImportBatchStatus,
    ElectronicCardMatchStatus,
    LaborExcelExportBatch,
    LaborExcelExportStatus,
    LaborExcelExportType,
    LaborConfirmedWorkDay,
    LaborConfirmedWorkSourceBasis,
    ElectronicCardWorkDay,
    ElectronicCardWorkRaw,
    LaborReconciliationResolution,
    LaborReconciliationResult,
    LaborReconciliationStatus,
    LaborMonthlyPayroll,
    LaborMonthlyPayrollPaymentStatus,
    LaborWorkLedger,
    LaborWorkLedgerSource,
    LaborWorkLedgerStatus,
    LaborRateScope,
    LaborRateTable,
    LaborRateType,
    LaborRole,
    PayrollAllocationBatch,
    PayrollAllocationLine,
    PayrollAllocationStatus,
    PayrollBatchNumberSequence,
    Timesheet,
    TimesheetLine,
    TimesheetNumberSequence,
    TimesheetStatus,
    WorkerMaster,
    _protect_sensitive_value,
)

logger = logging.getLogger(__name__)

# Discovery:
# - CostItem: apps/cost/models.py CostItem
# - RBAC: apps/core/rbac/permissions.py require_role
# - HQ master UI pattern: templates/app/hq/master_* + apps/master/web_urls.py

FALLBACK_LABOR_CBS_CODE = "LABOR-GENERAL"
FALLBACK_LABOR_CBS_NAME = "인건비(공통)"
ELECTRONIC_CARD_IMPORT_UPLOAD = "LABOR_ECARD_IMPORT_UPLOAD"
ELECTRONIC_CARD_IMPORT_PARSE = "LABOR_ECARD_IMPORT_PARSE"
ELECTRONIC_CARD_RECONCILE = "LABOR_ECARD_RECONCILE"
LABOR_RECONCILIATION_RESOLVE = "LABOR_RECONCILIATION_RESOLVE"
LABOR_RECONCILIATION_BULK_RESOLVE = "LABOR_RECONCILIATION_BULK_RESOLVE"
LABOR_RECONCILIATION_WORKER_MATCH = "LABOR_RECONCILIATION_WORKER_MATCH"
LABOR_CONFIRMED_WORKDAY_GENERATE = "LABOR_CONFIRMED_WORKDAY_GENERATE"
LABOR_ECARD_RECONCILIATION_CONFIRM = "LABOR_ECARD_RECONCILIATION_CONFIRM"
LABOR_EXCEL_EXPORT_GENERATE = "LABOR_EXCEL_EXPORT_GENERATE"
LABOR_EXCEL_EXPORT_DOWNLOAD = "LABOR_EXCEL_EXPORT_DOWNLOAD"
LABOR_WORKER_DELETE = "LABOR_WORKER_DELETE"
LABOR_WORKER_DEACTIVATE = "LABOR_WORKER_DEACTIVATE"
LABOR_WORKER_ALREADY_INACTIVE = "LABOR_WORKER_ALREADY_INACTIVE"

E_CARD_HEADER_LABELS = {
    "row_no": ("No", "번호"),
    "year_month": ("근로년월", "근로연월", "근로월", "연월", "귀속연월"),
    "work_history_status": ("근로내역상태", "근로내역 상태"),
    "report_status": ("신고상태", "신고 상태"),
    "project_name": ("공사명", "현장명", "프로젝트명"),
    "deduction_join_no": ("공제가입번호", "공제 가입번호"),
    "company_name": ("업체명", "회사명", "사업장명"),
    "job_type": ("직종", "직무"),
    "worker_name": ("성명", "근로자명", "이름"),
    "card_issued": ("전자카드발급여부", "전자카드 발급여부"),
    "resident_no": ("주민등록번호", "주민번호", "생년월일"),
    "phone": ("연락처", "휴대전화", "전화번호"),
    "retirement_deduction": ("퇴직공제여부", "퇴직공제 여부"),
    "exclusion_reason": ("비퇴직사유", "제외사유"),
    "error_message": ("오류확인내역", "오류/확인내역", "오류내역"),
    "auto_work_days": ("자동집계", "자동 출역일수"),
    "reported_days": ("신고일수",),
    "declaration_days": ("이월예정일수", "이월 예정일수", "예정일수"),
    "confirmed_days": ("확정일수",),
    "note": ("비고",),
}
E_CARD_ESSENTIAL_FIELDS = {
    "year_month",
    "project_name",
    "job_type",
    "worker_name",
    "resident_no",
}
E_CARD_HEADER_LABELS["exclusion_reason"] = (
    "비대상사유",
    "비대상 사유",
    "비퇴직사유",
    "제외사유",
)
E_CARD_OPTIONAL_FIELDS = set(E_CARD_HEADER_LABELS.keys()) - E_CARD_ESSENTIAL_FIELDS


def _log_action_safe(*, actor, action, object_id, summary, metadata, object_type):
    if actor is None:
        return
    try:
        log_action(
            actor=actor,
            action=action,
            object_type=object_type,
            object_id=object_id,
            meta=metadata,
        )
    except Exception:
        logger.warning("AuditLog failed for %s.", action, exc_info=True)


def _resolve_cost_item(value):
    if not value:
        return None
    if isinstance(value, CostItem):
        return value
    return CostItem.objects.filter(id=value).first()


def _resolve_project(value):
    if not value:
        return None
    if isinstance(value, Project):
        return value
    return Project.objects.filter(id=value).first()


def _resolve_worker(value):
    if not value:
        return None
    if isinstance(value, WorkerMaster):
        return value
    return WorkerMaster.objects.filter(id=value).first()


def _resolve_timesheet(value):
    if not value:
        return None
    if isinstance(value, Timesheet):
        return value
    return Timesheet.objects.filter(id=value).first()


def _resolve_cost_actual(value):
    if not value:
        return None
    if isinstance(value, CostActual):
        return value
    return CostActual.objects.filter(id=value).first()


def _normalize_identity_value(raw_value: str) -> str:
    return "".join(ch for ch in str(raw_value or "") if ch.isdigit())


def _mask_rrn_value(raw_value: str) -> str:
    normalized = _normalize_identity_value(raw_value)
    if len(normalized) >= 7:
        return f"{normalized[:6]}-{normalized[6]}******"
    if not normalized:
        return ""
    return f"{normalized[:1]}***"


def _mask_account_number_value(raw_value: str) -> str:
    value = str(raw_value or "").strip()
    if not value:
        return ""
    visible = value[-4:] if len(value) > 4 else value
    return f"***{visible}"


def _build_identity_hash(raw_value: str) -> str:
    normalized = _normalize_identity_value(raw_value)
    if not normalized:
        return ""
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _store_sensitive_value(raw_value: str) -> str:
    value = str(raw_value or "").strip()
    if not value:
        return ""
    # TODO(LABPAY-1): replace isolated plaintext storage with shared encryption util.
    return value


def _worker_audit_snapshot(worker: WorkerMaster) -> dict:
    return {
        "name": worker.name,
        "rrn_masked": worker.rrn_masked,
        "phone": worker.phone,
        "bank_name": worker.bank_name,
        "bank_code": worker.bank_code,
        "account_number_masked": worker.account_number_masked,
        "account_holder": worker.account_holder,
        "nationality_code": worker.nationality_code,
        "visa_code": worker.visa_code,
        "comwel_job_code": worker.comwel_job_code,
        "cwma_job_name": worker.cwma_job_name,
        "default_labor_role_id": worker.default_labor_role_id,
        "retirement_deduction_eligible": worker.retirement_deduction_eligible,
        "active": worker.active,
        "identity_hash": worker.identity_hash,
    }


def _labor_work_ledger_snapshot(ledger: LaborWorkLedger) -> dict:
    return {
        "work_date": str(ledger.work_date),
        "work_month": str(ledger.work_month),
        "worker_id": ledger.worker_id,
        "actual_project_id": ledger.actual_project_id,
        "report_project_id": ledger.report_project_id,
        "labor_role_id": ledger.labor_role_id,
        "work_unit": str(ledger.work_unit),
        "work_hours": str(ledger.work_hours),
        "unit_wage": ledger.unit_wage,
        "gross_wage": ledger.gross_wage,
        "income_tax": ledger.income_tax,
        "local_tax": ledger.local_tax,
        "employment_insurance": ledger.employment_insurance,
        "pension": ledger.pension,
        "health_insurance": ledger.health_insurance,
        "net_pay": ledger.net_pay,
        "detail_work_type": ledger.detail_work_type,
        "source": ledger.source,
        "status": ledger.status,
        "timesheet_ref_id": ledger.timesheet_ref_id,
        "cost_ref_id": ledger.cost_ref_id,
    }


def _labor_monthly_payroll_snapshot(payroll: LaborMonthlyPayroll) -> dict:
    return {
        "year_month": str(payroll.year_month),
        "worker_id": payroll.worker_id,
        "project_id": payroll.project_id,
        "report_project_id": payroll.report_project_id,
        "total_work_unit": str(payroll.total_work_unit),
        "gross_wage": payroll.gross_wage,
        "income_tax": payroll.income_tax,
        "local_tax": payroll.local_tax,
        "insurance_deductions": payroll.insurance_deductions,
        "net_pay": payroll.net_pay,
        "bank_name": payroll.bank_name,
        "account_number_masked": payroll.account_number_masked,
        "payment_status": payroll.payment_status,
        "paid_at": payroll.paid_at.isoformat() if payroll.paid_at else "",
    }


def _overlap_q(start, end):
    if end:
        return Q(effective_from__lte=end) & (
            Q(effective_to__isnull=True) | Q(effective_to__gte=start)
        )
    return Q(effective_to__isnull=True) | Q(effective_to__gte=start)


def _month_range(year: int, month: int) -> tuple[date_type, date_type]:
    month = int(month)
    year = int(year)
    start = date_type(year, month, 1)
    if month == 12:
        end = date_type(year, 12, 31)
    else:
        end = date_type(year, month + 1, 1) - timedelta(days=1)
    return start, end


def _get_or_create_fallback_cbs() -> CostItem:
    fallback = CostItem.objects.filter(code=FALLBACK_LABOR_CBS_CODE).first()
    if fallback:
        return fallback
    return CostItem.objects.create(
        code=FALLBACK_LABOR_CBS_CODE,
        name=FALLBACK_LABOR_CBS_NAME,
        category=CostItemCategory.LABOR,
        cost_type="L",
        work_type="99",
        is_direct=False,
        is_active=True,
    )


def compute_labor_data_quality(project_id: int, year: int, month: int) -> str:
    period_date = date_type(int(year), int(month), 1)
    is_closed = is_month_closed(period_date)
    has_timesheet = Timesheet.objects.filter(
        project_id=project_id,
        status=TimesheetStatus.APPROVED,
        work_date__year=year,
        work_date__month=month,
    ).exists()
    has_payroll = PayrollAllocationLine.objects.filter(
        project_id=project_id,
        batch__status__in=[
            PayrollAllocationStatus.SUBMITTED,
            PayrollAllocationStatus.APPROVED,
        ],
        batch__period_year=year,
        batch__period_month=month,
    ).exists()
    if is_closed:
        return "OK" if has_timesheet and has_payroll else "ESTIMATED"
    return "UNSURE"


def get_labor_totals_for_projects(project_ids, *, as_of_date=None) -> dict:
    if not project_ids:
        return {}
    if as_of_date is None:
        as_of_date = timezone.localdate()
    fallback_cbs = _get_or_create_fallback_cbs()

    totals = {
        int(project_id): {
            "total": Decimal("0"),
            "by_cbs": {},
            "sources": {"timesheet_total": Decimal("0"), "payroll_total": Decimal("0")},
        }
        for project_id in project_ids
    }

    timesheet_rows = (
        TimesheetLine.objects.filter(
            timesheet__project_id__in=project_ids,
            timesheet__status=TimesheetStatus.APPROVED,
            timesheet__work_date__lte=as_of_date,
        )
        .values("timesheet__project_id", "labor_role__default_cbs")
        .annotate(total=Sum("amount"))
    )
    for row in timesheet_rows:
        project_id = int(row["timesheet__project_id"])
        cbs_id = row["labor_role__default_cbs"] or fallback_cbs.id
        amount = row["total"] or Decimal("0")
        entry = totals.setdefault(
            project_id,
            {
                "total": Decimal("0"),
                "by_cbs": {},
                "sources": {"timesheet_total": Decimal("0"), "payroll_total": Decimal("0")},
            },
        )
        entry["total"] += amount
        entry["sources"]["timesheet_total"] += amount
        entry["by_cbs"][cbs_id] = entry["by_cbs"].get(cbs_id, Decimal("0")) + amount

    payroll_rows = (
        PayrollAllocationLine.objects.filter(
            project_id__in=project_ids,
            batch__status__in=[
                PayrollAllocationStatus.SUBMITTED,
                PayrollAllocationStatus.APPROVED,
            ],
        )
        .filter(
            Q(batch__period_year__lt=as_of_date.year)
            | Q(
                batch__period_year=as_of_date.year,
                batch__period_month__lte=as_of_date.month,
            )
        )
        .values("project_id", "cbs_id")
        .annotate(total=Sum("amount"))
    )
    for row in payroll_rows:
        project_id = int(row["project_id"])
        cbs_id = row["cbs_id"] or fallback_cbs.id
        amount = row["total"] or Decimal("0")
        entry = totals.setdefault(
            project_id,
            {
                "total": Decimal("0"),
                "by_cbs": {},
                "sources": {"timesheet_total": Decimal("0"), "payroll_total": Decimal("0")},
            },
        )
        entry["total"] += amount
        entry["sources"]["payroll_total"] += amount
        entry["by_cbs"][cbs_id] = entry["by_cbs"].get(cbs_id, Decimal("0")) + amount

    return totals


def get_labor_actual_by_month(project_id: int, year: int, month: int) -> dict:
    fallback_cbs = _get_or_create_fallback_cbs()
    start_date, end_date = _month_range(year, month)

    timesheet_rows = (
        TimesheetLine.objects.filter(
            timesheet__project_id=project_id,
            timesheet__status=TimesheetStatus.APPROVED,
            timesheet__work_date__gte=start_date,
            timesheet__work_date__lte=end_date,
        )
        .values("labor_role__default_cbs")
        .annotate(total=Sum("amount"))
    )
    payroll_rows = (
        PayrollAllocationLine.objects.filter(
            project_id=project_id,
            batch__status__in=[
                PayrollAllocationStatus.SUBMITTED,
                PayrollAllocationStatus.APPROVED,
            ],
            batch__period_year=year,
            batch__period_month=month,
        )
        .values("cbs_id")
        .annotate(total=Sum("amount"))
    )

    by_cbs = {}
    timesheet_total = Decimal("0")
    payroll_total = Decimal("0")
    for row in timesheet_rows:
        cbs_id = row["labor_role__default_cbs"] or fallback_cbs.id
        amount = row["total"] or Decimal("0")
        timesheet_total += amount
        by_cbs[cbs_id] = by_cbs.get(cbs_id, Decimal("0")) + amount
    for row in payroll_rows:
        cbs_id = row["cbs_id"] or fallback_cbs.id
        amount = row["total"] or Decimal("0")
        payroll_total += amount
        by_cbs[cbs_id] = by_cbs.get(cbs_id, Decimal("0")) + amount

    cbs_ids = list(by_cbs.keys())
    cbs_map = {item.id: item for item in CostItem.objects.filter(id__in=cbs_ids)}
    by_cbs_list = []
    for cbs_id, amount in by_cbs.items():
        cbs = cbs_map.get(cbs_id) or fallback_cbs
        by_cbs_list.append(
            {
                "cbs_id": cbs.id,
                "cbs_code": cbs.code,
                "cbs_name": cbs.get_display_name(),
                "amount": amount,
            }
        )

    total = timesheet_total + payroll_total
    quality = compute_labor_data_quality(project_id, year, month)
    return {
        "total": int(total),
        "by_cbs": by_cbs_list,
        "sources": {
            "timesheet_total": int(timesheet_total),
            "payroll_total": int(payroll_total),
        },
        "data_quality": quality,
    }


def _validate_overlap(*, role, rate_type, scope_type, project, start, end, exclude_id=None):
    qs = LaborRateTable.objects.filter(
        labor_role=role,
        rate_type=rate_type,
        scope_type=scope_type,
    )
    if scope_type == LaborRateScope.PROJECT:
        qs = qs.filter(project=project)
    else:
        qs = qs.filter(project__isnull=True)
    qs = qs.filter(_overlap_q(start, end))
    if exclude_id:
        qs = qs.exclude(id=exclude_id)
    if qs.exists():
        raise ValidationError("동일 범위에서 기간이 겹치는 단가가 이미 존재합니다.")


def _assert_month_open(project: Project | None, target_date: date_type, *, message_prefix: str) -> None:
    guard_write(
        project=project,
        target_date=target_date,
        message_context=message_prefix,
        exc=PermissionDenied,
    )


def _ensure_timesheet_editable(timesheet: Timesheet, *, actor):
    if timesheet.status not in (TimesheetStatus.DRAFT, TimesheetStatus.REJECTED):
        raise PermissionDenied(
            "\uc81c\ucd9c \uc774\ud6c4\uc5d0\ub294 \uc218\uc815\ud560 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4."
        )
    if actor and actor != timesheet.created_by:
        raise PermissionDenied(
            "\ubcf8\uc778\uc774 \uc791\uc131\ud55c \ucd9c\uc5ed\ubd80\ub9cc \uc218\uc815 \uac00\ub2a5\ud569\ub2c8\ub2e4."
        )


def create_labor_role(data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    code = str(data.get("code") or "").strip().upper()
    name = str(data.get("name") or "").strip()
    if not code:
        raise ValidationError({"code": "code is required."})
    if not name:
        raise ValidationError({"name": "name is required."})
    role_group = str(data.get("role_group") or "").strip()
    default_cbs = _resolve_cost_item(data.get("default_cbs"))
    role = LaborRole.objects.create(
        code=code,
        name=name,
        role_group=role_group,
        is_active=bool(data.get("is_active", True)),
        sort_order=int(data.get("sort_order") or 0),
        default_cbs=default_cbs,
    )
    _log_action_safe(
        actor=actor,
        action="LABOR_ROLE_CREATE",
        object_id=role.id,
        summary=f"LaborRole create: {role.code}",
        metadata={"role_id": role.id, "role_code": role.code},
        object_type="LaborRole",
    )
    return role


def update_labor_role(role: LaborRole, data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    before_active = role.is_active
    if "code" in data and data["code"]:
        role.code = str(data["code"]).strip().upper()
    if "name" in data and data["name"] is not None:
        role.name = str(data["name"]).strip()
    if "role_group" in data:
        role.role_group = str(data.get("role_group") or "").strip()
    if "sort_order" in data:
        role.sort_order = int(data.get("sort_order") or 0)
    if "is_active" in data:
        role.is_active = bool(data.get("is_active"))
    if "default_cbs" in data:
        role.default_cbs = _resolve_cost_item(data.get("default_cbs"))
    role.save()
    action = "LABOR_ROLE_UPDATE"
    if before_active and not role.is_active:
        action = "LABOR_ROLE_DEACTIVATE"
    _log_action_safe(
        actor=actor,
        action=action,
        object_id=role.id,
        summary=f"LaborRole update: {role.code}",
        metadata={"role_id": role.id, "role_code": role.code, "is_active": role.is_active},
        object_type="LaborRole",
    )
    return role


def create_worker_master(data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    name = str(data.get("name") or "").strip()
    rrn = str(data.get("rrn") or "").strip()
    if not name:
        raise ValidationError({"name": "성명을 입력해 주세요."})
    if not rrn:
        raise ValidationError({"rrn": "주민등록번호를 입력해 주세요."})
    default_labor_role = data.get("default_labor_role")
    if default_labor_role and not isinstance(default_labor_role, LaborRole):
        default_labor_role = LaborRole.objects.filter(id=default_labor_role).first()
    worker = WorkerMaster.objects.create(
        name=name,
        rrn_encrypted=_store_sensitive_value(rrn),
        rrn_masked=_mask_rrn_value(rrn),
        identity_hash=_build_identity_hash(rrn),
        phone=str(data.get("phone") or "").strip(),
        address=str(data.get("address") or "").strip(),
        bank_name=str(data.get("bank_name") or "").strip(),
        bank_code=str(data.get("bank_code") or "").strip(),
        account_number_encrypted=_store_sensitive_value(data.get("account_number") or ""),
        account_holder=str(data.get("account_holder") or "").strip(),
        nationality_code=str(data.get("nationality_code") or "").strip().upper(),
        visa_code=str(data.get("visa_code") or "").strip().upper(),
        comwel_job_code=str(data.get("comwel_job_code") or "").strip().upper(),
        cwma_job_name=str(data.get("cwma_job_name") or "").strip(),
        default_labor_role=default_labor_role,
        retirement_deduction_eligible=bool(data.get("retirement_deduction_eligible")),
        active=bool(data.get("active", True)),
    )
    try:
        log_action(
            actor=actor,
            action="LABOR_WORKER_CREATE",
            object_type="WorkerMaster",
            object_id=worker.id,
            after=_worker_audit_snapshot(worker),
            meta={"worker_id": worker.id, "worker_name": worker.name},
        )
    except Exception:
        logger.warning("AuditLog failed for LABOR_WORKER_CREATE snapshot.", exc_info=True)
    return worker


def update_worker_master(worker: WorkerMaster, data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    before = _worker_audit_snapshot(worker)
    worker.name = str(data.get("name") or worker.name or "").strip()
    worker.phone = str(data.get("phone") or "").strip()
    worker.address = str(data.get("address") or "").strip()
    worker.bank_name = str(data.get("bank_name") or "").strip()
    worker.bank_code = str(data.get("bank_code") or "").strip()
    worker.account_holder = str(data.get("account_holder") or "").strip()
    worker.nationality_code = str(data.get("nationality_code") or "").strip().upper()
    worker.visa_code = str(data.get("visa_code") or "").strip().upper()
    worker.comwel_job_code = str(data.get("comwel_job_code") or "").strip().upper()
    worker.cwma_job_name = str(data.get("cwma_job_name") or "").strip()
    worker.retirement_deduction_eligible = bool(data.get("retirement_deduction_eligible"))
    worker.active = bool(data.get("active", True))
    default_labor_role = data.get("default_labor_role")
    if default_labor_role and not isinstance(default_labor_role, LaborRole):
        default_labor_role = LaborRole.objects.filter(id=default_labor_role).first()
    worker.default_labor_role = default_labor_role

    rrn = str(data.get("rrn") or "").strip()
    if rrn:
        worker.rrn_encrypted = _store_sensitive_value(rrn)
        worker.rrn_masked = _mask_rrn_value(rrn)
        worker.identity_hash = _build_identity_hash(rrn)

    account_number = str(data.get("account_number") or "").strip()
    if account_number:
        worker.account_number_encrypted = _store_sensitive_value(account_number)

    if not worker.name:
        raise ValidationError({"name": "성명을 입력해 주세요."})
    if not worker.rrn_encrypted:
        raise ValidationError({"rrn": "주민등록번호를 입력해 주세요."})

    worker.save()
    after = _worker_audit_snapshot(worker)
    try:
        log_action(
            actor=actor,
            action="LABOR_WORKER_UPDATE",
            object_type="WorkerMaster",
            object_id=worker.id,
            before=before,
            after=after,
            meta={"worker_id": worker.id, "worker_name": worker.name},
        )
    except Exception:
        logger.warning("AuditLog failed for LABOR_WORKER_UPDATE snapshot.", exc_info=True)
    return worker


def _worker_usage_breakdown(worker: WorkerMaster) -> dict:
    return {
        "work_ledgers": LaborWorkLedger.objects.filter(worker=worker).count(),
        "electronic_card_raw_rows": ElectronicCardWorkRaw.objects.filter(matched_worker=worker).count(),
        "electronic_card_day_rows": ElectronicCardWorkDay.objects.filter(worker=worker).count(),
        "reconciliation_results": LaborReconciliationResult.objects.filter(worker=worker).count(),
        "confirmed_work_days": LaborConfirmedWorkDay.objects.filter(worker=worker).count(),
        "monthly_payrolls": LaborMonthlyPayroll.objects.filter(worker=worker).count(),
    }


def _worker_usage_total(worker: WorkerMaster) -> int:
    return sum(_worker_usage_breakdown(worker).values())


def delete_or_deactivate_worker_master(worker: WorkerMaster, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    before = _worker_audit_snapshot(worker)
    usage_breakdown = _worker_usage_breakdown(worker)
    usage_count = sum(usage_breakdown.values())

    if not worker.active:
        _log_action_safe(
            actor=actor,
            action=LABOR_WORKER_ALREADY_INACTIVE,
            object_id=worker.id,
            summary=f"WorkerMaster already inactive: {worker.id}",
            metadata={
                "worker_id": worker.id,
                "worker_name": worker.name,
                "rrn_masked": worker.rrn_masked,
                "active_before": before["active"],
                "active_after": worker.active,
                "usage_count": usage_count,
                "usage_breakdown": usage_breakdown,
            },
            object_type="WorkerMaster",
        )
        return {"action": "already_inactive", "usage_count": usage_count}

    if usage_count == 0:
        worker_id = worker.id
        worker_name = worker.name
        rrn_masked = worker.rrn_masked
        worker.delete()
        _log_action_safe(
            actor=actor,
            action=LABOR_WORKER_DELETE,
            object_id=worker_id,
            summary=f"WorkerMaster delete: {worker_id}",
            metadata={
                "worker_id": worker_id,
                "worker_name": worker_name,
                "rrn_masked": rrn_masked,
                "active_before": before["active"],
                "active_after": None,
                "usage_count": 0,
                "usage_breakdown": usage_breakdown,
                "deleted": True,
            },
            object_type="WorkerMaster",
        )
        return {"action": "deleted", "usage_count": 0}

    worker.active = False
    worker.save(update_fields=["active", "updated_at"])
    after = _worker_audit_snapshot(worker)
    _log_action_safe(
        actor=actor,
        action=LABOR_WORKER_DEACTIVATE,
        object_id=worker.id,
        summary=f"WorkerMaster deactivate: {worker.id}",
        metadata={
            "worker_id": worker.id,
            "worker_name": worker.name,
            "rrn_masked": worker.rrn_masked,
            "active_before": before["active"],
            "active_after": after["active"],
            "usage_count": usage_count,
            "usage_breakdown": usage_breakdown,
            "deactivated": True,
        },
        object_type="WorkerMaster",
    )
    return {"action": "deactivated", "usage_count": usage_count}


def _coerce_decimal(value, *, field_name: str):
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except Exception as exc:
        raise ValidationError({field_name: f"{field_name} 값이 올바르지 않습니다."}) from exc


def _coerce_int(value, *, field_name: str):
    if value in (None, ""):
        return 0
    try:
        return int(value)
    except Exception as exc:
        raise ValidationError({field_name: f"{field_name} 값이 올바르지 않습니다."}) from exc


def _normalize_e_card_header(value) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    return re.sub(r"[\s\r\n\t()\-_/]+", "", text).lower()


def _parse_year_month_value(value):
    text = str(value or "").strip()
    if not text:
        return None
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) < 6:
        return None
    year = int(digits[:4])
    month = int(digits[4:6])
    if 1 <= month <= 12:
        return date_type(year, month, 1)
    return None


def _resolve_header_field(header_value):
    normalized = _normalize_e_card_header(header_value)
    if not normalized:
        return None
    for field_name, candidates in E_CARD_HEADER_LABELS.items():
        if normalized in {_normalize_e_card_header(candidate) for candidate in candidates}:
            return field_name
    return None


def _resolve_day_column(header_value):
    normalized = _normalize_e_card_header(header_value)
    if not normalized:
        return None
    if normalized.isdigit():
        day = int(normalized)
        if 1 <= day <= 31:
            return day
    match = re.fullmatch(r"(\d{1,2})일", normalized)
    if match:
        day = int(match.group(1))
        if 1 <= day <= 31:
            return day
    return None


def _resolve_header_map(header_row):
    field_map = {}
    day_columns = {}
    for col_index, value in enumerate(header_row):
        field_name = _resolve_header_field(value)
        if field_name and field_name not in field_map:
            field_map[field_name] = col_index
        day_value = _resolve_day_column(value)
        if day_value and day_value not in day_columns:
            day_columns[day_value] = col_index
    return field_map, day_columns


def _first_visible_sheet(workbook):
    active = workbook.active
    if getattr(active, "sheet_state", "visible") == "visible":
        return active
    for sheet in workbook.worksheets:
        if getattr(sheet, "sheet_state", "visible") == "visible":
            return sheet
    return workbook.worksheets[0]


def _build_e_card_header_summary(uploaded_file, *, selected_month):
    uploaded_file.seek(0)
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    sheet = _first_visible_sheet(workbook)
    rows = list(sheet.iter_rows(min_row=1, max_row=200, values_only=True))

    header_row_index = None
    header_field_map = {}
    day_columns = {}
    header_samples = []
    best_score = -1

    for row_index, row in enumerate(rows[:20], start=1):
        field_map = {}
        local_day_columns = {}
        for col_index, value in enumerate(row):
            field_name = _resolve_header_field(value)
            if field_name and field_name not in field_map:
                field_map[field_name] = col_index
            day_value = _resolve_day_column(value)
            if day_value and day_value not in local_day_columns:
                local_day_columns[day_value] = col_index
        score = len(field_map) + len(local_day_columns)
        if score > best_score:
            best_score = score
            header_row_index = row_index
            header_field_map = field_map
            day_columns = local_day_columns
            header_samples = [str(cell or "").strip() for cell in row if str(cell or "").strip()]

    if not header_row_index or best_score <= 0:
        raise ValidationError("전자카드 파일에서 헤더 행을 찾지 못했습니다.")

    missing_essential = [
        field_name for field_name in sorted(E_CARD_ESSENTIAL_FIELDS) if field_name not in header_field_map
    ]
    if missing_essential:
        missing_labels = [
            E_CARD_HEADER_LABELS[field_name][0] for field_name in missing_essential
        ]
        raise ValidationError(
            f"필수 헤더가 누락되었습니다: {', '.join(missing_labels)}"
        )

    warnings = []
    missing_optional = [
        field_name for field_name in sorted(E_CARD_OPTIONAL_FIELDS) if field_name not in header_field_map
    ]
    if missing_optional:
        missing_labels = [
            E_CARD_HEADER_LABELS[field_name][0] for field_name in missing_optional
        ]
        warnings.append(
            f"선택 헤더가 누락되었습니다: {', '.join(missing_labels)}"
        )
    if len(day_columns) < 28:
        warnings.append("일자 컬럼(1일~31일) 구성이 불완전합니다. 원본 파일 형식을 확인해 주세요.")

    workbook_months = []
    if "year_month" in header_field_map:
        month_col = header_field_map["year_month"]
        seen = set()
        for row in rows[header_row_index:]:
            if month_col >= len(row):
                continue
            parsed_month = _parse_year_month_value(row[month_col])
            if parsed_month and parsed_month not in seen:
                seen.add(parsed_month)
                workbook_months.append(parsed_month)
        workbook_months = sorted(workbook_months)

    if workbook_months:
        if len(workbook_months) == 1 and workbook_months[0] != selected_month:
            raise ValidationError(
                f"파일의 근로연월({workbook_months[0]:%Y-%m})이 선택한 기준월({selected_month:%Y-%m})과 다릅니다."
            )
        if selected_month not in workbook_months:
            warnings.append(
                "파일 안의 근로연월 값이 선택한 기준월과 완전히 일치하지 않습니다."
            )

    def _first_data_value(field_name):
        col_index = header_field_map.get(field_name)
        if col_index is None:
            return ""
        for row in rows[header_row_index:]:
            if col_index >= len(row):
                continue
            value = str(row[col_index] or "").strip()
            if value:
                return value
        return ""

    return {
        "sheet_name": sheet.title,
        "header_row_index": header_row_index,
        "detected_headers": {
            field_name: E_CARD_HEADER_LABELS[field_name][0]
            for field_name in header_field_map.keys()
        },
        "missing_optional_headers": [
            E_CARD_HEADER_LABELS[field_name][0] for field_name in missing_optional
        ],
        "day_column_count": len(day_columns),
        "day_columns": sorted(day_columns.keys()),
        "warnings": warnings,
        "header_samples": header_samples[:12],
        "workbook_months": [value.strftime("%Y-%m") for value in workbook_months],
        "cwma_project_name": _first_data_value("project_name"),
        "deduction_join_no": _first_data_value("deduction_join_no"),
        "company_name": _first_data_value("company_name"),
    }


def _parse_decimal_or_zero(value):
    text = str(value or "").strip().replace(",", "")
    if not text:
        return Decimal("0")
    try:
        return Decimal(text)
    except Exception:
        return Decimal("0")


def _normalize_phone(value):
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def _normalize_worker_name(value):
    return "".join(str(value or "").split())


def _coerce_cell_text(value):
    return str(value or "").strip()


def _resolve_card_value(value):
    text = _coerce_cell_text(value)
    if not text:
        return Decimal("0"), ""
    numeric = _parse_decimal_or_zero(text)
    if numeric > 0:
        return numeric, ""
    normalized = text.upper()
    if normalized in {"Y", "O", "○", "출", "근무", "1"}:
        return Decimal("1"), ""
    return Decimal("0"), f"일자값 해석 불가: {text}"


def _find_worker_match(*, identity_hash, worker_name, phone):
    if identity_hash:
        worker = WorkerMaster.objects.filter(identity_hash=identity_hash, active=True).first()
        if worker:
            return worker, "identity"
    normalized_name = _normalize_worker_name(worker_name)
    if normalized_name:
        normalized_phone = _normalize_phone(phone)
        candidates = WorkerMaster.objects.filter(active=True).order_by("id")
        for candidate in candidates:
            if _normalize_worker_name(candidate.name) != normalized_name:
                continue
            candidate_phone = _normalize_phone(candidate.phone)
            if normalized_phone and candidate_phone == normalized_phone:
                return candidate, "name_phone"
    return None, ""


def _row_has_parse_target(row, field_map):
    for field_name in ("worker_name", "resident_no", "project_name", "phone", "job_type"):
        col_index = field_map.get(field_name)
        if col_index is None or col_index >= len(row):
            continue
        if _coerce_cell_text(row[col_index]):
            return True
    return False


def parse_electronic_card_import_batch(batch, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(batch, ElectronicCardImportBatch):
        batch = ElectronicCardImportBatch.objects.filter(id=batch).first()
    if batch is None:
        raise ValidationError("전자카드 업로드 배치를 찾을 수 없습니다.")
    if batch.status in {
        ElectronicCardImportBatchStatus.CONFIRMED,
        ElectronicCardImportBatchStatus.DISCARDED,
    }:
        raise ValidationError("확정 또는 폐기된 배치는 다시 파싱할 수 없습니다.")
    if not batch.source_file:
        raise ValidationError("원본 파일이 없어 파싱할 수 없습니다.")

    batch.source_file.open("rb")
    try:
        workbook = load_workbook(batch.source_file, read_only=True, data_only=True)
    except Exception as exc:
        logger.exception("Failed to open electronic card workbook for parse.")
        raise ValidationError("전자카드 원본 파일을 열지 못했습니다.") from exc

    sheet = _first_visible_sheet(workbook)
    rows = list(sheet.iter_rows(min_row=1, max_row=5000, values_only=True))
    summary = batch.header_check_summary or {}
    header_row_index = int(summary.get("header_row_index") or 0)
    if header_row_index <= 0 or header_row_index > len(rows):
        _, _, _, header_row_index = None, None, None, 0
        best_score = -1
        for row_index, row in enumerate(rows[:30], start=1):
            field_map, day_columns = _resolve_header_map(row)
            score = len(field_map) + len(day_columns)
            if score > best_score:
                best_score = score
                header_row_index = row_index
    if header_row_index <= 0:
        raise ValidationError("전자카드 헤더 행을 찾지 못했습니다.")

    header_row = rows[header_row_index - 1]
    field_map, day_columns = _resolve_header_map(header_row)
    missing_essential = [
        field_name for field_name in sorted(E_CARD_ESSENTIAL_FIELDS) if field_name not in field_map
    ]
    if missing_essential:
        labels = [E_CARD_HEADER_LABELS[name][0] for name in missing_essential]
        raise ValidationError(f"필수 헤더가 누락되었습니다: {', '.join(labels)}")

    raw_rows_to_create = []
    day_rows_to_create = []
    matched_worker_count = 0
    unmatched_worker_count = 0
    matched_by_identity_count = 0
    matched_by_name_phone_count = 0
    unmatched_preview = []

    for excel_row_no, row in enumerate(rows[header_row_index:], start=header_row_index + 1):
        if not _row_has_parse_target(row, field_map):
            continue

        work_month = batch.year_month
        month_col = field_map.get("year_month")
        if month_col is not None and month_col < len(row):
            parsed_month = _parse_year_month_value(row[month_col])
            if parsed_month:
                work_month = parsed_month

        worker_name_raw = _coerce_cell_text(row[field_map["worker_name"]]) if field_map.get("worker_name") is not None and field_map["worker_name"] < len(row) else ""
        rrn_raw = _coerce_cell_text(row[field_map["resident_no"]]) if field_map.get("resident_no") is not None and field_map["resident_no"] < len(row) else ""
        phone = _coerce_cell_text(row[field_map["phone"]]) if field_map.get("phone") is not None and field_map["phone"] < len(row) else ""
        identity_hash = WorkerMaster.make_identity_hash(rrn_raw) if rrn_raw else ""
        rrn_masked = WorkerMaster.mask_rrn(rrn_raw) if rrn_raw else ""
        rrn_encrypted = _protect_sensitive_value(rrn_raw, "rrn") if rrn_raw else ""
        matched_worker, matched_by = _find_worker_match(
            identity_hash=identity_hash,
            worker_name=worker_name_raw,
            phone=phone,
        )
        match_status = (
            ElectronicCardMatchStatus.MATCHED
            if matched_worker
            else ElectronicCardMatchStatus.UNMATCHED
        )
        if matched_worker:
            matched_worker_count += 1
            if matched_by == "identity":
                matched_by_identity_count += 1
            elif matched_by == "name_phone":
                matched_by_name_phone_count += 1
        else:
            unmatched_worker_count += 1
            if len(unmatched_preview) < 10:
                unmatched_preview.append(
                    {
                        "row_no": excel_row_no,
                        "worker_name_raw": worker_name_raw,
                        "rrn_masked": rrn_masked,
                        "phone": phone,
                    }
                )

        day_values = {}
        day_notes = []
        for day in range(1, 32):
            col_index = day_columns.get(day)
            cell_value = row[col_index] if col_index is not None and col_index < len(row) else ""
            cell_text = _coerce_cell_text(cell_value)
            day_values[f"day_{day:02d}"] = cell_text

        raw = ElectronicCardWorkRaw(
            batch=batch,
            row_no=excel_row_no,
            work_month=work_month,
            project_name_raw=_coerce_cell_text(row[field_map["project_name"]]) if field_map.get("project_name") is not None and field_map["project_name"] < len(row) else "",
            deduction_join_no=_coerce_cell_text(row[field_map["deduction_join_no"]]) if field_map.get("deduction_join_no") is not None and field_map["deduction_join_no"] < len(row) else "",
            company_name=_coerce_cell_text(row[field_map["company_name"]]) if field_map.get("company_name") is not None and field_map["company_name"] < len(row) else "",
            worker_name_raw=worker_name_raw,
            rrn_encrypted=rrn_encrypted,
            rrn_masked=rrn_masked,
            identity_hash=identity_hash,
            phone=phone,
            job_name_raw=_coerce_cell_text(row[field_map["job_type"]]) if field_map.get("job_type") is not None and field_map["job_type"] < len(row) else "",
            card_issued=_coerce_cell_text(row[field_map["card_issued"]]) if field_map.get("card_issued") is not None and field_map["card_issued"] < len(row) else "",
            retirement_target=_coerce_cell_text(row[field_map["retirement_deduction"]]) if field_map.get("retirement_deduction") is not None and field_map["retirement_deduction"] < len(row) else "",
            exclusion_reason=_coerce_cell_text(row[field_map["exclusion_reason"]]) if field_map.get("exclusion_reason") is not None and field_map["exclusion_reason"] < len(row) else "",
            work_status=_coerce_cell_text(row[field_map["work_history_status"]]) if field_map.get("work_history_status") is not None and field_map["work_history_status"] < len(row) else "",
            report_status=_coerce_cell_text(row[field_map["report_status"]]) if field_map.get("report_status") is not None and field_map["report_status"] < len(row) else "",
            error_message=_coerce_cell_text(row[field_map["error_message"]]) if field_map.get("error_message") is not None and field_map["error_message"] < len(row) else "",
            auto_work_days=_coerce_cell_text(row[field_map["auto_work_days"]]) if field_map.get("auto_work_days") is not None and field_map["auto_work_days"] < len(row) else "",
            reported_days=_coerce_cell_text(row[field_map["reported_days"]]) if field_map.get("reported_days") is not None and field_map["reported_days"] < len(row) else "",
            declaration_days=_coerce_cell_text(row[field_map["declaration_days"]]) if field_map.get("declaration_days") is not None and field_map["declaration_days"] < len(row) else "",
            confirmed_days=_coerce_cell_text(row[field_map["confirmed_days"]]) if field_map.get("confirmed_days") is not None and field_map["confirmed_days"] < len(row) else "",
            note=_coerce_cell_text(row[field_map["note"]]) if field_map.get("note") is not None and field_map["note"] < len(row) else "",
            matched_worker=matched_worker,
            match_status=match_status,
            **day_values,
        )
        raw_rows_to_create.append(raw)

    with transaction.atomic():
        batch.day_rows.all().delete()
        batch.raw_rows.all().delete()
        created_raws = ElectronicCardWorkRaw.objects.bulk_create(raw_rows_to_create)
        created_day_count = 0
        for raw in created_raws:
            last_day = monthrange(raw.work_month.year, raw.work_month.month)[1]
            for day in range(1, last_day + 1):
                raw_value = getattr(raw, f"day_{day:02d}")
                card_value, note = _resolve_card_value(raw_value)
                day_rows_to_create.append(
                    ElectronicCardWorkDay(
                        raw=raw,
                        batch=batch,
                        worker=raw.matched_worker,
                        work_date=date_type(raw.work_month.year, raw.work_month.month, day),
                        card_value=card_value,
                        is_worked=card_value > 0,
                        card_project=batch.project,
                        match_status=raw.match_status,
                        note=note,
                    )
                )
            created_day_count += last_day
        ElectronicCardWorkDay.objects.bulk_create(day_rows_to_create)
        parsed_worked_day_count = 0
        parsed_zero_day_count = 0
        parsed_unreadable_day_count = 0
        worked_day_preview = []
        for day_row in day_rows_to_create:
            if day_row.note and "해석 불가" in day_row.note:
                parsed_unreadable_day_count += 1
            if day_row.card_value > 0:
                parsed_worked_day_count += 1
                if len(worked_day_preview) < 10:
                    worked_day_preview.append(
                        {
                            "work_date": day_row.work_date.isoformat(),
                            "worker_name_raw": day_row.raw.worker_name_raw,
                            "card_value": str(day_row.card_value),
                            "match_status": day_row.match_status,
                        }
                    )
            elif day_row.card_value == 0:
                parsed_zero_day_count += 1

        parse_summary = {
            "parsed_raw_count": len(created_raws),
            "parsed_day_count": len(day_rows_to_create),
            "parsed_worked_day_count": parsed_worked_day_count,
            "parsed_zero_day_count": parsed_zero_day_count,
            "parsed_unreadable_day_count": parsed_unreadable_day_count,
            "worked_day_preview": worked_day_preview,
            "matched_worker_count": matched_worker_count,
            "matched_by_identity_count": matched_by_identity_count,
            "matched_by_name_phone_count": matched_by_name_phone_count,
            "unmatched_worker_count": unmatched_worker_count,
            "unmatched_preview": unmatched_preview,
            "parsed_at": timezone.now().isoformat(),
        }
        updated_summary = {**summary, **parse_summary}
        batch.header_check_summary = updated_summary
        batch.status = ElectronicCardImportBatchStatus.PARSE_READY
        batch.save(update_fields=["header_check_summary", "status", "updated_at"])

        _log_action_safe(
            actor=actor,
            action=ELECTRONIC_CARD_IMPORT_PARSE,
            object_id=batch.id,
            summary=f"ElectronicCardImportBatch parse: {batch.id}",
            metadata={
                "batch_id": batch.id,
                "year_month": batch.year_month.isoformat(),
                "project_id": batch.project_id,
                "raw_count": len(created_raws),
                "day_count": len(day_rows_to_create),
                "matched_worker_count": matched_worker_count,
                "unmatched_worker_count": unmatched_worker_count,
            },
            object_type="ElectronicCardImportBatch",
        )
    return batch


def _ledger_key(worker_id, work_date):
    return worker_id, work_date


def reconcile_electronic_card_import_batch(batch, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(batch, ElectronicCardImportBatch):
        batch = ElectronicCardImportBatch.objects.filter(id=batch).first()
    if batch is None:
        raise ValidationError("전자카드 업로드 배치를 찾을 수 없습니다.")
    if batch.status in {
        ElectronicCardImportBatchStatus.CONFIRMED,
        ElectronicCardImportBatchStatus.DISCARDED,
    }:
        raise ValidationError("확정 또는 폐기된 배치는 대사할 수 없습니다.")
    if not batch.day_rows.exists():
        raise ValidationError("먼저 파싱을 실행해 주세요.")

    card_days = list(
        batch.day_rows.select_related("worker", "card_project", "raw").order_by("work_date", "id")
    )
    reconcile_card_worked_day_count = sum(1 for row in card_days if Decimal(row.card_value or 0) > 0)

    report_rows = (
        LaborWorkLedger.objects.filter(
            work_month=batch.year_month,
            report_project=batch.project,
        )
        .values("worker_id", "work_date")
        .annotate(total=Sum("work_unit"))
    )
    reconcile_erp_ledger_count = len(report_rows)
    actual_rows = (
        LaborWorkLedger.objects.filter(
            work_month=batch.year_month,
            actual_project=batch.project,
        )
        .values("worker_id", "work_date")
        .annotate(total=Sum("work_unit"))
    )
    report_map = {
        _ledger_key(row["worker_id"], row["work_date"]): Decimal(row["total"] or 0)
        for row in report_rows
        if row["worker_id"] and Decimal(row["total"] or 0) > 0
    }
    actual_map = {
        _ledger_key(row["worker_id"], row["work_date"]): Decimal(row["total"] or 0)
        for row in actual_rows
        if row["worker_id"] and Decimal(row["total"] or 0) > 0
    }
    reconcile_erp_worked_count = len(set(report_map.keys()) | set(actual_map.keys()))

    results = []
    status_counts = {
        LaborReconciliationStatus.MATCH: 0,
        LaborReconciliationStatus.ERP_ONLY: 0,
        LaborReconciliationStatus.CARD_ONLY: 0,
        LaborReconciliationStatus.DIFF: 0,
        LaborReconciliationStatus.UNMATCHED: 0,
    }

    def register_result(**kwargs):
        status = kwargs["status"]
        status_counts[status] = status_counts.get(status, 0) + 1
        results.append(LaborReconciliationResult(**kwargs))

    for card_day in card_days:
        card_unit = Decimal(card_day.card_value or 0)
        if card_day.worker_id is None or card_day.match_status == ElectronicCardMatchStatus.UNMATCHED:
            if card_unit <= 0:
                continue
            register_result(
                batch=batch,
                year_month=batch.year_month,
                project=batch.project,
                worker=None,
                work_date=card_day.work_date,
                erp_work_unit=Decimal("0"),
                card_work_unit=card_unit,
                difference=Decimal("0") - card_unit,
                status=LaborReconciliationStatus.UNMATCHED,
                resolution="",
                final_work_unit=None,
                export_included=True,
                export_value=None,
                card_day=card_day,
                export_note="",
            )
            continue

        key = _ledger_key(card_day.worker_id, card_day.work_date)
        erp_unit = report_map.pop(key, None)
        if erp_unit is None:
            erp_unit = actual_map.pop(key, Decimal("0"))
        else:
            actual_map.pop(key, None)

        if card_unit <= 0 and erp_unit <= 0:
            continue
        if card_unit > 0 and erp_unit > 0:
            status = (
                LaborReconciliationStatus.MATCH
                if erp_unit == card_unit
                else LaborReconciliationStatus.DIFF
            )
        elif erp_unit > 0:
            status = LaborReconciliationStatus.ERP_ONLY
        else:
            status = LaborReconciliationStatus.CARD_ONLY

        register_result(
            batch=batch,
            year_month=batch.year_month,
            project=batch.project,
            worker=card_day.worker,
            work_date=card_day.work_date,
            erp_work_unit=erp_unit,
            card_work_unit=card_unit,
            difference=erp_unit - card_unit,
            status=status,
            resolution=(
                LaborReconciliationResolution.ERP
                if status == LaborReconciliationStatus.MATCH
                else ""
            ),
            final_work_unit=erp_unit if status == LaborReconciliationStatus.MATCH else None,
            export_included=True,
            export_value=erp_unit if status == LaborReconciliationStatus.MATCH else None,
            export_note="",
            card_day=card_day,
        )

    remaining_keys = set(report_map.keys()) | set(actual_map.keys())
    workers = {
        worker.id: worker
        for worker in WorkerMaster.objects.filter(id__in=[key[0] for key in remaining_keys if key[0]])
    }
    for key in sorted(remaining_keys, key=lambda item: (item[1], item[0] or 0)):
        erp_unit = report_map.get(key)
        if erp_unit is None:
            erp_unit = actual_map.get(key, Decimal("0"))
        if erp_unit <= 0:
            continue
        worker_id, work_date = key
        register_result(
            batch=batch,
            year_month=batch.year_month,
            project=batch.project,
            worker=workers.get(worker_id),
            work_date=work_date,
            erp_work_unit=erp_unit,
            card_work_unit=Decimal("0"),
            difference=erp_unit,
            status=LaborReconciliationStatus.ERP_ONLY,
            resolution="",
            final_work_unit=None,
            export_included=True,
            export_value=None,
            export_note="",
            card_day=None,
        )

    with transaction.atomic():
        batch.reconciliation_results.all().delete()
        LaborReconciliationResult.objects.bulk_create(results)
        no_result_reason = ""
        if not results:
            if reconcile_card_worked_day_count == 0 and reconcile_erp_worked_count == 0:
                no_result_reason = "전자카드 출역값과 ERP 원장이 모두 없습니다."
            elif reconcile_card_worked_day_count > 0 and reconcile_erp_worked_count == 0:
                no_result_reason = "전자카드 출역값은 있으나 대사 결과가 생성되지 않았습니다. 점검 필요."
            elif reconcile_card_worked_day_count == 0 and reconcile_erp_worked_count > 0:
                no_result_reason = "ERP 원장은 있으나 대사 결과가 생성되지 않았습니다. 점검 필요."

        updated_summary = {
            **(batch.header_check_summary or {}),
            "reconcile_card_worked_day_count": reconcile_card_worked_day_count,
            "reconcile_erp_ledger_count": reconcile_erp_ledger_count,
            "reconcile_erp_worked_count": reconcile_erp_worked_count,
            "reconciliation_total_count": len(results),
            "reconciliation_match_count": status_counts[LaborReconciliationStatus.MATCH],
            "reconciliation_erp_only_count": status_counts[LaborReconciliationStatus.ERP_ONLY],
            "reconciliation_card_only_count": status_counts[LaborReconciliationStatus.CARD_ONLY],
            "reconciliation_diff_count": status_counts[LaborReconciliationStatus.DIFF],
            "reconciliation_unmatched_count": status_counts[LaborReconciliationStatus.UNMATCHED],
            "reconciliation_no_result_reason": no_result_reason,
            "reconciled_at": timezone.now().isoformat(),
        }
        batch.header_check_summary = updated_summary
        batch.save(update_fields=["header_check_summary", "updated_at"])
        _log_action_safe(
            actor=actor,
            action=ELECTRONIC_CARD_RECONCILE,
            object_id=batch.id,
            summary=f"ElectronicCardImportBatch reconcile: {batch.id}",
            metadata={
                "batch_id": batch.id,
                "year_month": batch.year_month.isoformat(),
                "project_id": batch.project_id,
                "total_count": len(results),
                "match_count": status_counts[LaborReconciliationStatus.MATCH],
                "erp_only_count": status_counts[LaborReconciliationStatus.ERP_ONLY],
                "card_only_count": status_counts[LaborReconciliationStatus.CARD_ONLY],
                "diff_count": status_counts[LaborReconciliationStatus.DIFF],
                "unmatched_count": status_counts[LaborReconciliationStatus.UNMATCHED],
            },
            object_type="ElectronicCardImportBatch",
        )
    return batch


def _resolve_reconciliation_action(action):
    value = str(action or "").strip().upper()
    mapping = {
        "USE_ERP": LaborReconciliationResolution.ERP,
        "ERP": LaborReconciliationResolution.ERP,
        "USE_CARD": LaborReconciliationResolution.CARD,
        "CARD": LaborReconciliationResolution.CARD,
        "MANUAL": LaborReconciliationResolution.MANUAL,
        "EXCLUDE": LaborReconciliationResolution.EXCLUDED,
        "EXCLUDED": LaborReconciliationResolution.EXCLUDED,
    }
    return mapping.get(value, value)


def _ensure_reconciliation_batch_editable(batch):
    if batch.status in {
        ElectronicCardImportBatchStatus.CONFIRMED,
        ElectronicCardImportBatchStatus.DISCARDED,
    }:
        raise ValidationError("확정 또는 폐기된 배치는 수정할 수 없습니다.")


def resolve_labor_reconciliation_result(
    result,
    action,
    actor,
    final_work_unit=None,
    comment="",
    export_note="",
):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(result, LaborReconciliationResult):
        result = LaborReconciliationResult.objects.select_related("batch").filter(id=result).first()
    if result is None:
        raise ValidationError("대사 결과를 찾을 수 없습니다.")
    _ensure_reconciliation_batch_editable(result.batch)

    resolution = _resolve_reconciliation_action(action)
    comment = str(comment or "").strip()
    export_note = str(export_note or "").strip()
    original_status = result.status
    requires_comment = original_status in {
        LaborReconciliationStatus.ERP_ONLY,
        LaborReconciliationStatus.CARD_ONLY,
        LaborReconciliationStatus.DIFF,
        LaborReconciliationStatus.UNMATCHED,
    }
    if requires_comment and not comment:
        raise ValidationError("비일치 대사 결과는 HQ 의견을 입력해 주세요.")

    new_final_work_unit = None
    export_included = True
    export_value = None
    if resolution == LaborReconciliationResolution.ERP:
        new_final_work_unit = result.erp_work_unit
        export_value = result.erp_work_unit
    elif resolution == LaborReconciliationResolution.CARD:
        new_final_work_unit = result.card_work_unit
        export_value = result.card_work_unit
    elif resolution == LaborReconciliationResolution.MANUAL:
        if final_work_unit in (None, ""):
            raise ValidationError("수동 조정은 최종 공수를 입력해 주세요.")
        new_final_work_unit = Decimal(str(final_work_unit))
        if new_final_work_unit < 0:
            raise ValidationError("최종 공수는 0 이상이어야 합니다.")
        export_value = new_final_work_unit
    elif resolution == LaborReconciliationResolution.EXCLUDED:
        if not comment:
            raise ValidationError("제외 처리 사유를 입력해 주세요.")
        new_final_work_unit = Decimal("0")
        export_included = False
        export_value = Decimal("0")
    else:
        raise ValidationError("처리 방식을 확인해 주세요.")

    result.resolution = resolution
    result.final_work_unit = new_final_work_unit
    result.export_included = export_included
    result.export_value = export_value
    result.export_note = export_note
    result.hq_comment = comment
    result.resolved_by = actor
    result.resolved_at = timezone.now()
    if original_status != LaborReconciliationStatus.MATCH:
        result.status = LaborReconciliationStatus.RESOLVED
    result.save(
        update_fields=[
            "status",
            "resolution",
            "final_work_unit",
            "hq_comment",
            "resolved_by",
            "resolved_at",
            "export_included",
            "export_value",
            "export_note",
            "updated_at",
        ]
    )
    _log_action_safe(
        actor=actor,
        action=LABOR_RECONCILIATION_RESOLVE,
        object_id=result.id,
        summary=f"Labor reconciliation resolve: {result.id}",
        metadata={
            "result_id": result.id,
            "batch_id": result.batch_id,
            "old_status": original_status,
            "new_status": result.status,
            "resolution": result.resolution,
            "erp_work_unit": str(result.erp_work_unit),
            "card_work_unit": str(result.card_work_unit),
            "final_work_unit": str(result.final_work_unit or ""),
            "export_included": result.export_included,
        },
        object_type="LaborReconciliationResult",
    )
    return result


def bulk_resolve_labor_reconciliation_results(
    result_ids,
    action,
    actor,
    comment,
    final_work_unit=None,
):
    require_role(actor, [Role.HQ, Role.CEO])
    results = list(LaborReconciliationResult.objects.filter(id__in=result_ids).select_related("batch"))
    if not results:
        raise ValidationError("선택한 대사 결과가 없습니다.")
    with transaction.atomic():
        for result in results:
            resolve_labor_reconciliation_result(
                result,
                action,
                actor,
                final_work_unit=final_work_unit,
                comment=comment,
            )
    _log_action_safe(
        actor=actor,
        action=LABOR_RECONCILIATION_BULK_RESOLVE,
        object_id=results[0].batch_id,
        summary=f"Labor reconciliation bulk resolve: {results[0].batch_id}",
        metadata={
            "batch_id": results[0].batch_id,
            "result_count": len(results),
            "resolution": _resolve_reconciliation_action(action),
        },
        object_type="ElectronicCardImportBatch",
    )
    return len(results)


def match_reconciliation_worker(result, worker, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(result, LaborReconciliationResult):
        result = (
            LaborReconciliationResult.objects.select_related("batch", "card_day__raw")
            .filter(id=result)
            .first()
        )
    if result is None:
        raise ValidationError("대사 결과를 찾을 수 없습니다.")
    _ensure_reconciliation_batch_editable(result.batch)
    if result.status != LaborReconciliationStatus.UNMATCHED or result.card_day is None:
        raise ValidationError("근로자 매칭은 미매칭 전자카드 행에서만 처리할 수 있습니다.")
    if not isinstance(worker, WorkerMaster):
        worker = WorkerMaster.objects.filter(id=worker).first()
    if worker is None:
        raise ValidationError("매칭할 근로자를 선택해 주세요.")
    if result.batch.reconciliation_results.exclude(id=result.id).filter(resolution__gt="").exists():
        raise ValidationError("이미 처리된 대사 결과가 있어 근로자 매칭 전에 안전하게 다시 정리할 수 없습니다.")

    raw = result.card_day.raw
    raw.matched_worker = worker
    raw.match_status = ElectronicCardMatchStatus.MATCHED
    raw.save(update_fields=["matched_worker", "match_status", "updated_at"])
    raw.days.update(worker=worker, match_status=ElectronicCardMatchStatus.MATCHED)
    reconcile_electronic_card_import_batch(result.batch, actor)
    _log_action_safe(
        actor=actor,
        action=LABOR_RECONCILIATION_WORKER_MATCH,
        object_id=result.id,
        summary=f"Labor reconciliation worker match: {result.id}",
        metadata={
            "result_id": result.id,
            "batch_id": result.batch_id,
            "worker_id": worker.id,
        },
        object_type="LaborReconciliationResult",
    )
    return result.batch


def _get_confirmed_row_payload(result, actor):
    resolution = result.resolution or ""
    if result.status == LaborReconciliationStatus.MATCH:
        source_basis = LaborConfirmedWorkSourceBasis.ERP
        final_work_unit = result.final_work_unit or result.erp_work_unit
        export_included = True
        export_value = (
            result.export_value if result.export_value is not None else final_work_unit
        )
    elif resolution == LaborReconciliationResolution.ERP:
        source_basis = LaborConfirmedWorkSourceBasis.ERP
        final_work_unit = (
            result.final_work_unit
            if result.final_work_unit is not None
            else result.erp_work_unit
        )
        export_included = result.export_included
        export_value = (
            result.export_value if result.export_value is not None else final_work_unit
        )
    elif resolution == LaborReconciliationResolution.CARD:
        source_basis = LaborConfirmedWorkSourceBasis.CARD
        final_work_unit = (
            result.final_work_unit
            if result.final_work_unit is not None
            else result.card_work_unit
        )
        export_included = result.export_included
        export_value = (
            result.export_value if result.export_value is not None else final_work_unit
        )
    elif resolution == LaborReconciliationResolution.MANUAL:
        source_basis = LaborConfirmedWorkSourceBasis.MANUAL
        final_work_unit = result.final_work_unit
        export_included = result.export_included
        export_value = (
            result.export_value if result.export_value is not None else final_work_unit
        )
    elif resolution == LaborReconciliationResolution.EXCLUDED:
        source_basis = LaborConfirmedWorkSourceBasis.EXCLUDED
        final_work_unit = Decimal("0")
        export_included = False
        export_value = Decimal("0")
    else:
        raise ValidationError("미해결 대사 결과가 남아 있어 확정할 수 없습니다.")

    if source_basis != LaborConfirmedWorkSourceBasis.EXCLUDED:
        if result.worker_id is None:
            raise ValidationError("근로자 미매칭 결과가 남아 있어 확정할 수 없습니다.")
        if final_work_unit is None:
            raise ValidationError("최종 공수가 확정되지 않은 결과가 있어 확정할 수 없습니다.")

    ledger = None
    if result.worker_id:
        ledger = (
            LaborWorkLedger.objects.filter(
                worker_id=result.worker_id,
                work_date=result.work_date,
                report_project=result.project,
            )
            .select_related("actual_project")
            .order_by("id")
            .first()
        )
    actual_project = (
        ledger.actual_project if ledger and ledger.actual_project_id else result.project
    )
    if result.card_day_id and result.card_day and result.card_day.card_project_id:
        card_project = result.card_day.card_project
    else:
        card_project = result.batch.project if result.batch.project_id else None

    return {
        "batch": result.batch,
        "year_month": result.year_month,
        "worker_id": result.worker_id,
        "actual_project": actual_project,
        "report_project": result.project,
        "card_project": card_project,
        "work_date": result.work_date,
        "final_work_unit": Decimal(str(final_work_unit or 0)),
        "source_basis": source_basis,
        "reconciliation_result": result,
        "export_included": export_included,
        "export_value": Decimal(str(export_value or 0)),
        "export_note": result.export_note or "",
        "confirmed_by": actor,
    }


def generate_labor_confirmed_work_days(batch, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(batch, ElectronicCardImportBatch):
        batch = (
            ElectronicCardImportBatch.objects.select_related("project")
            .filter(id=batch)
            .first()
        )
    if batch is None:
        raise ValidationError("전자카드 업로드 배치를 찾을 수 없습니다.")
    _ensure_reconciliation_batch_editable(batch)
    if batch.confirmed_work_days.exists():
        raise ValidationError("이미 확정 근로내역이 생성되었습니다.")

    results = list(
        batch.reconciliation_results.select_related(
            "worker",
            "project",
            "card_day__card_project",
            "batch",
        ).order_by("work_date", "id")
    )
    if not results:
        raise ValidationError("대사 결과가 없어 확정할 수 없습니다.")

    unresolved_exists = any(
        item.status
        in {
            LaborReconciliationStatus.ERP_ONLY,
            LaborReconciliationStatus.CARD_ONLY,
            LaborReconciliationStatus.DIFF,
            LaborReconciliationStatus.UNMATCHED,
        }
        and not item.resolution
        and item.final_work_unit is None
        for item in results
    )
    if unresolved_exists:
        raise ValidationError("미해결 대사 결과가 남아 있어 확정할 수 없습니다.")

    _assert_month_open(
        batch.project,
        batch.year_month,
        message_prefix="마감된 월은 확정 근로내역을 생성할 수 없습니다.",
    )
    confirmed_at = timezone.now()
    confirmed_rows = []
    for result in results:
        _assert_month_open(
            result.project,
            result.work_date,
            message_prefix="마감된 월은 확정 근로내역을 생성할 수 없습니다.",
        )
        payload = _get_confirmed_row_payload(result, actor)
        payload["confirmed_at"] = confirmed_at
        confirmed_rows.append(LaborConfirmedWorkDay(**payload))

    LaborConfirmedWorkDay.objects.bulk_create(confirmed_rows)
    included_count = sum(1 for row in confirmed_rows if row.export_included)
    excluded_count = len(confirmed_rows) - included_count
    total_final_work_unit = sum(
        (row.final_work_unit or Decimal("0")) for row in confirmed_rows
    )
    batch.header_check_summary = {
        **(batch.header_check_summary or {}),
        "confirmed_work_day_count": len(confirmed_rows),
        "confirmed_included_count": included_count,
        "confirmed_excluded_count": excluded_count,
        "confirmed_total_final_work_unit": str(total_final_work_unit),
        "confirmed_generated_at": confirmed_at.isoformat(),
    }
    batch.save(update_fields=["header_check_summary", "updated_at"])
    _log_action_safe(
        actor=actor,
        action=LABOR_CONFIRMED_WORKDAY_GENERATE,
        object_id=batch.id,
        summary=f"LaborConfirmedWorkDay generate: {batch.id}",
        metadata={
            "batch_id": batch.id,
            "year_month": batch.year_month.isoformat(),
            "project_id": batch.project_id,
            "created_count": len(confirmed_rows),
            "included_count": included_count,
            "excluded_count": excluded_count,
            "total_final_work_unit": str(total_final_work_unit),
            "confirmed_by_id": actor.id if actor else None,
        },
        object_type="ElectronicCardImportBatch",
    )
    return {
        "created_count": len(confirmed_rows),
        "included_count": included_count,
        "excluded_count": excluded_count,
        "total_final_work_unit": total_final_work_unit,
    }


def confirm_electronic_card_reconciliation_batch(batch, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(batch, ElectronicCardImportBatch):
        batch = (
            ElectronicCardImportBatch.objects.select_related("project")
            .filter(id=batch)
            .first()
        )
    if batch is None:
        raise ValidationError("전자카드 업로드 배치를 찾을 수 없습니다.")
    _ensure_reconciliation_batch_editable(batch)

    with transaction.atomic():
        generation = generate_labor_confirmed_work_days(batch, actor)
        batch.status = ElectronicCardImportBatchStatus.CONFIRMED
        batch.confirmed_by = actor
        batch.confirmed_at = timezone.now()
        batch.save(update_fields=["status", "confirmed_by", "confirmed_at", "updated_at"])

    _log_action_safe(
        actor=actor,
        action=LABOR_ECARD_RECONCILIATION_CONFIRM,
        object_id=batch.id,
        summary=f"ElectronicCardImportBatch reconciliation confirm: {batch.id}",
        metadata={
            "batch_id": batch.id,
            "year_month": batch.year_month.isoformat(),
            "project_id": batch.project_id,
            "result_count": batch.reconciliation_results.count(),
            "created_count": generation["created_count"],
        },
        object_type="ElectronicCardImportBatch",
    )
    return batch


def _format_export_cell_value(value):
    amount = Decimal(str(value or 0)).quantize(Decimal("0.01"))
    if amount <= 0:
        return None
    if amount == amount.quantize(Decimal("1")):
        return int(amount)
    return float(amount)


def _normalize_excel_compare_value(value):
    if value in (None, ""):
        return ""
    if isinstance(value, Decimal):
        return str(value.quantize(Decimal("0.01")).normalize())
    if isinstance(value, (int, float)):
        return str(Decimal(str(value)).quantize(Decimal("0.01")).normalize())
    return str(value).strip()


def _set_excel_cell_value(cell, new_value, *, force_write=False):
    current = _normalize_excel_compare_value(cell.value)
    incoming = _normalize_excel_compare_value(new_value)
    if current == incoming and not force_write:
        return False
    previous = cell.value
    cell.value = new_value
    return previous != new_value


def _build_export_note_summary(confirmed_rows):
    notes = []
    for row in confirmed_rows:
        note = str(row.export_note or "").strip()
        if note and note not in notes:
            notes.append(note)
    if not confirmed_rows:
        return "", ""
    excluded_rows = [row for row in confirmed_rows if not row.export_included]
    exclusion_reason = ""
    remark_parts = []
    if excluded_rows and len(excluded_rows) == len(confirmed_rows):
        exclusion_reason = notes[0] if notes else "신고 제외"
    elif excluded_rows:
        remark_parts.append("일부 일자 신고 제외")
    if notes:
        remark_parts.append("; ".join(notes[:3]))
    return exclusion_reason, " / ".join(remark_parts)


def _build_cwma_export_filename(batch):
    project_part = (batch.project.code or "").strip() or str(batch.project_id)
    return f"cwma_card_reupload_{batch.year_month.strftime('%Y%m')}_{project_part}_{batch.id}.xlsx"


def generate_cwma_card_reupload_excel(source_batch, actor, note=""):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(source_batch, ElectronicCardImportBatch):
        source_batch = (
            ElectronicCardImportBatch.objects.select_related("project")
            .filter(id=source_batch)
            .first()
        )
    if source_batch is None:
        raise ValidationError("전자카드 업로드 배치를 찾을 수 없습니다.")
    if source_batch.status != ElectronicCardImportBatchStatus.CONFIRMED:
        raise ValidationError("확정된 배치만 재업로드 엑셀을 생성할 수 있습니다.")
    if not source_batch.source_file:
        raise ValidationError("원본 전자카드 파일이 없어 재업로드 엑셀을 생성할 수 없습니다.")

    confirmed_rows = list(
        source_batch.confirmed_work_days.select_related(
            "worker",
            "reconciliation_result__card_day__raw",
        ).order_by("work_date", "id")
    )
    if not confirmed_rows:
        raise ValidationError("확정 근로내역이 없어 재업로드용 엑셀을 생성할 수 없습니다.")

    source_batch.source_file.open("rb")
    try:
        workbook = load_workbook(source_batch.source_file, read_only=False, data_only=False)
    except Exception as exc:
        logger.exception("Failed to open source workbook for CWMA reupload export.")
        raise ValidationError("원본 전자카드 파일을 다시 열지 못했습니다.") from exc

    sheet = _first_visible_sheet(workbook)
    summary = source_batch.header_check_summary or {}
    header_row_index = int(summary.get("header_row_index") or 0)
    if header_row_index <= 0 or header_row_index > sheet.max_row:
        best_score = -1
        for row_index in range(1, min(sheet.max_row, 30) + 1):
            row_values = [sheet.cell(row=row_index, column=col).value for col in range(1, sheet.max_column + 1)]
            field_map, day_columns = _resolve_header_map(row_values)
            score = len(field_map) + len(day_columns)
            if score > best_score:
                best_score = score
                header_row_index = row_index
    if header_row_index <= 0:
        raise ValidationError("원본 전자카드 파일의 헤더 행을 찾지 못했습니다.")

    header_values = [sheet.cell(row=header_row_index, column=col).value for col in range(1, sheet.max_column + 1)]
    field_map, day_columns = _resolve_header_map(header_values)

    raw_rows = list(
        source_batch.raw_rows.select_related("matched_worker").order_by("row_no", "id")
    )
    confirmed_by_raw_id = {}
    for row in confirmed_rows:
        raw_id = None
        if row.reconciliation_result_id and row.reconciliation_result and row.reconciliation_result.card_day_id:
            raw_id = row.reconciliation_result.card_day.raw_id
        if raw_id:
            confirmed_by_raw_id.setdefault(raw_id, []).append(row)

    included_count = sum(1 for row in confirmed_rows if row.export_included)
    excluded_count = len(confirmed_rows) - included_count
    total_export_work_unit = sum(
        (
            row.export_value
            if row.export_included and row.export_value is not None
            else row.final_work_unit
            if row.export_included
            else Decimal("0")
        )
        for row in confirmed_rows
    )

    changed_count = 0
    for raw in raw_rows:
        if raw.row_no <= 0 or raw.row_no > sheet.max_row:
            continue
        row_confirmed = confirmed_by_raw_id.get(raw.id, [])
        if not row_confirmed:
            continue
        by_day = {item.work_date.day: item for item in row_confirmed}
        last_day = monthrange(raw.work_month.year, raw.work_month.month)[1]
        for day in range(1, 32):
            col_index = day_columns.get(day)
            if col_index is None:
                continue
            cell = sheet.cell(row=raw.row_no, column=col_index + 1)
            confirmed = by_day.get(day)
            if day > last_day:
                continue
            if confirmed:
                export_value = (
                    confirmed.export_value
                    if confirmed.export_value is not None
                    else confirmed.final_work_unit
                )
                new_value = (
                    _format_export_cell_value(export_value)
                    if confirmed.export_included
                    else None
                )
            elif raw.matched_worker_id:
                new_value = None
            else:
                new_value = cell.value
            if _set_excel_cell_value(cell, new_value, force_write=confirmed is not None):
                changed_count += 1

        included_work_days = [
            item for item in row_confirmed if item.export_included and Decimal(str(item.export_value if item.export_value is not None else item.final_work_unit or 0)) > 0
        ]
        day_count_value = len(included_work_days)
        exclusion_reason, note_summary = _build_export_note_summary(row_confirmed)
        if "reported_days" in field_map:
            if _set_excel_cell_value(
                sheet.cell(row=raw.row_no, column=field_map["reported_days"] + 1),
                day_count_value if day_count_value > 0 else None,
            ):
                changed_count += 1
        if "confirmed_days" in field_map:
            if _set_excel_cell_value(
                sheet.cell(row=raw.row_no, column=field_map["confirmed_days"] + 1),
                day_count_value if day_count_value > 0 else None,
            ):
                changed_count += 1
        if "exclusion_reason" in field_map:
            if _set_excel_cell_value(
                sheet.cell(row=raw.row_no, column=field_map["exclusion_reason"] + 1),
                exclusion_reason or None,
            ):
                changed_count += 1
        if "note" in field_map:
            if _set_excel_cell_value(
                sheet.cell(row=raw.row_no, column=field_map["note"] + 1),
                note_summary or None,
            ):
                changed_count += 1

    generated_filename = _build_cwma_export_filename(source_batch)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    with transaction.atomic():
        export_batch = LaborExcelExportBatch.objects.create(
            export_type=LaborExcelExportType.CWMA_CARD_REUPLOAD,
            year_month=source_batch.year_month,
            project=source_batch.project,
            source_batch=source_batch,
            generated_file=ContentFile(output.getvalue(), name=generated_filename),
            original_filename=source_batch.original_filename,
            generated_filename=generated_filename,
            generated_by=actor,
            status=LaborExcelExportStatus.GENERATED,
            changed_count=changed_count,
            included_count=included_count,
            excluded_count=excluded_count,
            total_export_work_unit=total_export_work_unit,
            note=str(note or "").strip(),
        )

    _log_action_safe(
        actor=actor,
        action=LABOR_EXCEL_EXPORT_GENERATE,
        object_id=export_batch.id,
        summary=f"LaborExcelExportBatch generate: {export_batch.id}",
        metadata={
            "export_type": export_batch.export_type,
            "export_id": export_batch.id,
            "source_batch_id": source_batch.id,
            "year_month": source_batch.year_month.isoformat(),
            "project_id": source_batch.project_id,
            "generated_filename": export_batch.generated_filename,
            "changed_count": changed_count,
            "included_count": included_count,
            "excluded_count": excluded_count,
            "total_export_work_unit": str(total_export_work_unit),
        },
        object_type="LaborExcelExportBatch",
    )
    return export_batch


def register_labor_excel_export_download(export_batch, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if not isinstance(export_batch, LaborExcelExportBatch):
        export_batch = LaborExcelExportBatch.objects.select_related("source_batch", "project").filter(id=export_batch).first()
    if export_batch is None:
        raise ValidationError("엑셀 내보내기 배치를 찾을 수 없습니다.")
    if not export_batch.generated_file:
        raise ValidationError("생성된 엑셀 파일이 없습니다.")
    download_time = timezone.now()
    export_batch.status = LaborExcelExportStatus.DOWNLOADED
    export_batch.downloaded_by = actor
    export_batch.downloaded_at = download_time
    export_batch.save(update_fields=["status", "downloaded_by", "downloaded_at", "updated_at"])
    _log_action_safe(
        actor=actor,
        action=LABOR_EXCEL_EXPORT_DOWNLOAD,
        object_id=export_batch.id,
        summary=f"LaborExcelExportBatch download: {export_batch.id}",
        metadata={
            "export_type": export_batch.export_type,
            "export_id": export_batch.id,
            "source_batch_id": export_batch.source_batch_id,
            "downloaded_by_id": actor.id if actor else None,
            "downloaded_at": download_time.isoformat(),
        },
        object_type="LaborExcelExportBatch",
    )
    return export_batch


def create_electronic_card_import_batch(data, file, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    project = _resolve_project(data.get("project"))
    if project is None:
        raise ValidationError({"project": "현장을 선택해 주세요."})

    year_month = _parse_year_month_value(data.get("year_month"))
    if year_month is None:
        raise ValidationError({"year_month": "기준월 형식은 YYYY-MM 이어야 합니다."})

    original_filename = str(getattr(file, "name", "") or "").strip()
    if not original_filename.lower().endswith(".xlsx"):
        raise ValidationError({"source_file": "전자카드 Excel은 .xlsx 파일만 업로드할 수 있습니다."})

    try:
        header_summary = _build_e_card_header_summary(file, selected_month=year_month)
    except ValidationError:
        raise
    except Exception as exc:
        logger.exception("Failed to read electronic card workbook header.")
        raise ValidationError("전자카드 파일 헤더를 확인하지 못했습니다. 파일 형식을 확인해 주세요.") from exc

    file.seek(0)
    with transaction.atomic():
        batch = ElectronicCardImportBatch.objects.create(
            year_month=year_month,
            project=project,
            cwma_project_name=header_summary.get("cwma_project_name", ""),
            deduction_join_no=header_summary.get("deduction_join_no", ""),
            company_name=header_summary.get("company_name", ""),
            source_file=file,
            original_filename=original_filename,
            status=ElectronicCardImportBatchStatus.PARSE_READY,
            uploaded_by=actor,
            header_check_summary=header_summary,
        )
        _log_action_safe(
            actor=actor,
            action=ELECTRONIC_CARD_IMPORT_UPLOAD,
            object_id=batch.id,
            summary=f"ElectronicCardImportBatch upload: {batch.id}",
            metadata={
                "batch_id": batch.id,
                "year_month": batch.year_month.isoformat(),
                "project_id": batch.project_id,
                "original_filename": batch.original_filename,
                "status": batch.status,
                "header_check_summary": {
                    "sheet_name": header_summary.get("sheet_name", ""),
                    "header_row_index": header_summary.get("header_row_index", 0),
                    "detected_headers": sorted(header_summary.get("detected_headers", {}).values()),
                    "missing_optional_headers": header_summary.get("missing_optional_headers", []),
                    "day_column_count": header_summary.get("day_column_count", 0),
                    "warnings": header_summary.get("warnings", []),
                    "workbook_months": header_summary.get("workbook_months", []),
                },
            },
            object_type="ElectronicCardImportBatch",
        )
    return batch


def _assert_labor_work_ledger_open(*, actual_project, report_project, work_date, message_prefix: str):
    _assert_month_open(actual_project, work_date, message_prefix=message_prefix)
    if report_project and report_project != actual_project:
        _assert_month_open(report_project, work_date, message_prefix=message_prefix)


def _prepare_labor_work_ledger_data(data: dict) -> dict:
    worker = _resolve_worker(data.get("worker"))
    actual_project = _resolve_project(data.get("actual_project"))
    report_project = _resolve_project(data.get("report_project")) or actual_project
    labor_role = data.get("labor_role")
    if labor_role and not isinstance(labor_role, LaborRole):
        labor_role = LaborRole.objects.filter(id=labor_role).first()
    timesheet_ref = _resolve_timesheet(data.get("timesheet_ref"))
    cost_ref = _resolve_cost_actual(data.get("cost_ref"))
    work_date = data.get("work_date")

    if timesheet_ref is not None:
        if work_date is None:
            work_date = timesheet_ref.work_date
        if actual_project is None:
            actual_project = timesheet_ref.project
        if report_project is None:
            report_project = actual_project
        if labor_role is None:
            line = timesheet_ref.lines.select_related("labor_role").first()
            if line is not None:
                labor_role = line.labor_role

    if cost_ref is not None:
        if work_date is None:
            work_date = cost_ref.report_date
        if actual_project is None:
            actual_project = cost_ref.project
        if report_project is None:
            report_project = actual_project

    if worker is None:
        raise ValidationError({"worker": "근로자를 선택해 주세요."})
    if work_date is None:
        raise ValidationError({"work_date": "근무일을 입력해 주세요."})
    if actual_project is None:
        raise ValidationError({"actual_project": "실제 근무 프로젝트를 선택해 주세요."})
    if report_project is None:
        report_project = actual_project
    if labor_role is None:
        raise ValidationError({"labor_role": "노무 역할을 선택해 주세요."})

    source = str(data.get("source") or "").strip() or LaborWorkLedgerSource.MANUAL
    if source == LaborWorkLedgerSource.MANUAL and timesheet_ref is not None:
        source = LaborWorkLedgerSource.TIMESHEET
    elif source == LaborWorkLedgerSource.MANUAL and cost_ref is not None:
        source = LaborWorkLedgerSource.COST

    status = str(data.get("status") or "").strip() or LaborWorkLedgerStatus.DRAFT
    work_unit = _coerce_decimal(data.get("work_unit"), field_name="work_unit")
    work_hours = _coerce_decimal(data.get("work_hours"), field_name="work_hours")
    unit_wage = _coerce_int(data.get("unit_wage"), field_name="unit_wage")
    income_tax = _coerce_int(data.get("income_tax"), field_name="income_tax")
    local_tax = _coerce_int(data.get("local_tax"), field_name="local_tax")
    employment_insurance = _coerce_int(
        data.get("employment_insurance"), field_name="employment_insurance"
    )
    pension = _coerce_int(data.get("pension"), field_name="pension")
    health_insurance = _coerce_int(
        data.get("health_insurance"), field_name="health_insurance"
    )

    return {
        "worker": worker,
        "work_date": work_date,
        "actual_project": actual_project,
        "report_project": report_project,
        "labor_role": labor_role,
        "work_unit": work_unit,
        "work_hours": work_hours,
        "unit_wage": unit_wage,
        "income_tax": income_tax,
        "local_tax": local_tax,
        "employment_insurance": employment_insurance,
        "pension": pension,
        "health_insurance": health_insurance,
        "detail_work_type": str(data.get("detail_work_type") or "").strip(),
        "source": source,
        "status": status,
        "timesheet_ref": timesheet_ref,
        "cost_ref": cost_ref,
    }


def create_labor_work_ledger(data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    payload = _prepare_labor_work_ledger_data(data)
    _assert_labor_work_ledger_open(
        actual_project=payload["actual_project"],
        report_project=payload["report_project"],
        work_date=payload["work_date"],
        message_prefix="노무 작업 원장 등록은 불가합니다.",
    )
    ledger = LaborWorkLedger.objects.create(**payload)
    try:
        log_action(
            actor=actor,
            action="LABOR_WORK_LEDGER_CREATE",
            object_type="LaborWorkLedger",
            object_id=ledger.id,
            after=_labor_work_ledger_snapshot(ledger),
            meta={"ledger_id": ledger.id, "worker_id": ledger.worker_id},
        )
    except Exception:
        logger.warning("AuditLog failed for LABOR_WORK_LEDGER_CREATE.", exc_info=True)
    return ledger


def update_labor_work_ledger(ledger: LaborWorkLedger, data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    before = _labor_work_ledger_snapshot(ledger)
    _assert_labor_work_ledger_open(
        actual_project=ledger.actual_project,
        report_project=ledger.report_project,
        work_date=ledger.work_date,
        message_prefix="노무 작업 원장 수정은 불가합니다.",
    )
    payload = _prepare_labor_work_ledger_data(data)
    _assert_labor_work_ledger_open(
        actual_project=payload["actual_project"],
        report_project=payload["report_project"],
        work_date=payload["work_date"],
        message_prefix="노무 작업 원장 수정은 불가합니다.",
    )
    for field_name, value in payload.items():
        setattr(ledger, field_name, value)
    ledger.save()
    after = _labor_work_ledger_snapshot(ledger)
    try:
        log_action(
            actor=actor,
            action="LABOR_WORK_LEDGER_UPDATE",
            object_type="LaborWorkLedger",
            object_id=ledger.id,
            before=before,
            after=after,
            meta={"ledger_id": ledger.id, "worker_id": ledger.worker_id},
        )
    except Exception:
        logger.warning("AuditLog failed for LABOR_WORK_LEDGER_UPDATE.", exc_info=True)
    return ledger


def update_labor_reporting_project(ledger_ids, report_project, reason, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if not ledger_ids:
        raise ValidationError("변경할 원장을 선택해 주세요.")
    new_report_project = _resolve_project(report_project)
    if new_report_project is None:
        raise ValidationError({"report_project": "신고용 현장을 선택해 주세요."})

    reason = str(reason or "").strip()
    ledgers = list(
        LaborWorkLedger.objects.select_related(
            "worker",
            "actual_project",
            "report_project",
            "labor_role",
        ).filter(id__in=ledger_ids)
    )
    if len(ledgers) != len(set(int(ledger_id) for ledger_id in ledger_ids)):
        raise ValidationError("선택한 원장 중 일부를 찾을 수 없습니다.")

    updated = []
    with transaction.atomic():
        for ledger in ledgers:
            if ledger.status == LaborWorkLedgerStatus.CONFIRMED:
                raise ValidationError("확정된 원장은 신고용 현장을 수정할 수 없습니다.")
            if new_report_project != ledger.actual_project and not reason:
                raise ValidationError("실제 현장과 다른 신고용 현장으로 변경할 때는 사유가 필요합니다.")
            try:
                _assert_labor_work_ledger_open(
                    actual_project=ledger.actual_project,
                    report_project=ledger.report_project,
                    work_date=ledger.work_date,
                    message_prefix="마감된 월은 수정할 수 없습니다. 정정으로 처리하세요.",
                )
                _assert_labor_work_ledger_open(
                    actual_project=ledger.actual_project,
                    report_project=new_report_project,
                    work_date=ledger.work_date,
                    message_prefix="마감된 월은 수정할 수 없습니다. 정정으로 처리하세요.",
                )
            except PermissionDenied as exc:
                raise PermissionDenied("마감된 월은 수정할 수 없습니다. 정정으로 처리하세요.") from exc

            old_report_project_id = ledger.report_project_id
            if old_report_project_id == new_report_project.id:
                continue
            before = _labor_work_ledger_snapshot(ledger)
            ledger.report_project = new_report_project
            ledger.save()
            after = _labor_work_ledger_snapshot(ledger)
            try:
                log_action(
                    actor=actor,
                    action="LABOR_REPORTING_PROJECT_UPDATE",
                    object_type="LaborWorkLedger",
                    object_id=ledger.id,
                    before=before,
                    after=after,
                    meta={
                        "ledger_id": ledger.id,
                        "worker_id": ledger.worker_id,
                        "work_date": str(ledger.work_date),
                        "actual_project_id": ledger.actual_project_id,
                        "old_report_project_id": old_report_project_id,
                        "new_report_project_id": new_report_project.id,
                        "reason": reason,
                    },
                )
            except Exception:
                logger.warning(
                    "AuditLog failed for LABOR_REPORTING_PROJECT_UPDATE.",
                    exc_info=True,
                )
            updated.append(ledger)
    return updated


def generate_labor_monthly_payroll(*, year: int, month: int, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    period = date_type(int(year), int(month), 1)
    _assert_month_open(None, period, message_prefix="월 급여 집계 생성은 불가합니다.")

    source_qs = LaborWorkLedger.objects.filter(
        work_month=period,
        status=LaborWorkLedgerStatus.CONFIRMED,
    )
    project_ids = set(
        source_qs.values_list("actual_project_id", flat=True)
    ) | set(source_qs.values_list("report_project_id", flat=True))
    project_ids.discard(None)
    if project_ids:
        for project in Project.objects.filter(id__in=project_ids):
            _assert_month_open(
                project,
                period,
                message_prefix="월 급여 집계 생성은 불가합니다.",
            )

    aggregates = list(
        source_qs.values("worker_id", "actual_project_id", "report_project_id").annotate(
            total_work_unit=Sum("work_unit"),
            gross_wage=Sum("gross_wage"),
            income_tax=Sum("income_tax"),
            local_tax=Sum("local_tax"),
            employment_insurance=Sum("employment_insurance"),
            pension=Sum("pension"),
            health_insurance=Sum("health_insurance"),
            net_pay=Sum("net_pay"),
        )
    )
    worker_ids = [row["worker_id"] for row in aggregates]
    workers = {
        worker.id: worker
        for worker in WorkerMaster.objects.filter(id__in=worker_ids).order_by("id")
    }
    existing = {
        (
            payroll.worker_id,
            payroll.project_id,
            payroll.report_project_id,
        ): payroll
        for payroll in LaborMonthlyPayroll.objects.filter(year_month=period)
    }
    created_count = 0
    updated_count = 0
    with transaction.atomic():
        seen_keys = set()
        for row in aggregates:
            key = (
                row["worker_id"],
                row["actual_project_id"],
                row["report_project_id"],
            )
            seen_keys.add(key)
            payroll = existing.get(key)
            worker = workers.get(row["worker_id"])
            if payroll is None:
                payroll = LaborMonthlyPayroll(
                    year_month=period,
                    worker_id=row["worker_id"],
                    project_id=row["actual_project_id"],
                    report_project_id=row["report_project_id"],
                    payment_status=LaborMonthlyPayrollPaymentStatus.PENDING,
                )
                created_count += 1
            else:
                updated_count += 1
            insurance_deductions = sum(
                int(row.get(field_name) or 0)
                for field_name in ("employment_insurance", "pension", "health_insurance")
            )
            payroll.total_work_unit = row["total_work_unit"] or Decimal("0")
            payroll.gross_wage = int(row["gross_wage"] or 0)
            payroll.income_tax = int(row["income_tax"] or 0)
            payroll.local_tax = int(row["local_tax"] or 0)
            payroll.insurance_deductions = insurance_deductions
            payroll.net_pay = int(row["net_pay"] or 0)
            payroll.bank_name = worker.bank_name if worker else ""
            payroll.account_number_masked = (
                worker.account_number_masked if worker else ""
            )
            payroll.save()
        stale_qs = LaborMonthlyPayroll.objects.filter(year_month=period)
        if seen_keys:
            keep_q = Q()
            for worker_id, project_id, report_project_id in seen_keys:
                keep_q |= Q(
                    worker_id=worker_id,
                    project_id=project_id,
                    report_project_id=report_project_id,
                )
            stale_qs = stale_qs.exclude(keep_q)
        deleted_count = stale_qs.count()
        if deleted_count:
            stale_qs.delete()
    _log_action_safe(
        actor=actor,
        action="LABOR_MONTHLY_PAYROLL_GENERATE",
        object_id=0,
        summary=f"LaborMonthlyPayroll generate: {period:%Y-%m}",
        metadata={
            "year": int(year),
            "month": int(month),
            "created_count": created_count,
            "updated_count": updated_count,
            "deleted_count": deleted_count,
            "row_count": len(aggregates),
        },
        object_type="LaborMonthlyPayroll",
    )
    return {
        "year_month": period,
        "created_count": created_count,
        "updated_count": updated_count,
        "deleted_count": deleted_count,
        "row_count": len(aggregates),
    }


def create_rate(data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    role = data.get("labor_role")
    if not isinstance(role, LaborRole):
        role = LaborRole.objects.filter(id=role).first()
    if not role:
        raise ValidationError({"labor_role": "labor_role is required."})
    rate_type = str(data.get("rate_type") or LaborRateType.DAY)
    scope_type = str(data.get("scope_type") or LaborRateScope.GLOBAL)
    project = _resolve_project(data.get("project"))
    unit_rate = data.get("unit_rate")
    effective_from = data.get("effective_from")
    effective_to = data.get("effective_to")
    if unit_rate in (None, ""):
        raise ValidationError({"unit_rate": "unit_rate is required."})
    unit_rate = int(unit_rate)
    if unit_rate <= 0:
        raise ValidationError({"unit_rate": "unit_rate must be > 0."})
    if not effective_from:
        raise ValidationError({"effective_from": "effective_from is required."})
    if scope_type == LaborRateScope.PROJECT and not project:
        raise ValidationError({"project": "project is required for PROJECT rates."})
    if scope_type == LaborRateScope.GLOBAL:
        project = None
    if effective_to and effective_to < effective_from:
        raise ValidationError({"effective_to": "effective_to must be >= effective_from."})
    try:
        _validate_overlap(
            role=role,
            rate_type=rate_type,
            scope_type=scope_type,
            project=project,
            start=effective_from,
            end=effective_to,
        )
    except ValidationError as exc:
        _log_action_safe(
            actor=actor,
            action="LABOR_RATE_OVERLAP_BLOCKED",
            object_id=role.id,
            summary=f"LaborRate overlap blocked: {role.code}",
            metadata={
                "role_id": role.id,
                "role_code": role.code,
                "project_id": project.id if project else None,
                "scope_type": scope_type,
                "rate_type": rate_type,
                "effective_from": str(effective_from),
                "effective_to": str(effective_to) if effective_to else None,
            },
            object_type="LaborRateTable",
        )
        raise
    with transaction.atomic():
        rate = LaborRateTable.objects.create(
            labor_role=role,
            rate_type=rate_type,
            unit_rate=unit_rate,
            currency=str(data.get("currency") or "KRW"),
            effective_from=effective_from,
            effective_to=effective_to,
            scope_type=scope_type,
            project=project,
            is_active=bool(data.get("is_active", True)),
            note=str(data.get("note") or "").strip(),
            created_by=actor,
        )
    _log_action_safe(
        actor=actor,
        action="LABOR_RATE_CREATE",
        object_id=rate.id,
        summary=f"LaborRate create: {role.code}",
        metadata={
            "role_id": role.id,
            "role_code": role.code,
            "project_id": project.id if project else None,
            "scope_type": scope_type,
            "rate_type": rate_type,
            "unit_rate": unit_rate,
            "effective_from": str(effective_from),
            "effective_to": str(effective_to) if effective_to else None,
        },
        object_type="LaborRateTable",
    )
    return rate


def update_rate(rate: LaborRateTable, data, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    role = rate.labor_role
    rate_type = str(data.get("rate_type") or rate.rate_type)
    scope_type = str(data.get("scope_type") or rate.scope_type)
    project = _resolve_project(data.get("project")) if "project" in data else rate.project
    unit_rate = data.get("unit_rate", rate.unit_rate)
    effective_from = data.get("effective_from", rate.effective_from)
    effective_to = data.get("effective_to", rate.effective_to)
    if unit_rate in (None, ""):
        raise ValidationError({"unit_rate": "unit_rate is required."})
    unit_rate = int(unit_rate)
    if unit_rate <= 0:
        raise ValidationError({"unit_rate": "unit_rate must be > 0."})
    if scope_type == LaborRateScope.PROJECT and not project:
        raise ValidationError({"project": "project is required for PROJECT rates."})
    if scope_type == LaborRateScope.GLOBAL:
        project = None
    if effective_to and effective_to < effective_from:
        raise ValidationError({"effective_to": "effective_to must be >= effective_from."})
    try:
        _validate_overlap(
            role=role,
            rate_type=rate_type,
            scope_type=scope_type,
            project=project,
            start=effective_from,
            end=effective_to,
            exclude_id=rate.id,
        )
    except ValidationError as exc:
        _log_action_safe(
            actor=actor,
            action="LABOR_RATE_OVERLAP_BLOCKED",
            object_id=rate.id,
            summary=f"LaborRate overlap blocked: {role.code}",
            metadata={
                "role_id": role.id,
                "role_code": role.code,
                "project_id": project.id if project else None,
                "scope_type": scope_type,
                "rate_type": rate_type,
                "effective_from": str(effective_from),
                "effective_to": str(effective_to) if effective_to else None,
            },
            object_type="LaborRateTable",
        )
        raise
    before_active = rate.is_active
    rate.rate_type = rate_type
    rate.scope_type = scope_type
    rate.project = project
    rate.unit_rate = unit_rate
    rate.effective_from = effective_from
    rate.effective_to = effective_to
    rate.currency = str(data.get("currency") or rate.currency or "KRW")
    if "note" in data:
        rate.note = str(data.get("note") or "").strip()
    if "is_active" in data:
        rate.is_active = bool(data.get("is_active"))
    rate.save()
    action = "LABOR_RATE_UPDATE"
    if before_active and not rate.is_active:
        action = "LABOR_RATE_DEACTIVATE"
    _log_action_safe(
        actor=actor,
        action=action,
        object_id=rate.id,
        summary=f"LaborRate update: {role.code}",
        metadata={
            "role_id": role.id,
            "role_code": role.code,
            "project_id": project.id if project else None,
            "scope_type": scope_type,
            "rate_type": rate_type,
            "unit_rate": unit_rate,
            "effective_from": str(effective_from),
            "effective_to": str(effective_to) if effective_to else None,
        },
        object_type="LaborRateTable",
    )
    return rate


def get_applicable_rate(project, labor_role, date, *, rate_type=LaborRateType.DAY):
    if not labor_role or not date:
        return None
    if not isinstance(labor_role, LaborRole):
        labor_role = LaborRole.objects.filter(id=labor_role).first()
    if not labor_role:
        return None
    if project and not isinstance(project, Project):
        project = Project.objects.filter(id=project).first()
    target_date = date if isinstance(date, date_type) else None
    if target_date is None:
        return None
    project_rate = (
        LaborRateTable.objects.filter(
            labor_role=labor_role,
            rate_type=rate_type,
            scope_type=LaborRateScope.PROJECT,
            project=project,
            is_active=True,
            effective_from__lte=target_date,
        )
        .filter(Q(effective_to__isnull=True) | Q(effective_to__gte=target_date))
        .order_by("-effective_from")
        .first()
    )
    if project_rate:
        return project_rate
    global_rate = (
        LaborRateTable.objects.filter(
            labor_role=labor_role,
            rate_type=rate_type,
            scope_type=LaborRateScope.GLOBAL,
            project__isnull=True,
            is_active=True,
            effective_from__lte=target_date,
        )
        .filter(Q(effective_to__isnull=True) | Q(effective_to__gte=target_date))
        .order_by("-effective_from")
        .first()
    )
    return global_rate


def _next_timesheet_no(work_date: date_type) -> str:
    key = f"TS-{work_date.strftime('%Y%m')}"
    with transaction.atomic():
        seq, _created = TimesheetNumberSequence.objects.select_for_update().get_or_create(
            key=key, defaults={"last_number": 0}
        )
        seq.last_number += 1
        seq.save(update_fields=["last_number"])
    return f"{key}-{seq.last_number:04d}"


def create_timesheet(*, project, work_date, actor, note="") -> Timesheet:
    require_role(actor, [Role.FIELD, Role.HQ, Role.CEO])
    require_project_access(actor, project.id)
    if not isinstance(work_date, date_type):
        raise ValidationError({"work_date": "\uc791\uc131\uc77c\uc744 \ud655\uc778\ud574 \uc8fc\uc138\uc694."})
    _assert_month_open(project, work_date, message_prefix="\ucd9c\uc5ed\ubd80 \uc785\ub825\uc740 \ubd88\uac00\ub2a5\ud569\ub2c8\ub2e4.")
    sheet_no = _next_timesheet_no(work_date)
    timesheet = Timesheet.objects.create(
        sheet_no=sheet_no,
        project=project,
        work_date=work_date,
        status=TimesheetStatus.DRAFT,
        note=note,
        created_by=actor,
    )
    _log_action_safe(
        actor=actor,
        action="TIMESHEET_CREATE",
        object_id=timesheet.id,
        summary=f"Timesheet create: {timesheet.sheet_no}",
        metadata={
            "timesheet_id": timesheet.id,
            "sheet_no": timesheet.sheet_no,
            "project_id": project.id,
            "work_date": str(work_date),
            "status": timesheet.status,
        },
        object_type="Timesheet",
    )
    return timesheet


def upsert_timesheet_lines(*, timesheet: Timesheet, lines_payload: list[dict], actor):
    require_role(actor, [Role.FIELD, Role.HQ, Role.CEO])
    _ensure_timesheet_editable(timesheet, actor=actor)
    _assert_month_open(timesheet.project, timesheet.work_date, message_prefix="\ucd9c\uc5ed\ubd80 \uc218\uc815\uc740 \ubd88\uac00\ub2a5\ud569\ub2c8\ub2e4.")
    cleaned_lines: list[dict] = []
    for raw in lines_payload:
        role_id = raw.get("labor_role_id") or raw.get("labor_role")
        if not role_id:
            continue
        labor_role = LaborRole.objects.filter(id=role_id, is_active=True).first()
        if not labor_role:
            raise ValidationError("\uc720\ud6a8\ud55c \uc9c1\uc885\uc744 \uc120\ud0dd\ud574 \uc8fc\uc138\uc694.")
        headcount = raw.get("headcount")
        if headcount in (None, ""):
            raise ValidationError("\uc778\uc6d0\uc744 \uc785\ub825\ud574 \uc8fc\uc138\uc694.")
        try:
            headcount_val = Decimal(str(headcount))
        except Exception:
            raise ValidationError("\uc778\uc6d0 \uac12\uc774 \uc62c\ubc14\ub974\uc9c0 \uc54a\uc2b5\ub2c8\ub2e4.")
        if headcount_val <= 0:
            raise ValidationError("\uc778\uc6d0\uc740 0\ubcf4\ub2e4 \ud070 \uac12\uc744 \uc785\ub825\ud574 \uc8fc\uc138\uc694.")
        rate_type = str(raw.get("rate_type") or LaborRateType.DAY)
        hours_val = None
        if rate_type == LaborRateType.HOUR:
            hours = raw.get("hours")
            if hours in (None, ""):
                raise ValidationError("\uc2dc\uac04\uc744 \uc785\ub825\ud574 \uc8fc\uc138\uc694.")
            try:
                hours_val = Decimal(str(hours))
            except Exception:
                raise ValidationError("\uc2dc\uac04 \uac12\uc774 \uc62c\ubc14\ub974\uc9c0 \uc54a\uc2b5\ub2c8\ub2e4.")
            if hours_val <= 0:
                raise ValidationError("\uc2dc\uac04\uc740 0\ubcf4\ub2e4 \ud070 \uac12\uc744 \uc785\ub825\ud574 \uc8fc\uc138\uc694.")
        cleaned_lines.append(
            {
                "labor_role": labor_role,
                "headcount": headcount_val,
                "hours": hours_val,
                "rate_type": rate_type,
                "memo": str(raw.get("memo") or "").strip(),
            }
        )
    if not cleaned_lines:
        raise ValidationError("\ucd9c\uc5ed \ub77c\uc778\uc740 \ucd5c\uc18c 1\uac1c \uc774\uc0c1 \ud544\uc694\ud569\ub2c8\ub2e4.")
    with transaction.atomic():
        TimesheetLine.objects.filter(timesheet=timesheet).delete()
        TimesheetLine.objects.bulk_create(
            [
                TimesheetLine(
                    timesheet=timesheet,
                    labor_role=line["labor_role"],
                    headcount=line["headcount"],
                    hours=line["hours"],
                    rate_type=line["rate_type"],
                    unit_rate=0,
                    amount=0,
                    memo=line["memo"],
                )
                for line in cleaned_lines
            ]
        )
        timesheet.save(update_fields=["updated_at"])
    _log_action_safe(
        actor=actor,
        action="TIMESHEET_UPDATE",
        object_id=timesheet.id,
        summary=f"Timesheet update: {timesheet.sheet_no}",
        metadata={
            "timesheet_id": timesheet.id,
            "sheet_no": timesheet.sheet_no,
            "project_id": timesheet.project_id,
            "work_date": str(timesheet.work_date),
            "line_count": len(cleaned_lines),
            "status": timesheet.status,
        },
        object_type="Timesheet",
    )
    return timesheet


def submit_timesheet(*, timesheet: Timesheet, actor):
    require_role(actor, [Role.FIELD, Role.HQ, Role.CEO])
    if timesheet.status not in (TimesheetStatus.DRAFT, TimesheetStatus.REJECTED):
        raise PermissionDenied("\uc81c\ucd9c \ud560 \uc218 \uc5c6\ub294 \uc0c1\ud0dc\uc785\ub2c8\ub2e4.")
    _assert_month_open(timesheet.project, timesheet.work_date, message_prefix="\ucd9c\uc5ed\ubd80 \uc81c\ucd9c\uc740 \ubd88\uac00\ub2a5\ud569\ub2c8\ub2e4.")
    lines = list(timesheet.lines.select_related("labor_role").all())
    if not lines:
        raise ValidationError("\ucd9c\uc5ed \ub77c\uc778\uc740 \ucd5c\uc18c 1\uac1c \uc774\uc0c1 \ud544\uc694\ud569\ub2c8\ub2e4.")
    total_amount = 0
    with transaction.atomic():
        for line in lines:
            rate = get_applicable_rate(
                timesheet.project,
                line.labor_role,
                timesheet.work_date,
                rate_type=line.rate_type,
            )
            if not rate:
                raise ValidationError("\uc9c1\uc885 \ub2e8\uac00\uac00 \ub4f1\ub85d\ub418\uc9c0 \uc54a\uc558\uc2b5\ub2c8\ub2e4.")
            unit_rate = int(rate.unit_rate)
            if line.rate_type == LaborRateType.HOUR:
                hours = line.hours or Decimal("0")
                amount = (line.headcount * hours * Decimal(unit_rate)).quantize(
                    Decimal("1"), rounding=ROUND_HALF_UP
                )
            else:
                amount = (line.headcount * Decimal(unit_rate)).quantize(
                    Decimal("1"), rounding=ROUND_HALF_UP
                )
            line.unit_rate = unit_rate
            line.amount = int(amount)
            line.save(update_fields=["unit_rate", "amount"])
            total_amount += int(amount)
        timesheet.status = TimesheetStatus.SUBMITTED
        timesheet.submitted_at = timezone.now()
        timesheet.save(update_fields=["status", "submitted_at", "updated_at"])
    _log_action_safe(
        actor=actor,
        action="TIMESHEET_SUBMIT",
        object_id=timesheet.id,
        summary=f"Timesheet submit: {timesheet.sheet_no}",
        metadata={
            "timesheet_id": timesheet.id,
            "sheet_no": timesheet.sheet_no,
            "project_id": timesheet.project_id,
            "work_date": str(timesheet.work_date),
            "line_count": len(lines),
            "total_amount": total_amount,
            "status": timesheet.status,
        },
        object_type="Timesheet",
    )
    return timesheet, total_amount


def approve_timesheet(*, timesheet: Timesheet, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    if timesheet.status != TimesheetStatus.SUBMITTED:
        raise ValidationError("\uc81c\ucd9c \ub300\uae30 \uc0c1\ud0dc\uc758 \ucd9c\uc5ed\ubd80\ub9cc \uc2b9\uc778\ud560 \uc218 \uc788\uc2b5\ub2c8\ub2e4.")
    _assert_month_open(timesheet.project, timesheet.work_date, message_prefix="\ucd9c\uc5ed\ubd80 \uc2b9\uc778\uc740 \ubd88\uac00\ub2a5\ud569\ub2c8\ub2e4.")
    timesheet.status = TimesheetStatus.APPROVED
    timesheet.approved_at = timezone.now()
    timesheet.approved_by = actor
    timesheet.save(update_fields=["status", "approved_at", "approved_by", "updated_at"])
    _log_action_safe(
        actor=actor,
        action="TIMESHEET_APPROVE",
        object_id=timesheet.id,
        summary=f"Timesheet approve: {timesheet.sheet_no}",
        metadata={
            "timesheet_id": timesheet.id,
            "sheet_no": timesheet.sheet_no,
            "project_id": timesheet.project_id,
            "work_date": str(timesheet.work_date),
            "status": timesheet.status,
        },
        object_type="Timesheet",
    )
    return timesheet


def reject_timesheet(*, timesheet: Timesheet, actor, reason: str):
    require_role(actor, [Role.HQ, Role.CEO])
    if timesheet.status != TimesheetStatus.SUBMITTED:
        raise ValidationError("\uc81c\ucd9c \ub300\uae30 \uc0c1\ud0dc\uc758 \ucd9c\uc5ed\ubd80\ub9cc \ubc18\ub824\ud560 \uc218 \uc788\uc2b5\ub2c8\ub2e4.")
    if not reason:
        raise ValidationError("\ubc18\ub824 \uc0ac\uc720\ub97c \uc785\ub825\ud574 \uc8fc\uc138\uc694.")
    _assert_month_open(timesheet.project, timesheet.work_date, message_prefix="\ucd9c\uc5ed\ubd80 \ubc18\ub824\ub294 \ubd88\uac00\ub2a5\ud569\ub2c8\ub2e4.")
    timesheet.status = TimesheetStatus.REJECTED
    timesheet.rejected_at = timezone.now()
    timesheet.rejected_by = actor
    timesheet.reject_reason = reason
    timesheet.save(
        update_fields=[
            "status",
            "rejected_at",
            "rejected_by",
            "reject_reason",
            "updated_at",
        ]
    )
    _log_action_safe(
        actor=actor,
        action="TIMESHEET_REJECT",
        object_id=timesheet.id,
        summary=f"Timesheet reject: {timesheet.sheet_no}",
        metadata={
            "timesheet_id": timesheet.id,
            "sheet_no": timesheet.sheet_no,
            "project_id": timesheet.project_id,
            "work_date": str(timesheet.work_date),
            "status": timesheet.status,
        },
        object_type="Timesheet",
    )
    return timesheet


def _period_date(year: int, month: int) -> date_type:
    return date_type(int(year), int(month), 1)


def _next_payroll_batch_no(*, year: int, month: int) -> str:
    key = f"PA-{year:04d}{month:02d}"
    seq, _ = PayrollBatchNumberSequence.objects.select_for_update().get_or_create(
        key=key, defaults={"last_number": 0}
    )
    seq.last_number += 1
    seq.save(update_fields=["last_number"])
    return f"{key}-{seq.last_number:04d}"


def _ensure_payroll_editable(batch: PayrollAllocationBatch):
    if batch.status not in (PayrollAllocationStatus.DRAFT, PayrollAllocationStatus.REJECTED):
        raise PermissionDenied("제출 이후에는 수정할 수 없습니다.")


def create_payroll_batch(*, year: int, month: int, total_amount: int, actor, note: str = ""):
    require_role(actor, [Role.HQ, Role.CEO])
    period_date = _period_date(year, month)
    _assert_month_open(None, period_date, message_prefix="급여 배부 생성은 불가능합니다.")
    if total_amount is None:
        raise ValidationError({"total_amount": "total_amount is required."})
    total_amount = int(total_amount)
    if total_amount < 0:
        raise ValidationError({"total_amount": "total_amount must be >= 0."})
    if PayrollAllocationBatch.objects.filter(
        period_year=year, period_month=month
    ).exists():
        raise ValidationError("해당 월 배부가 이미 존재합니다.")
    with transaction.atomic():
        batch_no = _next_payroll_batch_no(year=year, month=month)
        batch = PayrollAllocationBatch.objects.create(
            batch_no=batch_no,
            period_year=year,
            period_month=month,
            total_amount=total_amount,
            note=note or "",
            created_by=actor,
        )
    _log_action_safe(
        actor=actor,
        action="PAYROLL_BATCH_CREATE",
        object_id=batch.id,
        summary=f"Payroll batch create: {batch.batch_no}",
        metadata={
            "batch_id": batch.id,
            "batch_no": batch.batch_no,
            "year": batch.period_year,
            "month": batch.period_month,
            "total_amount": batch.total_amount,
            "status": batch.status,
        },
        object_type="PayrollAllocationBatch",
    )
    return batch


def update_payroll_batch(batch: PayrollAllocationBatch, data: dict, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    _ensure_payroll_editable(batch)
    period_date = _period_date(batch.period_year, batch.period_month)
    _assert_month_open(None, period_date, message_prefix="급여 배부 수정은 불가능합니다.")
    if "total_amount" in data and data["total_amount"] not in (None, ""):
        total_amount = int(data["total_amount"])
        if total_amount < 0:
            raise ValidationError({"total_amount": "total_amount must be >= 0."})
        batch.total_amount = total_amount
    if "note" in data:
        batch.note = str(data.get("note") or "")
    batch.save(update_fields=["total_amount", "note", "updated_at"])
    _log_action_safe(
        actor=actor,
        action="PAYROLL_BATCH_UPDATE",
        object_id=batch.id,
        summary=f"Payroll batch update: {batch.batch_no}",
        metadata={
            "batch_id": batch.id,
            "batch_no": batch.batch_no,
            "year": batch.period_year,
            "month": batch.period_month,
            "total_amount": batch.total_amount,
            "status": batch.status,
        },
        object_type="PayrollAllocationBatch",
    )
    return batch


def upsert_payroll_lines(
    batch: PayrollAllocationBatch, lines_payload: list[dict], *, actor
):
    require_role(actor, [Role.HQ, Role.CEO])
    _ensure_payroll_editable(batch)
    period_date = _period_date(batch.period_year, batch.period_month)
    _assert_month_open(None, period_date, message_prefix="급여 배부 수정은 불가능합니다.")
    cleaned: list[dict] = []
    for raw in lines_payload:
        project_id = raw.get("project_id") or raw.get("project")
        amount = raw.get("amount")
        if not project_id and amount in (None, "", 0):
            continue
        if not project_id:
            raise ValidationError("프로젝트를 선택해 주세요.")
        project = Project.objects.filter(id=project_id).first()
        if not project:
            raise ValidationError("유효한 프로젝트를 선택해 주세요.")
        if amount in (None, ""):
            raise ValidationError("배부 금액을 입력해 주세요.")
        amount_val = int(amount)
        if amount_val < 0:
            raise ValidationError("배부 금액은 0 이상이어야 합니다.")
        cbs = _resolve_cost_item(raw.get("cbs") or raw.get("cbs_id"))
        if cbs and hasattr(cbs, "is_active") and not cbs.is_active:
            raise ValidationError("비활성 CBS는 선택할 수 없습니다.")
        cleaned.append(
            {
                "project": project,
                "amount": amount_val,
                "cbs": cbs,
                "memo": str(raw.get("memo") or ""),
            }
        )
    if not cleaned:
        raise ValidationError("배부 라인은 최소 1개 이상 입력해야 합니다.")
    with transaction.atomic():
        PayrollAllocationLine.objects.filter(batch=batch).delete()
        PayrollAllocationLine.objects.bulk_create(
            [
                PayrollAllocationLine(
                    batch=batch,
                    project=line["project"],
                    cbs=line["cbs"],
                    amount=line["amount"],
                    memo=line["memo"],
                )
                for line in cleaned
            ]
        )
        batch.save(update_fields=["updated_at"])
    _log_action_safe(
        actor=actor,
        action="PAYROLL_LINES_UPDATE",
        object_id=batch.id,
        summary=f"Payroll lines update: {batch.batch_no}",
        metadata={
            "batch_id": batch.id,
            "batch_no": batch.batch_no,
            "year": batch.period_year,
            "month": batch.period_month,
            "line_count": len(cleaned),
        },
        object_type="PayrollAllocationBatch",
    )
    return batch


def validate_payroll_batch(batch: PayrollAllocationBatch):
    total_lines = (
        PayrollAllocationLine.objects.filter(batch=batch).aggregate(total=Sum("amount"))[
            "total"
        ]
        or 0
    )
    diff = int(batch.total_amount) - int(total_lines)
    return diff == 0, diff, int(total_lines)


def submit_payroll_batch(batch: PayrollAllocationBatch, *, actor):
    require_role(actor, [Role.HQ, Role.CEO])
    _ensure_payroll_editable(batch)
    period_date = _period_date(batch.period_year, batch.period_month)
    _assert_month_open(None, period_date, message_prefix="급여 배부 제출은 불가능합니다.")
    line_count = PayrollAllocationLine.objects.filter(batch=batch).count()
    if line_count == 0:
        raise ValidationError("배부 라인은 최소 1개 이상 입력해야 합니다.")
    ok, diff, sum_lines = validate_payroll_batch(batch)
    if not ok:
        _log_action_safe(
            actor=actor,
            action="PAYROLL_BATCH_VALIDATE_FAIL",
            object_id=batch.id,
            summary=f"Payroll batch validate fail: {batch.batch_no}",
            metadata={
                "batch_id": batch.id,
                "batch_no": batch.batch_no,
                "year": batch.period_year,
                "month": batch.period_month,
                "sum_lines": sum_lines,
                "total_amount": batch.total_amount,
                "diff": diff,
            },
            object_type="PayrollAllocationBatch",
        )
        raise ValidationError("배부 합계가 총액과 일치하지 않습니다.")
    batch.status = PayrollAllocationStatus.SUBMITTED
    batch.submitted_at = timezone.now()
    batch.save(update_fields=["status", "submitted_at", "updated_at"])
    _log_action_safe(
        actor=actor,
        action="PAYROLL_BATCH_SUBMIT",
        object_id=batch.id,
        summary=f"Payroll batch submit: {batch.batch_no}",
        metadata={
            "batch_id": batch.id,
            "batch_no": batch.batch_no,
            "year": batch.period_year,
            "month": batch.period_month,
            "sum_lines": sum_lines,
            "total_amount": batch.total_amount,
            "status": batch.status,
        },
        object_type="PayrollAllocationBatch",
    )
    return batch, sum_lines, diff
