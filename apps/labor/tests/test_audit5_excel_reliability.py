import json
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.audit.models import AuditLog
from apps.core.rbac.models import Role, UserProfile
from apps.labor.models import (
    ElectronicCardImportBatch,
    ElectronicCardImportBatchStatus,
    ElectronicCardWorkDay,
    LaborExcelExportStatus,
    LaborRole,
    LaborWorkLedgerStatus,
    WorkerMaster,
)
from apps.labor.services import (
    E_CARD_HEADER_LABELS,
    confirm_electronic_card_reconciliation_batch,
    create_electronic_card_import_batch,
    create_labor_work_ledger,
    generate_cwma_card_reupload_excel,
    parse_electronic_card_import_batch,
    reconcile_electronic_card_import_batch,
    register_labor_excel_export_download,
)
from apps.projects.models import Project


IDENTITY_A = "AUDIT5-A-IDENTITY"
IDENTITY_B = "AUDIT5-B-IDENTITY"
ACCOUNT_A = "AUDIT5-ACCOUNT-A"


@pytest.fixture(autouse=True)
def _disable_two_factor_enforce(settings):
    settings.MIDDLEWARE = [
        middleware
        for middleware in settings.MIDDLEWARE
        if middleware != "apps.core.middleware.two_factor_enforce.TwoFactorEnforceMiddleware"
    ]


def _user(role, username):
    user = get_user_model().objects.create_user(username=username, password="pass")
    UserProfile.objects.create(user=user, role=role)
    return user


def _project(code="AUDIT5-PROJ", name="AUDIT5 실제 현장"):
    return Project.objects.create(code=code, name=name, project_type="civil")


def _labor_role(code="AUDIT5-LABOR", name="보통인부"):
    return LaborRole.objects.create(code=code, name=name, is_active=True)


def _worker(name, identity, phone, role=None, *, account_number=""):
    worker = WorkerMaster(
        name=name,
        phone=phone,
        default_labor_role=role,
        active=True,
        bank_name="테스트은행",
        bank_code="999",
        account_holder=name,
    )
    worker.set_rrn(identity)
    worker.set_account_number(account_number)
    worker.save()
    return worker


def _header(field_name):
    return E_CARD_HEADER_LABELS[field_name][0]


def _ecard_upload(rows, *, name="audit5_ecard.xlsx", headers=None, sheet_name="전자카드"):
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet_name
    base_headers = headers or [
        _header("year_month"),
        _header("project_name"),
        _header("deduction_join_no"),
        _header("company_name"),
        _header("job_type"),
        _header("worker_name"),
        _header("resident_no"),
        _header("phone"),
        _header("retirement_deduction"),
        _header("exclusion_reason"),
        _header("reported_days"),
        _header("confirmed_days"),
        _header("report_status"),
        _header("note"),
    ]
    ws.append([*base_headers, *range(1, 32)])
    for row in rows:
        if row.get("_blank"):
            ws.append(["" for _ in range(len(base_headers) + 31)])
            continue
        day_values = row.get("days", {})
        ws.append(
            [
                row.get("year_month", "2026-05"),
                row.get("project_name", "AUDIT5 신고 현장"),
                row.get("deduction_join_no", "AUDIT5-JOIN"),
                row.get("company_name", "AUDIT5 건설"),
                row.get("job_name", "보통인부"),
                row.get("name", ""),
                row.get("identity", ""),
                row.get("phone", ""),
                "Y",
                row.get("exclusion_reason", ""),
                row.get("reported_days", ""),
                row.get("confirmed_days", ""),
                row.get("report_status", "정상"),
                row.get("note", ""),
                *[day_values.get(day, "") for day in range(1, 32)],
            ]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _corrupt_xlsx(name="audit5_corrupt.xlsx"):
    return SimpleUploadedFile(
        name,
        b"not-a-real-openxml-workbook",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _create_ledger(*, actor, worker, project, role, work_date, work_unit="1.00"):
    return create_labor_work_ledger(
        {
            "worker": worker,
            "work_date": work_date,
            "actual_project": project,
            "report_project": project,
            "labor_role": role,
            "work_unit": Decimal(str(work_unit)),
            "work_hours": Decimal("8.00"),
            "unit_wage": 100000,
            "status": LaborWorkLedgerStatus.CONFIRMED,
        },
        actor=actor,
    )


def _create_confirmed_batch(settings, tmp_path, *, username="audit5-hq-confirmed"):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, username)
    project = _project(f"AUDIT5-PROJ-{username[-8:]}", "AUDIT5 실제 현장")
    role = _labor_role(f"AUDIT5-ROLE-{username[-8:]}")
    worker = _worker("홍길동", IDENTITY_A, "010-1111-2222", role, account_number=ACCOUNT_A)
    _create_ledger(
        actor=hq,
        worker=worker,
        project=project,
        role=role,
        work_date=date(2026, 5, 1),
        work_unit="1.00",
    )
    upload = _ecard_upload(
        [
            {
                "name": "홍길동",
                "identity": IDENTITY_A,
                "phone": "010-1111-2222",
                "days": {1: 1},
            }
        ]
    )
    batch = create_electronic_card_import_batch(
        {"year_month": "2026-05", "project": project.id},
        upload,
        hq,
    )
    parse_electronic_card_import_batch(batch, hq)
    reconcile_electronic_card_import_batch(batch, hq)
    confirm_electronic_card_reconciliation_batch(batch, hq)
    batch.refresh_from_db()
    return {"hq": hq, "project": project, "role": role, "worker": worker, "batch": batch}


def _audit_payload():
    rows = list(
        AuditLog.objects.filter(action__startswith="LABOR")
        .values("action", "before_json", "after_json", "meta_json")
        .order_by("id")
    )
    return json.dumps(rows, ensure_ascii=False, default=str)


def _assert_no_raw_pii_in_audit_logs(*workers):
    payload = _audit_payload()
    for secret in (IDENTITY_A, IDENTITY_B, ACCOUNT_A):
        assert secret not in payload
    for worker in workers:
        if worker.identity_hash:
            assert worker.identity_hash not in payload
    assert "enc1:" not in payload


def _streaming_bytes(response):
    return b"".join(response.streaming_content)


@pytest.mark.django_db
def test_audit5_ecard_import_rejects_non_xlsx_or_corrupt_workbook(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, "audit5-hq-invalid")
    project = _project("AUDIT5-INVALID", "AUDIT5 오류 현장")

    with pytest.raises(ValidationError):
        create_electronic_card_import_batch(
            {"year_month": "2026-05", "project": project.id},
            SimpleUploadedFile("ecard.txt", b"not xlsx", content_type="text/plain"),
            hq,
        )
    with pytest.raises(ValidationError):
        create_electronic_card_import_batch(
            {"year_month": "2026-05", "project": project.id},
            _corrupt_xlsx(),
            hq,
        )

    assert ElectronicCardImportBatch.objects.count() == 0


@pytest.mark.django_db
def test_audit5_ecard_import_rejects_missing_required_headers(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, "audit5-hq-missing-header")
    project = _project("AUDIT5-MISSING", "AUDIT5 누락 현장")
    headers = [_header("year_month"), _header("project_name"), _header("worker_name")]

    with pytest.raises(ValidationError):
        create_electronic_card_import_batch(
            {"year_month": "2026-05", "project": project.id},
            _ecard_upload(
                [{"name": "홍길동", "identity": IDENTITY_A, "days": {1: 1}}],
                headers=headers,
            ),
            hq,
        )

    assert ElectronicCardImportBatch.objects.count() == 0


@pytest.mark.django_db
def test_audit5_ecard_import_preserves_korean_headers_and_values(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, "audit5-hq-korean")
    project = _project("AUDIT5-KO", "AUDIT5 한글 현장")
    role = _labor_role("AUDIT5-KO-ROLE")
    worker = _worker("홍길동", IDENTITY_A, "01011112222", role)
    batch = create_electronic_card_import_batch(
        {"year_month": "2026-05", "project": project.id},
        _ecard_upload(
            [
                {
                    "project_name": "국도46호선 포장 보수공사",
                    "company_name": "대한건설",
                    "job_name": "보통인부",
                    "name": "홍길동",
                    "identity": IDENTITY_A,
                    "phone": "010-1111-2222",
                    "days": {1: 1},
                }
            ],
            sheet_name="전자카드",
        ),
        hq,
    )
    parse_electronic_card_import_batch(batch, hq)

    raw = batch.raw_rows.get()
    assert raw.worker_name_raw == "홍길동"
    assert raw.project_name_raw == "국도46호선 포장 보수공사"
    assert raw.company_name == "대한건설"
    assert raw.matched_worker == worker
    assert batch.header_check_summary["sheet_name"] == "전자카드"
    assert chr(0xFFFD) not in json.dumps(batch.header_check_summary, ensure_ascii=False, default=str)


@pytest.mark.django_db
def test_audit5_ecard_import_parses_numeric_day_values_deterministically(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, "audit5-hq-numeric")
    project = _project("AUDIT5-NUM", "AUDIT5 숫자 현장")
    role = _labor_role("AUDIT5-NUM-ROLE")
    worker = _worker("홍길동", IDENTITY_A, "01011112222", role)
    batch = create_electronic_card_import_batch(
        {"year_month": "2026-05", "project": project.id},
        _ecard_upload(
            [
                {
                    "name": "홍길동",
                    "identity": IDENTITY_A,
                    "phone": "010-1111-2222",
                    "days": {1: 1, 2: 0.5, 3: "abc"},
                },
                {"_blank": True},
            ]
        ),
        hq,
    )
    parse_electronic_card_import_batch(batch, hq)

    assert batch.raw_rows.count() == 1
    day1 = batch.day_rows.get(work_date=date(2026, 5, 1))
    day2 = batch.day_rows.get(work_date=date(2026, 5, 2))
    day3 = batch.day_rows.get(work_date=date(2026, 5, 3))
    assert day1.worker == worker
    assert day1.card_value == Decimal("1.00")
    assert day2.card_value == Decimal("0.50")
    assert day3.card_value == Decimal("0.00")
    assert day3.note
    batch.refresh_from_db()
    assert batch.header_check_summary["parsed_raw_count"] == 1
    assert batch.header_check_summary["parsed_worked_day_count"] == 2
    assert batch.header_check_summary["parsed_unreadable_day_count"] == 1


@pytest.mark.django_db
def test_audit5_ecard_import_does_not_log_raw_pii_on_failure(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, "audit5-hq-privacy-failure")
    project = _project("AUDIT5-PRIVACY", "AUDIT5 개인정보 현장")

    with pytest.raises(ValidationError):
        create_electronic_card_import_batch(
            {"year_month": "2026-05", "project": project.id},
            _corrupt_xlsx("audit5-privacy.xlsx"),
            hq,
        )

    _assert_no_raw_pii_in_audit_logs()


@pytest.mark.django_db
def test_audit5_cwma_export_generates_valid_xlsx_with_korean_sheet_and_headers(settings, tmp_path):
    context = _create_confirmed_batch(settings, tmp_path)
    batch = context["batch"]
    export_batch = generate_cwma_card_reupload_excel(batch, context["hq"], note="AUDIT5 export")

    export_batch.generated_file.open("rb")
    workbook = load_workbook(export_batch.generated_file, data_only=True)
    sheet = workbook.active

    assert sheet.title == "전자카드"
    assert sheet.cell(row=1, column=1).value == _header("year_month")
    assert sheet.cell(row=2, column=6).value == "홍길동"
    assert sheet.cell(row=2, column=15).value == 1
    assert sheet.cell(row=2, column=11).value == 1
    assert sheet.cell(row=2, column=12).value == 1
    assert export_batch.status == LaborExcelExportStatus.GENERATED
    assert export_batch.changed_count > 0
    assert export_batch.included_count == 1
    assert export_batch.excluded_count == 0
    assert export_batch.total_export_work_unit == Decimal("1.00")


@pytest.mark.django_db
def test_audit5_cwma_export_download_registers_metadata_and_audit(settings, tmp_path):
    context = _create_confirmed_batch(settings, tmp_path, username="audit5-hq-download-service")
    export_batch = generate_cwma_card_reupload_excel(context["batch"], context["hq"])

    register_labor_excel_export_download(export_batch, context["hq"])
    export_batch.refresh_from_db()

    assert export_batch.status == LaborExcelExportStatus.DOWNLOADED
    assert export_batch.downloaded_by == context["hq"]
    assert export_batch.downloaded_at is not None
    assert AuditLog.objects.filter(action="LABOR_EXCEL_EXPORT_DOWNLOAD", object_id=export_batch.id).exists()


@pytest.mark.django_db
def test_audit5_cwma_export_requires_confirmed_batch(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, "audit5-hq-export-precondition")
    project = _project("AUDIT5-NOT-CONFIRMED", "AUDIT5 미확정 현장")
    batch = create_electronic_card_import_batch(
        {"year_month": "2026-05", "project": project.id},
        _ecard_upload([{"name": "홍길동", "identity": IDENTITY_A, "days": {1: 1}}]),
        hq,
    )
    assert batch.status == ElectronicCardImportBatchStatus.PARSE_READY

    with pytest.raises(ValidationError):
        generate_cwma_card_reupload_excel(batch, hq)


@pytest.mark.django_db
def test_audit5_cwma_export_does_not_corrupt_source_file(settings, tmp_path):
    context = _create_confirmed_batch(settings, tmp_path, username="audit5-hq-source-safe")
    batch = context["batch"]
    batch.source_file.open("rb")
    source_before = batch.source_file.read()

    generate_cwma_card_reupload_excel(batch, context["hq"])

    batch.source_file.open("rb")
    source_after = batch.source_file.read()
    assert source_after == source_before
    workbook = load_workbook(BytesIO(source_after), data_only=True)
    assert workbook.active.cell(row=2, column=15).value == 1


@pytest.mark.django_db
def test_audit5_excel_download_requires_hq_or_ceo(settings, tmp_path, client):
    context = _create_confirmed_batch(settings, tmp_path, username="audit5-hq-field-block")
    export_batch = generate_cwma_card_reupload_excel(context["batch"], context["hq"])
    field = _user(Role.FIELD, "audit5-field-download")
    client.force_login(field)

    response = client.get(f"/app/hq/labor/excel-exports/{export_batch.id}/download/")

    assert response.status_code == 403
    export_batch.refresh_from_db()
    assert export_batch.status == LaborExcelExportStatus.GENERATED
    assert export_batch.downloaded_by is None


@pytest.mark.django_db
def test_audit5_excel_download_returns_404_or_validation_for_missing_file(settings, tmp_path, client):
    context = _create_confirmed_batch(settings, tmp_path, username="audit5-hq-missing-file")
    export_batch = generate_cwma_card_reupload_excel(context["batch"], context["hq"])
    missing_path = Path(export_batch.generated_file.path)
    missing_path.unlink()
    client.force_login(context["hq"])

    response = client.get(f"/app/hq/labor/excel-exports/{export_batch.id}/download/")

    assert response.status_code == 404
    export_batch.refresh_from_db()
    assert export_batch.status == LaborExcelExportStatus.GENERATED
    assert export_batch.downloaded_by is None


@pytest.mark.django_db
def test_audit5_excel_download_sets_xlsx_content_type_or_attachment_header(settings, tmp_path, client):
    context = _create_confirmed_batch(settings, tmp_path, username="audit5-hq-download-view")
    export_batch = generate_cwma_card_reupload_excel(context["batch"], context["hq"])
    client.force_login(context["hq"])

    response = client.get(f"/app/hq/labor/excel-exports/{export_batch.id}/download/")

    assert response.status_code == 200
    assert ".xlsx" in response.headers["Content-Disposition"]
    assert "attachment" in response.headers["Content-Disposition"]
    downloaded_workbook = load_workbook(BytesIO(_streaming_bytes(response)), data_only=True)
    assert downloaded_workbook.active.title == "전자카드"
    export_batch.refresh_from_db()
    assert export_batch.status == LaborExcelExportStatus.DOWNLOADED
    assert export_batch.downloaded_by == context["hq"]


@pytest.mark.django_db
def test_audit5_excel_import_export_audit_logs_do_not_expose_raw_pii(settings, tmp_path):
    context = _create_confirmed_batch(settings, tmp_path, username="audit5-hq-audit-privacy")
    export_batch = generate_cwma_card_reupload_excel(context["batch"], context["hq"])
    register_labor_excel_export_download(export_batch, context["hq"])

    _assert_no_raw_pii_in_audit_logs(context["worker"])
