import json
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.audit.models import AuditLog
from apps.core.rbac.models import Role, UserProfile
from apps.labor.models import (
    ElectronicCardImportBatch,
    ElectronicCardImportBatchStatus,
    LaborConfirmedWorkDay,
    LaborExcelExportBatch,
    LaborExcelExportStatus,
    LaborMonthlyPayroll,
    LaborReconciliationResolution,
    LaborReconciliationResult,
    LaborReconciliationStatus,
    LaborRole,
    LaborWorkLedger,
    LaborWorkLedgerStatus,
    PayrollAllocationBatch,
    PayrollAllocationStatus,
    WorkerMaster,
)
from apps.labor.services import (
    E_CARD_HEADER_LABELS,
    bulk_resolve_labor_reconciliation_results,
    confirm_electronic_card_reconciliation_batch,
    create_electronic_card_import_batch,
    create_labor_work_ledger,
    create_payroll_batch,
    create_worker_master,
    generate_cwma_card_reupload_excel,
    generate_labor_monthly_payroll,
    match_reconciliation_worker,
    parse_electronic_card_import_batch,
    reconcile_electronic_card_import_batch,
    register_labor_excel_export_download,
    resolve_labor_reconciliation_result,
    submit_payroll_batch,
    update_labor_reporting_project,
    upsert_payroll_lines,
    validate_payroll_batch,
)
from apps.projects.models import Project


IDENTITY_A = "AUDIT6-A-IDENTITY"
IDENTITY_B = "AUDIT6-B-IDENTITY"
IDENTITY_C = "AUDIT6-C-IDENTITY"
ACCOUNT_A = "AUDIT6-ACCOUNT-A"


@pytest.fixture(autouse=True)
def _disable_two_factor_enforce(settings):
    settings.MIDDLEWARE = [
        middleware
        for middleware in settings.MIDDLEWARE
        if middleware != "apps.core.middleware.two_factor_enforce.TwoFactorEnforceMiddleware"
    ]


def _user(role, username):
    user = get_user_model().objects.create_user(username=username, password="pass")
    UserProfile.objects.create(user=user, role=role)
    return user


def _project(code, name):
    return Project.objects.create(code=code, name=name, project_type="civil")


def _labor_role(code="AUDIT6-LABOR", name="보통인부"):
    return LaborRole.objects.create(code=code, name=name, is_active=True)


def _worker(actor, name, identity, phone, role=None, account_number=""):
    return create_worker_master(
        {
            "name": name,
            "rrn": identity,
            "phone": phone,
            "bank_name": "테스트은행",
            "bank_code": "999",
            "account_number": account_number,
            "account_holder": name,
            "default_labor_role": role,
            "retirement_deduction_eligible": True,
            "active": True,
        },
        actor=actor,
    )


def _header(field_name):
    return E_CARD_HEADER_LABELS[field_name][0]


def _ecard_upload(rows, *, name="audit6_ecard.xlsx"):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "전자카드"
    base_headers = [
        _header("year_month"),
        _header("project_name"),
        _header("deduction_join_no"),
        _header("company_name"),
        _header("job_type"),
        _header("worker_name"),
        _header("resident_no"),
        _header("phone"),
        _header("retirement_deduction"),
        _header("exclusion_reason"),
        _header("reported_days"),
        _header("confirmed_days"),
        _header("report_status"),
        _header("note"),
    ]
    ws.append([*base_headers, *range(1, 32)])
    for row in rows:
        day_values = row.get("days", {})
        ws.append(
            [
                "2026-05",
                row.get("project_name", "AUDIT6 신고 현장"),
                "AUDIT6-JOIN",
                "AUDIT6 건설",
                row.get("job_name", "보통인부"),
                row["name"],
                row["identity"],
                row.get("phone", ""),
                "Y",
                "",
                "",
                "",
                "정상",
                row.get("note", ""),
                *[day_values.get(day, "") for day in range(1, 32)],
            ]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _create_ledger(
    *,
    actor,
    worker,
    project,
    role,
    work_date,
    work_unit,
    report_project=None,
    status=LaborWorkLedgerStatus.CONFIRMED,
    unit_wage=100000,
):
    ledger = create_labor_work_ledger(
        {
            "worker": worker,
            "work_date": work_date,
            "actual_project": project,
            "report_project": report_project or project,
            "labor_role": role,
            "work_unit": Decimal(str(work_unit)),
            "work_hours": Decimal("8.00"),
            "unit_wage": unit_wage,
            "status": status,
        },
        actor=actor,
    )
    return ledger


def _create_confirmed_e_card_flow(settings, tmp_path, *, username="audit6-hq-flow"):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, username)
    project = _project(f"AUDIT6-PROJ-{username}", "AUDIT6 실제 현장")
    role = _labor_role(f"AUDIT6-ROLE-{username[-8:]}")
    worker_a = _worker(hq, "홍길동", IDENTITY_A, "010-1111-2222", role, ACCOUNT_A)
    worker_b = _worker(hq, "김현장", IDENTITY_B, "01033334444", role)
    _create_ledger(
        actor=hq,
        worker=worker_a,
        project=project,
        role=role,
        work_date=date(2026, 5, 1),
        work_unit="1.00",
    )
    _create_ledger(
        actor=hq,
        worker=worker_a,
        project=project,
        role=role,
        work_date=date(2026, 5, 2),
        work_unit="1.00",
    )
    _create_ledger(
        actor=hq,
        worker=worker_b,
        project=project,
        role=role,
        work_date=date(2026, 5, 1),
        work_unit="0.50",
    )
    upload = _ecard_upload(
        [
            {
                "name": "홍길동",
                "identity": IDENTITY_A,
                "phone": "010-1111-2222",
                "days": {1: 1, 2: 1},
            },
            {
                "name": "김현장",
                "identity": IDENTITY_B,
                "phone": "010-3333-4444",
                "days": {1: 1},
            },
            {
                "name": "미등록자",
                "identity": IDENTITY_C,
                "phone": "010-5555-6666",
                "days": {3: 1},
                "note": "제외 검토",
            },
        ]
    )
    batch = create_electronic_card_import_batch(
        {"year_month": "2026-05", "project": project.id},
        upload,
        hq,
    )
    parse_electronic_card_import_batch(batch, hq)
    batch.refresh_from_db()
    reconcile_electronic_card_import_batch(batch, hq)
    batch.refresh_from_db()
    return {
        "hq": hq,
        "project": project,
        "role": role,
        "worker_a": worker_a,
        "worker_b": worker_b,
        "batch": batch,
    }


def _resolve_non_match_results(batch, actor):
    diff = batch.reconciliation_results.get(status=LaborReconciliationStatus.DIFF)
    resolve_labor_reconciliation_result(
        diff,
        LaborReconciliationResolution.CARD,
        actor,
        comment="카드 기준으로 확정",
    )
    unmatched = batch.reconciliation_results.get(status=LaborReconciliationStatus.UNMATCHED)
    resolve_labor_reconciliation_result(
        unmatched,
        LaborReconciliationResolution.EXCLUDED,
        actor,
        comment="미등록 근로자 신고 제외",
        export_note="신고 제외",
    )


def _confirm_flow(context):
    batch = context["batch"]
    hq = context["hq"]
    _resolve_non_match_results(batch, hq)
    confirm_electronic_card_reconciliation_batch(batch, hq)
    batch.refresh_from_db()
    return batch


def _audit_payload():
    rows = list(
        AuditLog.objects.filter(action__startswith="LABOR")
        .values("action", "before_json", "after_json", "meta_json")
        .order_by("id")
    )
    rows += list(
        AuditLog.objects.filter(action__startswith="PAYROLL")
        .values("action", "before_json", "after_json", "meta_json")
        .order_by("id")
    )
    return json.dumps(rows, ensure_ascii=False, default=str)


def _assert_no_raw_pii_in_audit_logs(*workers):
    payload = _audit_payload()
    for secret in (IDENTITY_A, IDENTITY_B, IDENTITY_C, ACCOUNT_A):
        assert secret not in payload
    for worker in workers:
        if worker.identity_hash:
            assert worker.identity_hash not in payload
    assert "enc1:" not in payload


@pytest.mark.django_db
def test_audit6_labpay_upload_parse_reconcile_resolve_confirm_export_e2e(settings, tmp_path):
    context = _create_confirmed_e_card_flow(settings, tmp_path)
    batch = context["batch"]
    hq = context["hq"]

    assert batch.raw_rows.count() == 3
    assert batch.day_rows.count() == 93
    assert batch.header_check_summary["matched_worker_count"] == 2
    assert batch.header_check_summary["unmatched_worker_count"] == 1
    assert batch.header_check_summary["parsed_worked_day_count"] == 4
    assert LaborReconciliationResult.objects.filter(
        batch=batch, status=LaborReconciliationStatus.MATCH
    ).count() == 2
    assert LaborReconciliationResult.objects.filter(
        batch=batch, status=LaborReconciliationStatus.DIFF
    ).count() == 1
    assert LaborReconciliationResult.objects.filter(
        batch=batch, status=LaborReconciliationStatus.UNMATCHED
    ).count() == 1
    with pytest.raises(ValidationError):
        confirm_electronic_card_reconciliation_batch(batch, hq)

    _confirm_flow(context)

    assert batch.status == ElectronicCardImportBatchStatus.CONFIRMED
    assert batch.confirmed_by == hq
    assert LaborConfirmedWorkDay.objects.filter(batch=batch).count() == 4
    assert LaborConfirmedWorkDay.objects.filter(batch=batch, export_included=True).count() == 3
    assert LaborConfirmedWorkDay.objects.filter(batch=batch, export_included=False).count() == 1
    assert AuditLog.objects.filter(action="LABOR_ECARD_IMPORT_PARSE", object_id=batch.id).exists()
    assert AuditLog.objects.filter(action="LABOR_ECARD_RECONCILE", object_id=batch.id).exists()
    assert AuditLog.objects.filter(action="LABOR_CONFIRMED_WORKDAY_GENERATE", object_id=batch.id).exists()
    assert AuditLog.objects.filter(action="LABOR_ECARD_RECONCILIATION_CONFIRM", object_id=batch.id).exists()

    export_batch = generate_cwma_card_reupload_excel(batch, hq, note="AUDIT6 export")
    assert export_batch.status == LaborExcelExportStatus.GENERATED
    assert export_batch.generated_filename.endswith(".xlsx")
    assert export_batch.generated_file
    assert export_batch.included_count == 3
    assert export_batch.excluded_count == 1
    assert export_batch.total_export_work_unit == Decimal("3.00")
    assert AuditLog.objects.filter(action="LABOR_EXCEL_EXPORT_GENERATE", object_id=export_batch.id).exists()

    export_batch.generated_file.open("rb")
    generated_workbook = load_workbook(export_batch.generated_file, data_only=True)
    generated_sheet = generated_workbook.active
    assert generated_sheet.title == "전자카드"
    assert generated_sheet.cell(row=2, column=6).value == "홍길동"
    assert generated_sheet.cell(row=2, column=15).value == 1
    assert generated_sheet.cell(row=2, column=16).value == 1
    assert generated_sheet.cell(row=2, column=11).value == 2
    assert generated_sheet.cell(row=2, column=12).value == 2
    assert generated_sheet.cell(row=4, column=17).value is None
    assert "신고 제외" in str(generated_sheet.cell(row=4, column=10).value)

    batch.source_file.open("rb")
    source_workbook = load_workbook(batch.source_file, data_only=True)
    assert source_workbook.active.cell(row=4, column=17).value == 1

    register_labor_excel_export_download(export_batch, hq)
    export_batch.refresh_from_db()
    assert export_batch.status == LaborExcelExportStatus.DOWNLOADED
    assert export_batch.downloaded_by == hq
    assert export_batch.downloaded_at is not None
    assert AuditLog.objects.filter(action="LABOR_EXCEL_EXPORT_DOWNLOAD", object_id=export_batch.id).exists()
    _assert_no_raw_pii_in_audit_logs(context["worker_a"], context["worker_b"])


@pytest.mark.django_db
def test_audit6_confirmed_batch_blocks_reparse_reconcile_resolve_and_regenerate_conflicts(settings, tmp_path):
    context = _create_confirmed_e_card_flow(settings, tmp_path, username="audit6-hq-lock")
    batch = _confirm_flow(context)
    hq = context["hq"]
    result = batch.reconciliation_results.order_by("id").first()
    confirmed_count = batch.confirmed_work_days.count()
    confirmed_total = sum(row.final_work_unit for row in batch.confirmed_work_days.all())

    with pytest.raises(ValidationError):
        parse_electronic_card_import_batch(batch, hq)
    with pytest.raises(ValidationError):
        reconcile_electronic_card_import_batch(batch, hq)
    with pytest.raises(ValidationError):
        resolve_labor_reconciliation_result(result, LaborReconciliationResolution.ERP, hq, comment="block")
    with pytest.raises(ValidationError):
        bulk_resolve_labor_reconciliation_results([result.id], LaborReconciliationResolution.ERP, hq, "block")
    with pytest.raises(ValidationError):
        match_reconciliation_worker(result, context["worker_a"], hq)
    with pytest.raises(ValidationError):
        confirm_electronic_card_reconciliation_batch(batch, hq)

    batch.refresh_from_db()
    assert batch.confirmed_work_days.count() == confirmed_count
    assert sum(row.final_work_unit for row in batch.confirmed_work_days.all()) == confirmed_total


@pytest.mark.django_db
def test_audit6_reporting_project_mapping_flows_into_monthly_payroll_and_allocation():
    hq = _user(Role.HQ, "audit6-hq-payroll")
    actual_project = _project("AUDIT6-ACTUAL", "AUDIT6 실제 현장")
    report_project = _project("AUDIT6-REPORT", "AUDIT6 신고 현장")
    role = _labor_role("AUDIT6-PAY-ROLE")
    worker = _worker(hq, "박급여", "AUDIT6-PAY-IDENTITY", "010-7777-8888", role)
    ledger = _create_ledger(
        actor=hq,
        worker=worker,
        project=actual_project,
        role=role,
        work_date=date(2026, 5, 4),
        work_unit="2.00",
        status=LaborWorkLedgerStatus.DRAFT,
    )

    updated = update_labor_reporting_project([ledger.id], report_project, "신고 현장 보정", hq)
    assert len(updated) == 1
    ledger.refresh_from_db()
    assert ledger.actual_project == actual_project
    assert ledger.report_project == report_project
    ledger.status = LaborWorkLedgerStatus.CONFIRMED
    ledger.save(update_fields=["status", "updated_at"])

    summary = generate_labor_monthly_payroll(year=2026, month=5, actor=hq)
    assert summary["row_count"] == 1
    payroll = LaborMonthlyPayroll.objects.get(year_month=date(2026, 5, 1))
    assert payroll.project == actual_project
    assert payroll.report_project == report_project
    assert payroll.total_work_unit == Decimal("2.00")
    assert payroll.gross_wage == 200000
    assert payroll.net_pay == 200000
    assert AuditLog.objects.filter(action="LABOR_REPORTING_PROJECT_UPDATE", object_id=ledger.id).exists()
    assert AuditLog.objects.filter(action="LABOR_MONTHLY_PAYROLL_GENERATE").exists()

    allocation = create_payroll_batch(
        year=2026,
        month=5,
        total_amount=payroll.gross_wage,
        actor=hq,
        note="AUDIT6 payroll allocation",
    )
    upsert_payroll_lines(
        allocation,
        [{"project": report_project.id, "amount": payroll.gross_wage - 1, "memo": "invalid"}],
        actor=hq,
    )
    ok, diff, _sum_lines = validate_payroll_batch(allocation)
    assert not ok
    assert diff == 1
    with pytest.raises(ValidationError):
        submit_payroll_batch(allocation, actor=hq)

    upsert_payroll_lines(
        allocation,
        [{"project": report_project.id, "amount": payroll.gross_wage, "memo": "valid"}],
        actor=hq,
    )
    ok, diff, sum_lines = validate_payroll_batch(allocation)
    assert ok
    assert diff == 0
    assert sum_lines == payroll.gross_wage
    submit_payroll_batch(allocation, actor=hq)
    allocation.refresh_from_db()
    assert allocation.status == PayrollAllocationStatus.SUBMITTED
    assert PayrollAllocationBatch.objects.filter(id=allocation.id).exists()
    assert AuditLog.objects.filter(action="PAYROLL_BATCH_CREATE", object_id=allocation.id).exists()
    assert AuditLog.objects.filter(action="PAYROLL_LINES_UPDATE", object_id=allocation.id).exists()
    assert AuditLog.objects.filter(action="PAYROLL_BATCH_VALIDATE_FAIL", object_id=allocation.id).exists()
    assert AuditLog.objects.filter(action="PAYROLL_BATCH_SUBMIT", object_id=allocation.id).exists()


@pytest.mark.django_db
def test_audit6_field_user_cannot_run_hq_labpay_e2e_services(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hq = _user(Role.HQ, "audit6-hq-rbac")
    field = _user(Role.FIELD, "audit6-field-rbac")
    project = _project("AUDIT6-RBAC-PROJ", "AUDIT6 RBAC 현장")
    upload = _ecard_upload(
        [{"name": "권한차단", "identity": "AUDIT6-RBAC-IDENTITY", "phone": "010-0000-0000", "days": {1: 1}}],
        name="audit6_rbac.xlsx",
    )
    batch = create_electronic_card_import_batch(
        {"year_month": "2026-05", "project": project.id},
        upload,
        hq,
    )
    payroll_batch = create_payroll_batch(year=2026, month=5, total_amount=1000, actor=hq)

    with pytest.raises(PermissionDenied):
        create_electronic_card_import_batch(
            {"year_month": "2026-05", "project": project.id},
            _ecard_upload([], name="blocked.xlsx"),
            field,
        )
    assert ElectronicCardImportBatch.objects.count() == 1
    with pytest.raises(PermissionDenied):
        parse_electronic_card_import_batch(batch, field)
    with pytest.raises(PermissionDenied):
        reconcile_electronic_card_import_batch(batch, field)
    with pytest.raises(PermissionDenied):
        confirm_electronic_card_reconciliation_batch(batch, field)
    with pytest.raises(PermissionDenied):
        generate_cwma_card_reupload_excel(batch, field)
    with pytest.raises(PermissionDenied):
        generate_labor_monthly_payroll(year=2026, month=5, actor=field)
    with pytest.raises(PermissionDenied):
        create_payroll_batch(year=2026, month=6, total_amount=1000, actor=field)
    with pytest.raises(PermissionDenied):
        upsert_payroll_lines(
            payroll_batch,
            [{"project": project.id, "amount": 1000}],
            actor=field,
        )
    with pytest.raises(PermissionDenied):
        submit_payroll_batch(payroll_batch, actor=field)

    batch.refresh_from_db()
    payroll_batch.refresh_from_db()
    assert batch.status == ElectronicCardImportBatchStatus.PARSE_READY
    assert payroll_batch.status == PayrollAllocationStatus.DRAFT
    assert LaborExcelExportBatch.objects.count() == 0


@pytest.mark.django_db
def test_audit6_labpay_e2e_audit_logs_do_not_expose_raw_pii(settings, tmp_path):
    context = _create_confirmed_e_card_flow(settings, tmp_path, username="audit6-hq-privacy")
    batch = _confirm_flow(context)
    export_batch = generate_cwma_card_reupload_excel(batch, context["hq"], note="privacy")
    register_labor_excel_export_download(export_batch, context["hq"])

    _assert_no_raw_pii_in_audit_logs(context["worker_a"], context["worker_b"])

