from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.sessions.middleware import SessionMiddleware
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import PermissionDenied
from django.test import RequestFactory
from openpyxl import Workbook, load_workbook

from apps.audit.models import AuditLog
from apps.core.rbac.models import Role, UserProfile
from apps.labor.models import (
    ElectronicCardImportBatch,
    ElectronicCardMatchStatus,
    LaborExcelExportBatch,
    LaborConfirmedWorkDay,
    LaborReconciliationResult,
    LaborReconciliationStatus,
    LaborRole,
    ElectronicCardWorkDay,
    ElectronicCardWorkRaw,
    LaborWorkLedger,
    WorkerMaster,
)
from apps.labor.web_views import (
    hq_e_card_import_batch_detail,
    hq_e_card_import_batch_list,
    hq_confirmed_work_day_list,
    hq_worker_master_delete,
    hq_labor_excel_export_download,
)
from apps.projects.models import Project


def _build_user(username, role):
    user = get_user_model().objects.create_user(username=username, password="pass")
    UserProfile.objects.create(user=user, role=role)
    return user


def _build_request(user, method="get", data=None, files=None):
    factory = RequestFactory()
    path = "/app/hq/labor/e-card-imports/"
    payload = data or {}
    if method.lower() == "post":
        request = factory.post(path, data={**payload, **(files or {})})
    else:
        request = factory.get(path, data=payload)
    SessionMiddleware(lambda req: None).process_request(request)
    request.session.save()
    setattr(request, "_messages", FallbackStorage(request))
    request.user = user
    return request


def _build_detail_request(user, batch, method="get", data=None, query=None):
    factory = RequestFactory()
    path = f"/app/hq/labor/e-card-imports/{batch.id}/"
    payload = data or {}
    if method.lower() == "post":
        request = factory.post(path, data=payload)
    else:
        request = factory.get(path, data=query or {})
    SessionMiddleware(lambda req: None).process_request(request)
    request.session.save()
    setattr(request, "_messages", FallbackStorage(request))
    request.user = user
    return request


def _xlsx_upload(
    name="ecard.xlsx",
    *,
    year_month_header="근로년월",
    worker_name="홍길동",
    rrn="900101-1234567",
    phone="010-1234-5678",
    day_values=None,
):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "전자카드"
    ws.append(
        [
            year_month_header,
            "공사명",
            "공제가입번호",
            "업체명",
            "직종",
            "성명",
            "주민등록번호",
            "연락처",
            "퇴직공제여부",
            "비대상사유",
            "비고",
            "신고일수",
            "확정일수",
            "신고상태",
            *[f"{day}일" for day in range(1, 32)],
        ]
    )
    ws.append(
        [
            "2026-05",
            "포장 보수공사",
            "JOIN-001",
            "건설 주식회사",
            "포장공",
            worker_name,
            rrn,
            phone,
            "Y",
            "",
            "",
            "",
            "",
            "정상",
            *[
                (day_values or {}).get(day, "")
                for day in range(1, 32)
            ],
        ]
    )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _xlsx_upload_with_exclusion_header(
    *,
    exclusion_header="비대상사유",
    exclusion_value="",
):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "전자카드"
    ws.append(
        [
            "근로년월",
            "공사명",
            "공제가입번호",
            "업체명",
            "직종",
            "성명",
            "주민등록번호",
            "연락처",
            "퇴직공제여부",
            exclusion_header,
            "신고상태",
            *[f"{day}일" for day in range(1, 32)],
        ]
    )
    ws.append(
        [
            "2026-05",
            "포장 보수공사",
            "JOIN-001",
            "건설 주식회사",
            "현장공",
            "홍길동",
            "900101-1234567",
            "010-1234-5678",
            "Y",
            exclusion_value,
            "정상",
            *["" for _day in range(1, 32)],
        ]
    )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        "ecard_exclusion.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@pytest.mark.django_db
def test_hq_can_open_e_card_imports_page():
    user = _build_user("labor-hq", Role.HQ)

    response = hq_e_card_import_batch_list(_build_request(user))

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert "전자카드 근로내역 업로드" in content
    assert "업로드 이력" in content


@pytest.mark.django_db
def test_e_card_import_list_displays_grouped_relation_counts():
    user = _build_user("labor-hq-counts", Role.HQ)
    project = Project.objects.create(code="PRJ-LAB-COUNTS", name="전자카드 건수 현장")
    batch = ElectronicCardImportBatch.objects.create(
        year_month=date(2026, 5, 1),
        project=project,
        source_file="labor/e_card_imports/counts.xlsx",
        uploaded_by=user,
    )
    raw_one = ElectronicCardWorkRaw.objects.create(
        batch=batch,
        row_no=1,
        work_month=date(2026, 5, 1),
        worker_name_raw="작업자1",
    )
    raw_two = ElectronicCardWorkRaw.objects.create(
        batch=batch,
        row_no=2,
        work_month=date(2026, 5, 1),
        worker_name_raw="작업자2",
    )
    for raw, work_day in ((raw_one, 1), (raw_one, 2), (raw_two, 3)):
        ElectronicCardWorkDay.objects.create(
            batch=batch,
            raw=raw,
            work_date=date(2026, 5, work_day),
        )
    for work_day in (1, 2):
        LaborReconciliationResult.objects.create(
            batch=batch,
            year_month=date(2026, 5, 1),
            project=project,
            work_date=date(2026, 5, work_day),
        )

    response = hq_e_card_import_batch_list(_build_request(user))

    content = response.content.decode("utf-8")
    assert "Raw 건수</span> 2건" in content
    assert "일별 건수</span> 3건" in content
    assert "대사 검토" in content


@pytest.mark.django_db
def test_field_cannot_open_e_card_imports_page():
    user = _build_user("labor-field", Role.FIELD)

    with pytest.raises(PermissionDenied):
        hq_e_card_import_batch_list(_build_request(user))


@pytest.mark.django_db
def test_hq_uploads_xlsx_and_batch_is_created(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    user = _build_user("labor-hq-upload", Role.HQ)
    project = Project.objects.create(code="PRJ-LAB-001", name="노무 업로드 현장")
    upload = _xlsx_upload()

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"year_month": "2026-05", "project": str(project.id)},
            files={"source_file": upload},
        )
    )

    assert response.status_code == 302
    batch = ElectronicCardImportBatch.objects.get()
    assert batch.project == project
    assert batch.original_filename == "ecard.xlsx"
    assert batch.source_file.name
    assert batch.header_check_summary["sheet_name"] == "전자카드"
    assert "year_month" in batch.header_check_summary["detected_headers"]
    assert batch.header_check_summary["workbook_months"] == ["2026-05"]
    assert batch.header_check_summary["cwma_project_name"] == "포장 보수공사"
    assert AuditLog.objects.filter(
        action="LABOR_ECARD_IMPORT_UPLOAD",
        object_type="ElectronicCardImportBatch",
        object_id=batch.id,
    ).exists()


@pytest.mark.django_db
def test_hq_upload_accepts_legacy_geunroyeonwol_header(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    user = _build_user("labor-hq-upload-legacy", Role.HQ)
    project = Project.objects.create(code="PRJ-LAB-002", name="레거시 헤더 현장")
    upload = _xlsx_upload(year_month_header="근로연월")

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"year_month": "2026-05", "project": str(project.id)},
            files={"source_file": upload},
        )
    )

    assert response.status_code == 302
    batch = ElectronicCardImportBatch.objects.get(project=project)
    assert "year_month" in batch.header_check_summary["detected_headers"]
    assert batch.header_check_summary["workbook_months"] == ["2026-05"]


@pytest.mark.django_db
@pytest.mark.parametrize("filename", ["ecard.pdf", "ecard.xls"])
def test_non_xlsx_upload_is_rejected(settings, tmp_path, filename):
    settings.MEDIA_ROOT = tmp_path
    user = _build_user(f"labor-hq-{filename}", Role.HQ)
    project = Project.objects.create(code=f"PRJ-{filename}", name="형식 검증 현장")
    upload = SimpleUploadedFile(filename, b"dummy", content_type="application/octet-stream")

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"year_month": "2026-05", "project": str(project.id)},
            files={"source_file": upload},
        )
    )

    assert response.status_code == 200
    assert ElectronicCardImportBatch.objects.count() == 0
    content = response.content.decode("utf-8")
    assert ".xlsx 파일만 업로드" in content


def _create_batch(settings, tmp_path, *, worker_name="홍길동", rrn="900101-1234567", phone="010-1234-5678", day_values=None):
    settings.MEDIA_ROOT = tmp_path
    user = _build_user(f"labor-hq-parse-{worker_name}", Role.HQ)
    project = Project.objects.create(code=f"PRJ-{worker_name}", name="파싱 현장")
    upload = _xlsx_upload(worker_name=worker_name, rrn=rrn, phone=phone, day_values=day_values)
    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"year_month": "2026-05", "project": str(project.id)},
            files={"source_file": upload},
        )
    )
    assert response.status_code == 302
    return user, project, ElectronicCardImportBatch.objects.get(project=project)


def _parse_batch(user, batch):
    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"action": "parse", "batch_id": str(batch.id)},
        )
    )
    assert response.status_code == 302
    batch.refresh_from_db()
    return batch


def _reconcile_batch(user, batch):
    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"action": "reconcile", "batch_id": str(batch.id)},
        )
    )
    batch.refresh_from_db()
    return response


def _ensure_labor_role():
    role, _created = LaborRole.objects.get_or_create(
        code="LAB-001",
        defaults={"name": "보통인부"},
    )
    return role


def _create_ledger(*, worker, project, work_date, work_unit):
    return LaborWorkLedger.objects.create(
        work_date=work_date,
        work_month=work_date.replace(day=1),
        worker=worker,
        actual_project=project,
        report_project=project,
        labor_role=_ensure_labor_role(),
        work_unit=Decimal(str(work_unit)),
        work_hours=Decimal("8.00"),
        unit_wage=100000,
        status="DRAFT",
    )


@pytest.mark.django_db
def test_hq_can_parse_uploaded_batch(settings, tmp_path):
    user, _project, batch = _create_batch(settings, tmp_path, day_values={1: "1", 2: "0.5", 3: ""})

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"action": "parse", "batch_id": str(batch.id)},
        )
    )

    assert response.status_code == 302
    batch.refresh_from_db()
    assert ElectronicCardWorkRaw.objects.filter(batch=batch).count() == 1
    assert ElectronicCardWorkDay.objects.filter(batch=batch).count() == 31
    assert batch.header_check_summary["parsed_raw_count"] == 1
    assert batch.header_check_summary["parsed_day_count"] == 31
    assert AuditLog.objects.filter(
        action="LABOR_ECARD_IMPORT_PARSE",
        object_type="ElectronicCardImportBatch",
        object_id=batch.id,
    ).exists()


@pytest.mark.django_db
def test_field_cannot_parse_batch(settings, tmp_path):
    user, _project, batch = _create_batch(settings, tmp_path)
    field_user = _build_user("labor-field-parse", Role.FIELD)

    with pytest.raises(PermissionDenied):
        hq_e_card_import_batch_list(
            _build_request(
                field_user,
                method="post",
                data={"action": "parse", "batch_id": str(batch.id)},
            )
        )


@pytest.mark.django_db
def test_parse_matches_existing_worker_by_identity_hash(settings, tmp_path):
    rrn = "900101-1234567"
    worker = WorkerMaster(name="홍길동", phone="010-1234-5678", active=True)
    worker.set_rrn(rrn)
    worker.save()
    user, _project, batch = _create_batch(settings, tmp_path, rrn=rrn)

    hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"action": "parse", "batch_id": str(batch.id)},
        )
    )

    raw = ElectronicCardWorkRaw.objects.get(batch=batch)
    day = ElectronicCardWorkDay.objects.filter(batch=batch, work_date__day=1).get()
    assert raw.matched_worker == worker
    assert raw.match_status == ElectronicCardMatchStatus.MATCHED
    assert day.worker == worker
    assert day.match_status == ElectronicCardMatchStatus.MATCHED
    assert WorkerMaster.objects.count() == 1
    batch.refresh_from_db()
    assert batch.header_check_summary["matched_by_identity_count"] == 1
    assert batch.header_check_summary["matched_by_name_phone_count"] == 0
    assert batch.header_check_summary["unmatched_preview"] == []


@pytest.mark.django_db
def test_parse_keeps_unmatched_worker_and_reparse_does_not_duplicate(settings, tmp_path):
    user, _project, batch = _create_batch(settings, tmp_path, worker_name="미등록자", rrn="850101-2345678", phone="010-9999-9999")

    for _ in range(2):
        response = hq_e_card_import_batch_list(
            _build_request(
                user,
                method="post",
                data={"action": "parse", "batch_id": str(batch.id)},
            )
        )
        assert response.status_code == 302

    batch.refresh_from_db()
    assert ElectronicCardWorkRaw.objects.filter(batch=batch).count() == 1
    assert ElectronicCardWorkDay.objects.filter(batch=batch).count() == 31
    raw = ElectronicCardWorkRaw.objects.get(batch=batch)
    assert raw.match_status == ElectronicCardMatchStatus.UNMATCHED
    assert raw.matched_worker is None
    assert batch.header_check_summary["matched_worker_count"] == 0
    assert batch.header_check_summary["unmatched_worker_count"] == 1
    assert batch.header_check_summary["matched_by_identity_count"] == 0
    assert batch.header_check_summary["matched_by_name_phone_count"] == 0
    assert batch.header_check_summary["unmatched_preview"][0]["worker_name_raw"] == "미등록자"
    assert "2345678" not in str(batch.header_check_summary["unmatched_preview"])
    assert WorkerMaster.objects.count() == 0


@pytest.mark.django_db
def test_parse_matches_worker_with_rrn_format_difference(settings, tmp_path):
    rrn = "9001011234567"
    worker = WorkerMaster(name="홍길동", phone="01012345678", active=True)
    worker.set_rrn("900101-1234567")
    worker.save()
    user, _project, batch = _create_batch(settings, tmp_path, rrn=rrn, phone="010-1234-5678")

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"action": "parse", "batch_id": str(batch.id)},
        )
    )

    assert response.status_code == 302
    raw = ElectronicCardWorkRaw.objects.get(batch=batch)
    assert raw.matched_worker == worker
    assert raw.match_status == ElectronicCardMatchStatus.MATCHED


@pytest.mark.django_db
def test_parse_matches_worker_created_via_direct_orm_save(settings, tmp_path):
    worker = WorkerMaster.objects.create(
        name="김철수",
        rrn_encrypted="880101-1234567",
        phone="010-7777-8888",
        active=True,
    )
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="김철수",
        rrn="8801011234567",
        phone="01077778888",
    )

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"action": "parse", "batch_id": str(batch.id)},
        )
    )

    assert response.status_code == 302
    raw = ElectronicCardWorkRaw.objects.get(batch=batch)
    assert raw.matched_worker == worker
    assert raw.match_status == ElectronicCardMatchStatus.MATCHED


@pytest.mark.django_db
def test_parse_matches_worker_by_name_and_phone_normalization(settings, tmp_path):
    worker = WorkerMaster.objects.create(
        name="홍 길 동",
        rrn_encrypted="",
        rrn_masked="",
        identity_hash="",
        phone="01012345678",
        active=True,
    )
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="홍길동",
        rrn="",
        phone="010-1234-5678",
    )

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"action": "parse", "batch_id": str(batch.id)},
        )
    )

    assert response.status_code == 302
    raw = ElectronicCardWorkRaw.objects.get(batch=batch)
    assert raw.matched_worker == worker
    assert raw.match_status == ElectronicCardMatchStatus.MATCHED
    batch.refresh_from_db()
    assert batch.header_check_summary["matched_by_identity_count"] == 0
    assert batch.header_check_summary["matched_by_name_phone_count"] == 1


@pytest.mark.django_db
def test_hq_can_reconcile_parsed_batch(settings, tmp_path):
    rrn = "900101-1234567"
    worker = WorkerMaster(name="홍길동", phone="010-1234-5678", active=True)
    worker.set_rrn(rrn)
    worker.save()
    user, project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="홍길동",
        rrn=rrn,
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="1")

    response = _reconcile_batch(user, batch)

    assert response.status_code == 302
    result = LaborReconciliationResult.objects.get(batch=batch)
    assert result.status == LaborReconciliationStatus.MATCH
    assert result.erp_work_unit == Decimal("1")
    assert result.card_work_unit == Decimal("1")
    assert batch.header_check_summary["reconciliation_match_count"] == 1
    assert AuditLog.objects.filter(
        action="LABOR_ECARD_RECONCILE",
        object_type="ElectronicCardImportBatch",
        object_id=batch.id,
    ).exists()


@pytest.mark.django_db
def test_field_cannot_reconcile_batch(settings, tmp_path):
    user, _project, batch = _create_batch(settings, tmp_path, day_values={1: "1"})
    _parse_batch(user, batch)
    field_user = _build_user("labor-field-reconcile", Role.FIELD)

    with pytest.raises(PermissionDenied):
        hq_e_card_import_batch_list(
            _build_request(
                field_user,
                method="post",
                data={"action": "reconcile", "batch_id": str(batch.id)},
            )
        )


@pytest.mark.django_db
def test_reconcile_creates_diff_card_only_erp_only_and_unmatched_rows(settings, tmp_path):
    worker = WorkerMaster(name="김철수", phone="010-1234-0000", active=True)
    worker.set_rrn("900101-1234567")
    worker.save()
    user, project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="김철수",
        rrn="900101-1234567",
        phone="010-1234-0000",
        day_values={1: "1", 2: "1"},
    )
    _parse_batch(user, batch)

    unmatched_user, _unused_project, unmatched_batch = _create_batch(
        settings,
        tmp_path,
        worker_name="미등록자",
        rrn="850101-2345678",
        phone="010-9999-9999",
        day_values={3: "1"},
    )
    _parse_batch(unmatched_user, unmatched_batch)

    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="0.5")
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 4), work_unit="1")

    response = _reconcile_batch(user, batch)
    unmatched_response = _reconcile_batch(unmatched_user, unmatched_batch)

    assert response.status_code == 302
    assert unmatched_response.status_code == 302

    statuses = list(
        LaborReconciliationResult.objects.filter(batch=batch)
        .order_by("work_date", "id")
        .values_list("status", flat=True)
    )
    assert statuses == [
        LaborReconciliationStatus.DIFF,
        LaborReconciliationStatus.CARD_ONLY,
        LaborReconciliationStatus.ERP_ONLY,
    ]

    unmatched = LaborReconciliationResult.objects.get(batch=unmatched_batch)
    assert unmatched.status == LaborReconciliationStatus.UNMATCHED
    assert unmatched.worker is None
    assert WorkerMaster.objects.filter(name="미등록자").count() == 0


@pytest.mark.django_db
def test_reconcile_twice_does_not_duplicate_rows(settings, tmp_path):
    worker = WorkerMaster(name="박정희", phone="010-5555-6666", active=True)
    worker.set_rrn("910101-1234567")
    worker.save()
    user, project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="박정희",
        rrn="910101-1234567",
        phone="010-5555-6666",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="1")

    first_response = _reconcile_batch(user, batch)
    first_ids = list(LaborReconciliationResult.objects.filter(batch=batch).values_list("id", flat=True))
    second_response = _reconcile_batch(user, batch)
    second_count = LaborReconciliationResult.objects.filter(batch=batch).count()

    assert first_response.status_code == 302
    assert second_response.status_code == 302
    assert len(first_ids) == 1
    assert second_count == 1


@pytest.mark.django_db
def test_header_alias_bidaesangsayu_is_accepted(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    user = _build_user("labor-hq-exclusion", Role.HQ)
    project = Project.objects.create(code="PRJ-EXC-001", name="헤더 별칭 현장")
    upload = _xlsx_upload_with_exclusion_header(exclusion_header="비대상사유", exclusion_value="대상 아님")

    response = hq_e_card_import_batch_list(
        _build_request(
            user,
            method="post",
            data={"year_month": "2026-05", "project": str(project.id)},
            files={"source_file": upload},
        )
    )

    assert response.status_code == 302
    batch = ElectronicCardImportBatch.objects.get(project=project)
    warnings = batch.header_check_summary.get("warnings", [])
    assert not any("비대상사유" in warning and "누락" in warning for warning in warnings)


@pytest.mark.django_db
def test_reconcile_zero_result_reason_when_all_card_days_are_zero_and_no_erp(settings, tmp_path):
    user, _project, batch = _create_batch(settings, tmp_path, day_values={})
    _parse_batch(user, batch)

    response = _reconcile_batch(user, batch)

    assert response.status_code == 302
    batch.refresh_from_db()
    assert batch.header_check_summary["parsed_worked_day_count"] == 0
    assert batch.header_check_summary["reconcile_card_worked_day_count"] == 0
    assert batch.header_check_summary["reconcile_erp_worked_count"] == 0
    assert batch.header_check_summary["reconciliation_total_count"] == 0
    assert batch.header_check_summary["reconciliation_no_result_reason"] == "전자카드 출역값과 ERP 원장이 모두 없습니다."


@pytest.mark.django_db
def test_reconcile_card_only_summary_count_when_card_work_exists_without_erp(settings, tmp_path):
    worker = WorkerMaster(name="이순신", phone="010-2222-3333", active=True)
    worker.set_rrn("920101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="이순신",
        rrn="920101-1234567",
        phone="010-2222-3333",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)

    response = _reconcile_batch(user, batch)

    assert response.status_code == 302
    batch.refresh_from_db()
    assert batch.header_check_summary["parsed_worked_day_count"] == 1
    assert batch.header_check_summary["reconcile_card_worked_day_count"] == 1
    assert batch.header_check_summary["reconciliation_card_only_count"] == 1


@pytest.mark.django_db
def test_reconcile_erp_only_summary_count_when_erp_exists_without_card_work(settings, tmp_path):
    worker = WorkerMaster(name="장보고", phone="010-4444-5555", active=True)
    worker.set_rrn("930101-1234567")
    worker.save()
    user, project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="장보고",
        rrn="930101-1234567",
        phone="010-4444-5555",
        day_values={},
    )
    _parse_batch(user, batch)
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="1")

    response = _reconcile_batch(user, batch)

    assert response.status_code == 302
    batch.refresh_from_db()
    assert batch.header_check_summary["reconcile_erp_ledger_count"] >= 1
    assert batch.header_check_summary["reconcile_erp_worked_count"] == 1
    assert batch.header_check_summary["reconciliation_erp_only_count"] == 1


@pytest.mark.django_db
def test_hq_can_open_reconciliation_detail_page(settings, tmp_path):
    worker = WorkerMaster(name="홍길동", phone="010-1234-5678", active=True)
    worker.set_rrn("900101-1234567")
    worker.save()
    user, project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="홍길동",
        rrn="900101-1234567",
        phone="010-1234-5678",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="1")
    _reconcile_batch(user, batch)

    response = hq_e_card_import_batch_detail(_build_detail_request(user, batch), batch.id)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert "대사 결과" in content
    assert "대사 전체 확정" in content


@pytest.mark.django_db
def test_field_cannot_open_reconciliation_detail_page(settings, tmp_path):
    user, _project, batch = _create_batch(settings, tmp_path, day_values={1: "1"})
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    field_user = _build_user("labor-field-detail", Role.FIELD)

    with pytest.raises(PermissionDenied):
        hq_e_card_import_batch_detail(_build_detail_request(field_user, batch), batch.id)


@pytest.mark.django_db
def test_resolve_card_only_with_card(settings, tmp_path):
    worker = WorkerMaster(name="이순신", phone="010-2222-3333", active=True)
    worker.set_rrn("920101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="이순신",
        rrn="920101-1234567",
        phone="010-2222-3333",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "CARD",
                f"row_comment__{result.id}": "전자카드 기준으로 확정",
                "next_query": "",
            },
        ),
        batch.id,
    )

    assert response.status_code == 302
    result.refresh_from_db()
    assert result.status == "RESOLVED"
    assert result.resolution == "CARD"
    assert result.final_work_unit == result.card_work_unit
    assert result.resolved_by == user


@pytest.mark.django_db
def test_manual_resolution_requires_final_work_unit(settings, tmp_path):
    worker = WorkerMaster(name="김철수", phone="010-3333-4444", active=True)
    worker.set_rrn("930101-1234567")
    worker.save()
    user, project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="김철수",
        rrn="930101-1234567",
        phone="010-3333-4444",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="0.5")
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "MANUAL",
                f"row_comment__{result.id}": "수동 조정 필요",
                "next_query": "",
            },
        ),
        batch.id,
    )

    assert response.status_code == 200
    result.refresh_from_db()
    assert result.resolution == ""


@pytest.mark.django_db
def test_exclude_requires_comment(settings, tmp_path):
    worker = WorkerMaster(name="장보고", phone="010-5555-6666", active=True)
    worker.set_rrn("940101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="장보고",
        rrn="940101-1234567",
        phone="010-5555-6666",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "EXCLUDED",
                "next_query": "",
            },
        ),
        batch.id,
    )

    assert response.status_code == 200
    result.refresh_from_db()
    assert result.resolution == ""


@pytest.mark.django_db
def test_match_worker_updates_raw_and_day_rows_without_creating_worker(settings, tmp_path):
    existing_worker = WorkerMaster(name="박영희", phone="010-7777-8888", active=True)
    existing_worker.set_rrn("950101-1234567")
    existing_worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="박영희",
        rrn="",
        phone="010-0000-0000",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "match_result_id": str(result.id),
                f"row_worker_id__{result.id}": str(existing_worker.id),
                "next_query": "",
            },
        ),
        batch.id,
    )

    assert response.status_code == 302
    raw = ElectronicCardWorkRaw.objects.get(batch=batch)
    assert raw.matched_worker == existing_worker
    assert ElectronicCardWorkDay.objects.filter(batch=batch, worker=existing_worker).exists()
    assert WorkerMaster.objects.filter(name="박영희").count() == 1


@pytest.mark.django_db
def test_cannot_confirm_batch_with_unresolved_rows(settings, tmp_path):
    worker = WorkerMaster(name="최무선", phone="010-8888-9999", active=True)
    worker.set_rrn("960101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="최무선",
        rrn="960101-1234567",
        phone="010-8888-9999",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={"action": "confirm_batch", "next_query": ""},
        ),
        batch.id,
    )

    assert response.status_code == 200
    batch.refresh_from_db()
    assert batch.status != "CONFIRMED"


@pytest.mark.django_db
def test_can_confirm_batch_after_resolving_non_match_rows(settings, tmp_path):
    worker = WorkerMaster(name="정약용", phone="010-0000-1111", active=True)
    worker.set_rrn("970101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="정약용",
        rrn="970101-1234567",
        phone="010-0000-1111",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)

    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "CARD",
                f"row_comment__{result.id}": "전자카드 기준 확정",
                "next_query": "",
            },
        ),
        batch.id,
    )
    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={"action": "confirm_batch", "next_query": ""},
        ),
        batch.id,
    )

    assert response.status_code == 302
    batch.refresh_from_db()
    assert batch.status == "CONFIRMED"
    assert batch.confirmed_by == user


@pytest.mark.django_db
def test_confirm_match_creates_confirmed_work_day_and_audit_log(settings, tmp_path):
    worker = WorkerMaster(name="확정일치", phone="010-1212-3434", active=True)
    worker.set_rrn("980101-1234567")
    worker.save()
    user, project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="확정일치",
        rrn="980101-1234567",
        phone="010-1212-3434",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="1")
    _reconcile_batch(user, batch)

    raw_count = ElectronicCardWorkRaw.objects.filter(batch=batch).count()
    day_count = ElectronicCardWorkDay.objects.filter(batch=batch).count()
    ledger_count = LaborWorkLedger.objects.filter(report_project=project).count()

    response = hq_e_card_import_batch_detail(
        _build_detail_request(user, batch, method="post", data={"action": "confirm_batch", "next_query": ""}),
        batch.id,
    )

    assert response.status_code == 302
    confirmed = LaborConfirmedWorkDay.objects.get(batch=batch)
    batch.refresh_from_db()
    assert batch.status == "CONFIRMED"
    assert confirmed.source_basis == "ERP"
    assert confirmed.final_work_unit == Decimal("1")
    assert confirmed.export_included is True
    assert confirmed.report_project == project
    assert confirmed.worker == worker
    assert ElectronicCardWorkRaw.objects.filter(batch=batch).count() == raw_count
    assert ElectronicCardWorkDay.objects.filter(batch=batch).count() == day_count
    assert LaborWorkLedger.objects.filter(report_project=project).count() == ledger_count
    assert AuditLog.objects.filter(action="LABOR_CONFIRMED_WORKDAY_GENERATE", object_id=batch.id).exists()


@pytest.mark.django_db
def test_confirm_uses_card_and_manual_and_excluded_source_basis(settings, tmp_path):
    worker = WorkerMaster(name="확정혼합", phone="010-7777-2222", active=True)
    worker.set_rrn("981231-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="확정혼합",
        rrn="981231-1234567",
        phone="010-7777-2222",
        day_values={1: "1", 2: "1", 3: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    results = list(LaborReconciliationResult.objects.filter(batch=batch).order_by("work_date", "id"))

    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(results[0].id),
                f"row_resolution__{results[0].id}": "CARD",
                f"row_comment__{results[0].id}": "카드 기준",
                "next_query": "",
            },
        ),
        batch.id,
    )
    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(results[1].id),
                f"row_resolution__{results[1].id}": "MANUAL",
                f"row_comment__{results[1].id}": "수동 조정",
                f"row_final_work_unit__{results[1].id}": "0.75",
                "next_query": "",
            },
        ),
        batch.id,
    )
    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(results[2].id),
                f"row_resolution__{results[2].id}": "EXCLUDED",
                f"row_comment__{results[2].id}": "제외 처리",
                "next_query": "",
            },
        ),
        batch.id,
    )

    response = hq_e_card_import_batch_detail(
        _build_detail_request(user, batch, method="post", data={"action": "confirm_batch", "next_query": ""}),
        batch.id,
    )

    assert response.status_code == 302
    rows = list(LaborConfirmedWorkDay.objects.filter(batch=batch).order_by("work_date", "id"))
    assert [row.source_basis for row in rows] == ["CARD", "MANUAL", "EXCLUDED"]
    assert rows[0].final_work_unit == Decimal("1")
    assert rows[1].final_work_unit == Decimal("0.75")
    assert rows[2].final_work_unit == Decimal("0")
    assert rows[2].export_included is False


@pytest.mark.django_db
def test_repeated_confirm_does_not_duplicate_confirmed_rows(settings, tmp_path):
    worker = WorkerMaster(name="재확정", phone="010-3131-4242", active=True)
    worker.set_rrn("990101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="재확정",
        rrn="990101-1234567",
        phone="010-3131-4242",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)
    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "CARD",
                f"row_comment__{result.id}": "카드 기준 확정",
                "next_query": "",
            },
        ),
        batch.id,
    )

    first = hq_e_card_import_batch_detail(
        _build_detail_request(user, batch, method="post", data={"action": "confirm_batch", "next_query": ""}),
        batch.id,
    )
    second = hq_e_card_import_batch_detail(
        _build_detail_request(user, batch, method="post", data={"action": "confirm_batch", "next_query": ""}),
        batch.id,
    )

    assert first.status_code == 302
    assert second.status_code == 200
    assert LaborConfirmedWorkDay.objects.filter(batch=batch).count() == 1


@pytest.mark.django_db
def test_confirmed_batch_blocks_followup_processing(settings, tmp_path):
    worker = WorkerMaster(name="후속차단", phone="010-5656-7878", active=True)
    worker.set_rrn("991231-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="후속차단",
        rrn="991231-1234567",
        phone="010-5656-7878",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)
    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "CARD",
                f"row_comment__{result.id}": "카드 기준 확정",
                "next_query": "",
            },
        ),
        batch.id,
    )
    hq_e_card_import_batch_detail(
        _build_detail_request(user, batch, method="post", data={"action": "confirm_batch", "next_query": ""}),
        batch.id,
    )

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "ERP",
                f"row_comment__{result.id}": "재처리 시도",
                "next_query": "",
            },
        ),
        batch.id,
    )

    assert response.status_code == 200
    assert LaborConfirmedWorkDay.objects.filter(batch=batch).count() == 1


@pytest.mark.django_db
def test_hq_can_open_confirmed_work_day_list_and_field_cannot(settings, tmp_path):
    worker = WorkerMaster(name="목록열람", phone="010-8181-9191", active=True)
    worker.set_rrn("000101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="목록열람",
        rrn="000101-1234567",
        phone="010-8181-9191",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    result = LaborReconciliationResult.objects.get(batch=batch)
    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={
                "resolve_result_id": str(result.id),
                f"row_resolution__{result.id}": "CARD",
                f"row_comment__{result.id}": "카드 기준 확정",
                "next_query": "",
            },
        ),
        batch.id,
    )
    hq_e_card_import_batch_detail(
        _build_detail_request(user, batch, method="post", data={"action": "confirm_batch", "next_query": ""}),
        batch.id,
    )

    response = hq_confirmed_work_day_list(_build_request(user, data={"batch_id": str(batch.id)}))
    assert response.status_code == 200
    assert "확정 근로내역" in response.content.decode("utf-8")

    field_user = _build_user("labor-field-confirmed", Role.FIELD)
    with pytest.raises(PermissionDenied):
        hq_confirmed_work_day_list(_build_request(field_user, data={"batch_id": str(batch.id)}))


def _confirm_batch_with_card_resolution(user, batch):
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)
    for result in LaborReconciliationResult.objects.filter(batch=batch).order_by("work_date", "id"):
        hq_e_card_import_batch_detail(
            _build_detail_request(
                user,
                batch,
                method="post",
                data={
                    "resolve_result_id": str(result.id),
                    f"row_resolution__{result.id}": "CARD",
                    f"row_comment__{result.id}": "전자카드 기준 확정",
                    "next_query": "",
                },
            ),
            batch.id,
        )
    response = hq_e_card_import_batch_detail(
        _build_detail_request(user, batch, method="post", data={"action": "confirm_batch", "next_query": ""}),
        batch.id,
    )
    assert response.status_code == 302
    batch.refresh_from_db()
    return batch


def _read_generated_row(batch_or_export_file):
    batch_or_export_file.open("rb")
    workbook = load_workbook(batch_or_export_file, data_only=False)
    sheet = workbook.active
    headers = [sheet.cell(row=1, column=idx).value for idx in range(1, sheet.max_column + 1)]
    values = [sheet.cell(row=2, column=idx).value for idx in range(1, sheet.max_column + 1)]
    row_map = dict(zip(headers, values))
    return workbook, row_map


@pytest.mark.django_db
def test_cannot_generate_cwma_export_before_batch_confirmed(settings, tmp_path):
    worker = WorkerMaster(name="생성전차단", phone="010-1010-1010", active=True)
    worker.set_rrn("990101-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="생성전차단",
        rrn="990101-1234567",
        phone="010-1010-1010",
        day_values={1: "1"},
    )
    _parse_batch(user, batch)
    _reconcile_batch(user, batch)

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={"action": "generate_cwma_reupload", "export_note": "사전 생성 시도"},
        ),
        batch.id,
    )

    assert response.status_code == 200
    assert LaborExcelExportBatch.objects.count() == 0


@pytest.mark.django_db
def test_generate_cwma_reupload_export_creates_file_and_preserves_original(settings, tmp_path):
    worker = WorkerMaster(name="재업로드대상", phone="010-2222-3333", active=True)
    worker.set_rrn("990102-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="재업로드대상",
        rrn="990102-1234567",
        phone="010-2222-3333",
        day_values={1: "1", 2: "1"},
    )
    _confirm_batch_with_card_resolution(user, batch)

    excluded = LaborConfirmedWorkDay.objects.filter(batch=batch, work_date=date(2026, 5, 2)).get()
    excluded.source_basis = "EXCLUDED"
    excluded.final_work_unit = Decimal("0")
    excluded.export_included = False
    excluded.export_value = Decimal("0")
    excluded.export_note = "신고 제외"
    excluded.save(update_fields=["source_basis", "final_work_unit", "export_included", "export_value", "export_note", "updated_at"])

    response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={"action": "generate_cwma_reupload", "export_note": "1차 생성"},
        ),
        batch.id,
    )

    assert response.status_code == 302
    export_batch = LaborExcelExportBatch.objects.get(source_batch=batch)
    assert export_batch.generated_file.name
    assert export_batch.changed_count > 0
    assert export_batch.included_count == 1
    assert export_batch.excluded_count == 1
    assert AuditLog.objects.filter(action="LABOR_EXCEL_EXPORT_GENERATE", object_id=export_batch.id).exists()

    _generated_wb, generated_row = _read_generated_row(export_batch.generated_file)
    _source_wb, source_row = _read_generated_row(batch.source_file)
    assert generated_row["1일"] == 1
    assert generated_row["2일"] in (None, "")
    assert generated_row["신고일수"] == 1
    assert generated_row["확정일수"] == 1
    assert source_row["2일"] == "1"

    second_response = hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={"action": "generate_cwma_reupload", "export_note": "2차 생성"},
        ),
        batch.id,
    )

    assert second_response.status_code == 302
    exports = list(LaborExcelExportBatch.objects.filter(source_batch=batch).order_by("id"))
    assert len(exports) == 2
    assert exports[0].generated_file.name != exports[1].generated_file.name


@pytest.mark.django_db
def test_hq_can_download_export_and_field_is_blocked(settings, tmp_path):
    worker = WorkerMaster(name="다운로드대상", phone="010-4444-5555", active=True)
    worker.set_rrn("990103-1234567")
    worker.save()
    user, _project, batch = _create_batch(
        settings,
        tmp_path,
        worker_name="다운로드대상",
        rrn="990103-1234567",
        phone="010-4444-5555",
        day_values={1: "1"},
    )
    _confirm_batch_with_card_resolution(user, batch)
    hq_e_card_import_batch_detail(
        _build_detail_request(
            user,
            batch,
            method="post",
            data={"action": "generate_cwma_reupload", "export_note": "다운로드 테스트"},
        ),
        batch.id,
    )
    export_batch = LaborExcelExportBatch.objects.get(source_batch=batch)

    response = hq_labor_excel_export_download(_build_request(user), export_batch.id)

    assert response.status_code == 200
    export_batch.refresh_from_db()
    assert export_batch.status == "DOWNLOADED"
    assert export_batch.downloaded_by == user
    assert AuditLog.objects.filter(action="LABOR_EXCEL_EXPORT_DOWNLOAD", object_id=export_batch.id).exists()

    field_user = _build_user("labor-field-export", Role.FIELD)
    with pytest.raises(PermissionDenied):
        hq_labor_excel_export_download(_build_request(field_user), export_batch.id)


@pytest.mark.django_db
def test_hq_can_delete_unused_worker_master():
    user = _build_user("labor-hq-delete", Role.HQ)
    worker = WorkerMaster(name="삭제대상", phone="010-1111-9999", active=True)
    worker.set_rrn("900101-1234567")
    worker.save()

    response = hq_worker_master_delete(_build_request(user, method="post"), worker.id)

    assert response.status_code == 302
    assert not WorkerMaster.objects.filter(id=worker.id).exists()
    log = AuditLog.objects.get(action="LABOR_WORKER_DELETE", object_id=worker.id)
    assert "900101-1234567" not in str(log.meta_json)


@pytest.mark.django_db
def test_hq_deactivates_used_worker_master_instead_of_deleting():
    user = _build_user("labor-hq-deactivate", Role.HQ)
    project = Project.objects.create(code="PRJ-WORKER-DEL", name="근로자 사용현장")
    worker = WorkerMaster(name="비활성대상", phone="010-2222-9999", active=True)
    worker.set_rrn("910101-1234567")
    worker.save()
    _create_ledger(worker=worker, project=project, work_date=date(2026, 5, 1), work_unit="1")

    response = hq_worker_master_delete(_build_request(user, method="post"), worker.id)

    assert response.status_code == 302
    worker.refresh_from_db()
    assert worker.active is False
    log = AuditLog.objects.get(action="LABOR_WORKER_DEACTIVATE", object_id=worker.id)
    assert log.meta_json["usage_count"] >= 1
    assert "910101-1234567" not in str(log.meta_json)


@pytest.mark.django_db
def test_field_cannot_delete_worker_master():
    user = _build_user("labor-field-delete", Role.FIELD)
    worker = WorkerMaster(name="권한차단", phone="010-3333-9999", active=True)
    worker.set_rrn("920101-1234567")
    worker.save()

    with pytest.raises(PermissionDenied):
        hq_worker_master_delete(_build_request(user, method="post"), worker.id)


@pytest.mark.django_db
def test_get_worker_delete_is_blocked():
    user = _build_user("labor-hq-delete-get", Role.HQ)
    worker = WorkerMaster(name="GET차단", phone="010-4444-9999", active=True)
    worker.set_rrn("930101-1234567")
    worker.save()

    with pytest.raises(PermissionDenied):
        hq_worker_master_delete(_build_request(user), worker.id)
