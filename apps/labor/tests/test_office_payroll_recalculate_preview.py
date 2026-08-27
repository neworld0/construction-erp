from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from apps.core.rbac.models import Role, UserProfile
from apps.labor.models import (
    OfficeEmployeeProfile,
    OfficePayrollCorrection,
    OfficePayrollDeductionPolicy,
    OfficePayrollRun,
    OfficePayrollStatus,
    OfficePayslip,
    IncomeTaxTableRow,
    IncomeTaxTableVersion,
)
from apps.labor.office_payroll_exports import DEDUCTION_FIELDS


@pytest.fixture(autouse=True)
def _disable_two_factor_enforce(settings):
    settings.MIDDLEWARE = [
        middleware
        for middleware in settings.MIDDLEWARE
        if middleware != "apps.core.middleware.two_factor_enforce.TwoFactorEnforceMiddleware"
    ]


@pytest.fixture(autouse=True)
def _official_income_tax_table(db):
    version = IncomeTaxTableVersion.objects.create(effective_from="2026-03-01", source_name="official.xlsx", is_active=True)
    for family_count in range(1, 12):
        IncomeTaxTableRow.objects.create(version=version, monthly_pay_from=0, monthly_pay_to=10_000_000, dependent_count=family_count, income_tax=100_000)


@pytest.mark.django_db
def test_recalculate_previews_posted_earnings_without_saving(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="payroll-preview-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="payroll-preview-employee", password="pass")
    employee = OfficeEmployeeProfile.objects.create(user=employee_user, employee_no="ASAN-HQ-2026-0001")
    run = OfficePayrollRun.objects.create(period_year=2026, period_month=8, created_by=hq)
    slip = OfficePayslip.objects.create(run=run, employee=employee)

    client.force_login(hq)
    response = client.post(
        f"/app/hq/labor/office-payroll/{run.id}/",
        {
            "action": "recalculate",
            f"slip-{slip.id}-base_pay": "5000000",
            f"slip-{slip.id}-meal_allowance_pay": "100000",
            f"slip-{slip.id}-fuel_allowance_pay": "0",
            f"slip-{slip.id}-site_allowance_pay": "100000",
            f"slip-{slip.id}-overtime_pay": "0",
            f"slip-{slip.id}-bonus_pay": "100000",
            f"slip-{slip.id}-other_deduction": "0",
        },
    )

    assert response.status_code == 200
    preview = list(response.context["slips"])[0]
    assert preview.base_pay == 5_000_000
    assert preview.meal_allowance_pay == 100_000
    assert preview.site_allowance_pay == 100_000
    assert preview.gross_pay == 5_300_000
    assert preview.national_pension == 247_000
    slip.refresh_from_db()
    assert slip.base_pay == 0


@pytest.mark.django_db
def test_hq_can_withdraw_a_submitted_payroll_run(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="payroll-withdraw-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    run = OfficePayrollRun.objects.create(
        period_year=2026,
        period_month=9,
        created_by=hq,
        status="SUBMITTED",
    )

    client.force_login(hq)
    response = client.post(
        f"/app/hq/labor/office-payroll/{run.id}/",
        {"action": "withdraw_submission"},
    )

    assert response.status_code == 302
    run.refresh_from_db()
    assert run.status == "DRAFT"
    assert run.submitted_at is None


@pytest.mark.django_db
def test_hq_can_void_an_unallocated_approved_run_and_recreate_the_month(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="payroll-void-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    run = OfficePayrollRun.objects.create(
        period_year=2026,
        period_month=2,
        created_by=hq,
        status=OfficePayrollStatus.APPROVED,
    )

    client.force_login(hq)
    response = client.post(
        f"/app/hq/labor/office-payroll/{run.id}/",
        {"action": "void", "void_reason": "테스트 급여대장 폐기"},
    )

    assert response.status_code == 302
    assert response.url == "/app/hq/labor/office-payroll/"
    run.refresh_from_db()
    assert run.status == OfficePayrollStatus.VOID
    assert run.voided_by == hq
    assert "테스트 급여대장 폐기" in run.note
    replacement = OfficePayrollRun.objects.create(period_year=2026, period_month=2, created_by=hq)
    assert replacement.status == OfficePayrollStatus.DRAFT


@pytest.mark.django_db
def test_hq_can_save_annual_automatic_deduction_policy(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="payroll-policy-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    client.force_login(hq)

    response = client.post(
        "/app/hq/labor/office-payroll/deduction-policies/",
        {
            "year": "2027",
            "national_pension_rate": "4.5",
            "health_insurance_rate": "3.595",
            "long_term_care_rate": "13.14",
            "employment_insurance_rate": "0.9",
            "income_tax_rate": "3",
        },
    )

    assert response.status_code == 302
    policy = OfficePayrollDeductionPolicy.objects.get(year=2027)
    assert policy.national_pension_rate == Decimal("0.045")
    assert policy.long_term_care_rate == Decimal("0.1314")


@pytest.mark.django_db
def test_hq_can_save_active_table_child_tax_credit_amounts(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="child-credit-policy-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    version = IncomeTaxTableVersion.objects.get(is_active=True)
    client.force_login(hq)

    response = client.post(
        "/app/hq/labor/office-payroll/deduction-policies/?year=2026",
        {
            "action": "save_child_tax_credit",
            "income_tax_table_version_id": str(version.id),
            "child_credit_one": "13000",
            "child_credit_two": "30000",
            "child_credit_per_additional": "26000",
        },
    )

    assert response.status_code == 302
    version.refresh_from_db()
    assert version.child_credit_one == 13_000
    assert version.child_credit_two == 30_000
    assert version.child_credit_per_additional == 26_000


@pytest.mark.django_db
def test_hq_payroll_list_renders_child_tax_inputs(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="child-credit-list-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="child-credit-list-employee", password="pass")
    OfficeEmployeeProfile.objects.create(
        user=employee_user,
        employee_no="ASAN-HQ-2026-0301",
        tax_child_count_8_to_20=2,
    )
    client.force_login(hq)

    response = client.get("/app/hq/labor/office-payroll/")

    assert response.status_code == 200
    assert "8~20세 자녀" in response.content.decode()


@pytest.mark.django_db
def test_finalized_payroll_changes_only_through_separate_correction_approval(client):
    user_model = get_user_model()
    requester = user_model.objects.create_user(username="payroll-correction-requester", password="pass")
    reviewer = user_model.objects.create_user(username="payroll-correction-reviewer", password="pass")
    UserProfile.objects.create(user=requester, role=Role.HQ)
    UserProfile.objects.create(user=reviewer, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="payroll-correction-employee", password="pass")
    employee = OfficeEmployeeProfile.objects.create(user=employee_user, employee_no="ASAN-HQ-2026-0101")
    run = OfficePayrollRun.objects.create(period_year=2026, period_month=10, created_by=requester, status="APPROVED")
    slip = OfficePayslip.objects.create(run=run, employee=employee, base_pay=3_000_000)

    client.force_login(requester)
    assert client.get(f"/app/hq/labor/office-payroll/{run.id}/corrections/new/").status_code == 302
    correction = OfficePayrollCorrection.objects.get(run=run)
    fields = ("base_pay", "meal_allowance_pay", "fuel_allowance_pay", "site_allowance_pay", "overtime_pay", "bonus_pay", "income_tax", "local_income_tax", "national_pension", "health_insurance", "long_term_care", "employment_insurance", "other_deduction")
    post_data = {"action": "submit", "reason": "기본급 산정 오류 정정"}
    for field in fields:
        post_data[f"line-{slip.id}-{field}"] = "3200000" if field == "base_pay" else "0"
    assert client.post(f"/app/hq/labor/office-payroll/corrections/{correction.id}/", post_data).status_code == 302
    correction.refresh_from_db()
    assert correction.status == "SUBMITTED"
    slip.refresh_from_db()
    assert slip.base_pay == 3_000_000

    client.force_login(reviewer)
    assert client.post(f"/app/hq/labor/office-payroll/corrections/{correction.id}/", {"action": "approve"}).status_code == 302
    correction.refresh_from_db()
    assert correction.status == "APPROVED"
    assert client.post(f"/app/hq/labor/office-payroll/corrections/{correction.id}/", {"action": "apply"}).status_code == 302
    correction.refresh_from_db()
    slip.refresh_from_db()
    assert correction.status == "APPLIED"
    assert slip.base_pay == 3_200_000


@pytest.mark.django_db
def test_hq_can_download_finalized_payroll_and_payslip_as_excel_and_pdf(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="payroll-export-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="payroll-export-employee", password="pass")
    employee = OfficeEmployeeProfile.objects.create(user=employee_user, employee_no="ASAN-HQ-2026-0201", department="관리부")
    run = OfficePayrollRun.objects.create(period_year=2026, period_month=11, created_by=hq, status="APPROVED")
    slip = OfficePayslip.objects.create(run=run, employee=employee, base_pay=3_000_000, national_pension=135_000)
    client.force_login(hq)

    xlsx_response = client.get(f"/app/hq/labor/office-payroll/{run.id}/downloads/payroll.xlsx")
    assert xlsx_response.status_code == 200
    book = load_workbook(BytesIO(b"".join(xlsx_response.streaming_content)))
    assert book["급여대장"]["A1"].value == "2026년 11월 본사 급여대장"
    payroll_pdf_response = client.get(f"/app/hq/labor/office-payroll/{run.id}/downloads/payroll.pdf")
    assert payroll_pdf_response.status_code == 200
    assert b"".join(payroll_pdf_response.streaming_content).startswith(b"%PDF")
    payslip_xlsx_response = client.get(f"/app/hq/labor/office-payroll/{run.id}/payslips/{slip.id}/downloads/payslip.xlsx")
    assert payslip_xlsx_response.status_code == 200
    payslip_book = load_workbook(BytesIO(b"".join(payslip_xlsx_response.streaming_content)))
    assert payslip_book["급여명세서"]["A1"].value == "2026년 11월 급여명세서"
    assert payslip_book["급여명세서"]["A7"].value == "지급일자"
    assert payslip_book["급여명세서"]["B7"].value == "2026년 12월 05일"
    assert [payslip_book["급여명세서"].cell(row, 3).value for row in range(10, 17)] == [
        label for _, label in DEDUCTION_FIELDS
    ]
    assert "연장수당" not in [payslip_book["급여명세서"].cell(row, 3).value for row in range(10, 17)]
    assert "상여" not in [payslip_book["급여명세서"].cell(row, 3).value for row in range(10, 17)]
    pdf_response = client.get(f"/app/hq/labor/office-payroll/{run.id}/payslips/{slip.id}/downloads/payslip.pdf")
    assert pdf_response.status_code == 200
    assert b"".join(pdf_response.streaming_content).startswith(b"%PDF")


@pytest.mark.django_db
def test_six_standard_payment_items_are_saved_and_shown(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="six-item-payroll-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="six-item-payroll-employee", password="pass")
    employee = OfficeEmployeeProfile.objects.create(user=employee_user, employee_no="ASAN-HQ-2026-0601")
    run = OfficePayrollRun.objects.create(period_year=2026, period_month=6, created_by=hq)
    slip = OfficePayslip.objects.create(
        run=run,
        employee=employee,
        base_pay=3_000_000,
        meal_allowance_pay=100_000,
        fuel_allowance_pay=50_000,
        site_allowance_pay=150_000,
        overtime_pay=200_000,
        bonus_pay=500_000,
    )
    assert slip.gross_pay == 4_000_000

    client.force_login(hq)
    response = client.get(f"/app/hq/labor/office-payroll/{run.id}/")
    content = response.content.decode()
    assert response.status_code == 200
    for label in ("기본급", "식대", "주유대", "현장수당", "연장수당", "상여"):
        assert label in content


@pytest.mark.django_db
def test_payroll_payment_date_is_the_fifth_of_next_month():
    user = get_user_model().objects.create_user(username="payment-date-run-user")
    july = OfficePayrollRun.objects.create(period_year=2026, period_month=7, created_by=user)
    december = OfficePayrollRun.objects.create(period_year=2026, period_month=12, created_by=user)

    assert july.payment_date.isoformat() == "2026-08-05"
    assert december.payment_date.isoformat() == "2027-01-05"


@pytest.mark.django_db
def test_new_payroll_run_seeds_all_previous_month_payment_and_deduction_values(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="recurring-seed-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="recurring-seed-employee", password="pass")
    employee = OfficeEmployeeProfile.objects.create(user=employee_user, employee_no="ASAN-HQ-2026-0701")
    prior = OfficePayrollRun.objects.create(period_year=2026, period_month=6, created_by=hq)
    OfficePayslip.objects.create(
        run=prior,
        employee=employee,
        base_pay=3_100_000,
        meal_allowance_pay=200_000,
        fuel_allowance_pay=100_000,
        site_allowance_pay=150_000,
        overtime_pay=75_000,
        bonus_pay=50_000,
        income_tax=81_000,
        local_income_tax=8_100,
        national_pension=147_250,
        health_insurance=111_440,
        long_term_care=14_640,
        employment_insurance=27_900,
        other_deduction=3_000,
    )
    client.force_login(hq)

    response = client.post(
        "/app/hq/labor/office-payroll/",
        {"period_year": "2026", "period_month": "7"},
    )

    assert response.status_code == 302
    july = OfficePayrollRun.objects.get(period_year=2026, period_month=7)
    seeded = OfficePayslip.objects.get(run=july, employee=employee)
    assert (
        seeded.base_pay, seeded.meal_allowance_pay, seeded.fuel_allowance_pay,
        seeded.site_allowance_pay, seeded.overtime_pay, seeded.bonus_pay,
        seeded.income_tax, seeded.local_income_tax, seeded.national_pension,
        seeded.health_insurance, seeded.long_term_care,
        seeded.employment_insurance, seeded.other_deduction,
    ) == (
        3_100_000, 200_000, 100_000, 150_000, 75_000, 50_000,
        81_000, 8_100, 147_250, 111_440, 14_640, 27_900, 3_000,
    )


@pytest.mark.django_db
def test_payroll_detail_keeps_link_to_submitted_or_approved_correction(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="correction-link-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    run = OfficePayrollRun.objects.create(
        period_year=2026,
        period_month=3,
        status="APPROVED",
        created_by=hq,
    )
    correction = OfficePayrollCorrection.objects.create(
        run=run,
        requested_by=hq,
        status="APPROVED",
        reason="국민연금 정정",
    )

    client.force_login(hq)
    response = client.get(f"/app/hq/labor/office-payroll/{run.id}/")

    assert response.status_code == 200
    assert f"corrections/{correction.id}/".encode() in response.content
    assert "승인 정정 적용".encode() in response.content


@pytest.mark.django_db
def test_requester_can_delete_only_draft_office_payroll_correction(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="correction-delete-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    run = OfficePayrollRun.objects.create(period_year=2026, period_month=8, status="APPROVED", created_by=hq)
    correction = OfficePayrollCorrection.objects.create(run=run, requested_by=hq, status="DRAFT")

    client.force_login(hq)
    response = client.post(
        f"/app/hq/labor/office-payroll/corrections/{correction.id}/",
        {"action": "delete"},
    )

    assert response.status_code == 302
    assert not OfficePayrollCorrection.objects.filter(id=correction.id).exists()


@pytest.mark.django_db
def test_apply_legacy_correction_snapshot_maps_allowance_to_site_allowance(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="legacy-correction-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="legacy-correction-employee", password="pass")
    employee = OfficeEmployeeProfile.objects.create(user=employee_user, employee_no="ASAN-HQ-LEGACY")
    run = OfficePayrollRun.objects.create(period_year=2026, period_month=8, status="APPROVED", created_by=hq)
    slip = OfficePayslip.objects.create(run=run, employee=employee, base_pay=3_000_000, site_allowance_pay=300_000)
    correction = OfficePayrollCorrection.objects.create(
        run=run,
        requested_by=hq,
        status="APPROVED",
        proposed_snapshot=[
            {
                "slip_id": slip.id,
                "employee_no": employee.employee_no,
                "employee_name": "Legacy Employee",
                "base_pay": 3_000_000,
                "allowance_pay": 300_000,
                "overtime_pay": 0,
                "bonus_pay": 0,
                "income_tax": 0,
                "local_income_tax": 0,
                "national_pension": 0,
                "health_insurance": 0,
                "long_term_care": 0,
                "employment_insurance": 0,
                "other_deduction": 0,
            }
        ],
    )

    client.force_login(hq)
    response = client.post(
        f"/app/hq/labor/office-payroll/corrections/{correction.id}/",
        {"action": "apply"},
    )

    assert response.status_code == 302
    slip.refresh_from_db()
    correction.refresh_from_db()
    assert slip.site_allowance_pay == 300_000
    assert slip.meal_allowance_pay == 200_000
    assert correction.status == "APPLIED"


@pytest.mark.django_db
def test_payroll_pdf_supports_long_employee_number_without_overflow(client):
    user_model = get_user_model()
    hq = user_model.objects.create_user(username="long-number-payroll-hq", password="pass")
    UserProfile.objects.create(user=hq, role=Role.HQ)
    employee_user = user_model.objects.create_user(username="long-number-payroll-employee", password="pass")
    employee = OfficeEmployeeProfile.objects.create(
        user=employee_user,
        employee_no="ASAN-HQ-2026-VERY-LONG-EMPLOYEE-NUMBER-0001",
    )
    run = OfficePayrollRun.objects.create(period_year=2026, period_month=7, created_by=hq, status="APPROVED")
    OfficePayslip.objects.create(run=run, employee=employee, base_pay=3_000_000)
    client.force_login(hq)

    response = client.get(f"/app/hq/labor/office-payroll/{run.id}/downloads/payroll.pdf")

    assert response.status_code == 200
    assert b"".join(response.streaming_content).startswith(b"%PDF")
