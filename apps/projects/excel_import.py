from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import re

from django.core.exceptions import ValidationError
from openpyxl import load_workbook


BUDGET_SHEET_PREFERENCES = ["\ub3c4\uae09\uacc4\uc57d\ub0b4\uc5ed\uc11c", "\uc124\uacc4\ub0b4\uc5ed\uc11c"]
BUDGET_SUMMARY_SHEET_PREFERENCES = ["\ub0b4\uc5ed\uc11c\ucd1d\uad04\ud45c(\ub3c4\uae09)"]
WBS_SHEET_PREFERENCES = ["\uc608\uc815\uacf5\uc815\ud45c"]
COMMENCEMENT_INFO_SHEET_PREFERENCES = ["\ucc29\uacf5\uc2e0\uace0\uc11c", "\ucc29\uacf5\uacc4"]

BUDGET_HEADER_ALIASES = {
    "code": ["\ucf54\ub4dc", "\uacf5\uc885", "\ubc88\ud638", "\ub0b4\uc5ed\ubc88\ud638", "\ube44\ubaa9"],
    "hierarchy": ["\uacc4\uce35", "\ub808\ubca8", "\uad6c\ubd84"],
    "item_name": ["\ud488\uba85", "\uc138\ubd80\uacf5\uc885", "\uacf5\uc885\uba85", "\ud56d\ubaa9\uba85", "\uba85\uce6d", "\uc791\uc5c5\uba85"],
    "spec": ["\uaddc\uaca9", "\uc0ac\uc591"],
    "quantity": ["\uc218\ub7c9"],
    "unit": ["\ub2e8\uc704"],
    "unit_price": ["\ub2e8\uac00", "\uacc4\uc57d\ub2e8\uac00"],
    "amount": ["\uae08\uc561", "\uacc4\uc57d\uae08\uc561", "\ud569\uacc4\uae08\uc561", "\ub3c4\uae09\uae08\uc561"],
    "labor_amount": ["\ub178\ubb34\ube44"],
    "material_amount": ["\uc7ac\ub8cc\ube44"],
    "expense_amount": ["\uacbd\ube44"],
}

WBS_HEADER_ALIASES = {
    "code": ["\ucf54\ub4dc", "\ubc88\ud638", "\uacf5\uc885\ucf54\ub4dc", "wbs"],
    "name": ["\uacf5\uc885", "\uc791\uc5c5\uba85", "\uc5c5\ubb34\uba85", "\uc138\ubd80\uacf5\uc815", "\ud56d\ubaa9\uba85"],
    "weight": ["\uac00\uc911\uce58", "\ube44\uc911", "weight", "\ubcf4\ud560"],
    "plan_start_date": ["\ucc29\uc218", "\uc2dc\uc791", "\uacc4\ud68d\uc2dc\uc791", "planned start"],
    "plan_end_date": ["\uc885\ub8cc", "\uc644\ub8cc", "\uacc4\ud68d\uc885\ub8cc", "planned end"],
    "memo": ["\ube44\uace0", "\uba54\ubaa8", "remarks"],
}

COMMENCEMENT_INFO_LABELS = {
    "project_name": ["\uacf5\uc0ac\uba85", "\uc0ac\uc5c5\uba85", "\ud604\uc7a5\uba85"],
    "start_date": ["\ucc29\uacf5\uc77c", "\uacf5\uc0ac\uc2dc\uc791\uc77c", "\uc2dc\uc791\uc77c"],
    "end_date": ["\uc900\uacf5\uc77c", "\uacf5\uc0ac\uc885\ub8cc\uc77c", "\uc885\ub8cc\uc77c"],
}


def _normalize_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", text)


def _normalize_key(value) -> str:
    return re.sub(r"[^0-9A-Za-z\uac00-\ud7a3]+", "", _normalize_text(value)).lower()


def _coerce_decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = _normalize_text(value)
    if not text:
        return None
    text = text.replace(",", "").replace("%", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return Decimal(match.group(0))
    except InvalidOperation:
        return None


def _coerce_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _normalize_text(value)
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _load_workbook_pair(source):
    try:
        if hasattr(source, "read"):
            payload = source.read()
            if hasattr(source, "seek"):
                source.seek(0)
        else:
            payload = Path(source).read_bytes()
        return (
            load_workbook(BytesIO(payload), data_only=True),
            load_workbook(BytesIO(payload), data_only=False),
        )
    except Exception as exc:
        raise ValidationError("\uc5d1\uc140 \ud30c\uc77c\uc744 \uc77d\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.") from exc


def _build_merged_value_map(worksheet) -> dict[tuple[int, int], object]:
    merged_map: dict[tuple[int, int], object] = {}
    for cell_range in worksheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = cell_range.bounds
        anchor_value = worksheet.cell(min_row, min_col).value
        for row_idx in range(min_row, max_row + 1):
            for col_idx in range(min_col, max_col + 1):
                merged_map[(row_idx, col_idx)] = anchor_value
    return merged_map


def _safe_cell_value(worksheet, row_idx: int, col_idx: int, merged_map) -> object:
    value = worksheet.cell(row_idx, col_idx).value
    if value is not None:
        return value
    return merged_map.get((row_idx, col_idx))


def _row_values(worksheet, row_idx: int, merged_map) -> list[object]:
    return [
        _safe_cell_value(worksheet, row_idx, col_idx, merged_map)
        for col_idx in range(1, (worksheet.max_column or 0) + 1)
    ]


def _row_has_values(row_values) -> bool:
    return any(_normalize_text(value) for value in row_values)


def _sheet_bonus(sheet_name: str, preferred_names: list[str]) -> int:
    normalized = _normalize_text(sheet_name)
    for idx, preferred in enumerate(preferred_names):
        if normalized == preferred:
            return 100 - idx * 10
        if preferred in normalized:
            return 80 - idx * 10
    return 0


def _sheet_title_matches(worksheet, preferred_names: list[str]) -> bool:
    title = _normalize_text(worksheet.title)
    return any(preferred in title for preferred in preferred_names)


def _detect_header_row(worksheet, aliases_map: dict[str, list[str]], max_scan_rows: int = 30):
    merged_map = _build_merged_value_map(worksheet)
    best_row = None
    best_mapping = {}
    best_score = 0
    max_row = min(worksheet.max_row or 0, max_scan_rows)
    for row_idx in range(1, max_row + 1):
        normalized = [_normalize_key(value) for value in _row_values(worksheet, row_idx, merged_map)]
        mapping = {}
        used_columns = set()
        for field_name, aliases in aliases_map.items():
            normalized_aliases = [_normalize_key(alias) for alias in aliases]
            best_column = None
            best_match_score = 0
            for col_idx, cell_text in enumerate(normalized, start=1):
                if not cell_text or col_idx in used_columns:
                    continue
                match_score = 0
                for alias in normalized_aliases:
                    if not alias:
                        continue
                    if cell_text == alias:
                        match_score = max(match_score, 3)
                    elif cell_text.startswith(alias):
                        match_score = max(match_score, 2)
                    elif alias in cell_text:
                        match_score = max(match_score, 1)
                if match_score > best_match_score:
                    best_column = col_idx
                    best_match_score = match_score
            if best_column is not None:
                mapping[field_name] = best_column
                used_columns.add(best_column)
        if len(mapping) > best_score:
            best_row = row_idx
            best_mapping = mapping
            best_score = len(mapping)
    if best_row is None or best_score < 3:
        return None, {}
    return best_row, best_mapping


def _find_best_sheet(workbook, aliases_map, preferred_names, *, max_scan_rows=30):
    candidates = []
    for worksheet in workbook.worksheets:
        header_row, mapping = _detect_header_row(
            worksheet,
            aliases_map,
            max_scan_rows=max_scan_rows,
        )
        if not header_row:
            continue
        score = len(mapping) + _sheet_bonus(worksheet.title, preferred_names)
        candidates.append((score, worksheet, header_row, mapping))
    if not candidates:
        return None, None, {}
    candidates.sort(key=lambda item: item[0], reverse=True)
    _, worksheet, header_row, mapping = candidates[0]
    return worksheet, header_row, mapping


def _detect_budget_multilevel_layout(worksheet):
    if not _sheet_title_matches(worksheet, BUDGET_SHEET_PREFERENCES):
        return None
    merged_map = _build_merged_value_map(worksheet)
    max_row = min(worksheet.max_row or 0, 8)
    for row_idx in range(1, max_row):
        parent = [_normalize_key(value) for value in _row_values(worksheet, row_idx, merged_map)]
        child = [_normalize_key(value) for value in _row_values(worksheet, row_idx + 1, merged_map)]
        if len(parent) < 14 or len(child) < 13:
            continue
        if parent[0] != _normalize_key("\uacf5\uc885") or parent[1] != _normalize_key("\ud488\uba85") or parent[2] != _normalize_key("\uaddc\uaca9"):
            continue
        if parent[3] != _normalize_key("\uc218\ub7c9") or parent[4] != _normalize_key("\ub2e8\uc704"):
            continue
        if parent[5] != _normalize_key("\ub3c4\uae09\uae08\uc561") or parent[7] != _normalize_key("\ub178\ubb34\ube44"):
            continue
        if parent[9] != _normalize_key("\uc7ac\ub8cc\ube44") or parent[11] != _normalize_key("\uacbd\ube44"):
            continue
        if child[5] != _normalize_key("\ub2e8\uac00") or child[6] != _normalize_key("\uae08\uc561"):
            continue
        if child[7] != _normalize_key("\ub2e8\uac00") or child[8] != _normalize_key("\uae08\uc561"):
            continue
        if child[9] != _normalize_key("\ub2e8\uac00") or child[10] != _normalize_key("\uae08\uc561"):
            continue
        if child[11] != _normalize_key("\ub2e8\uac00") or child[12] != _normalize_key("\uae08\uc561"):
            continue
        return {
            "worksheet": worksheet,
            "header_row": row_idx + 1,
            "data_start_row": row_idx + 2,
            "mapping": {
                "code": 1,
                "item_name": 2,
                "spec": 3,
                "quantity": 4,
                "unit": 5,
                "unit_price": 6,
                "amount": 7,
                "labor_amount": 9,
                "material_amount": 11,
                "expense_amount": 13,
                "memo": 14,
            },
        }
    return None


def _detect_wbs_schedule_layout(worksheet):
    if "\uc608\uc815\uacf5\uc815\ud45c" not in _normalize_text(worksheet.title):
        return None
    merged_map = _build_merged_value_map(worksheet)
    max_row = min(worksheet.max_row or 0, 8)
    for row_idx in range(1, max_row):
        parent = [_normalize_key(value) for value in _row_values(worksheet, row_idx, merged_map)]
        child_values = _row_values(worksheet, row_idx + 1, merged_map)
        if len(parent) < 13 or len(child_values) < 13:
            continue
        if parent[0] != _normalize_key("\uacf5\uc885") or parent[2] != _normalize_key("\ubcf4\ud560"):
            continue
        if parent[3] != _normalize_key("\ucc29\uc218") or parent[12] != _normalize_key("\ube44\uace0"):
            continue
        period_columns = []
        for col_idx in range(4, 10):
            day_value = _coerce_decimal(child_values[col_idx - 1])
            if day_value is None:
                period_columns = []
                break
            label_suffix = str(int(day_value)) if day_value == int(day_value) else _normalize_text(day_value)
            period_columns.append(
                {
                    "column": col_idx,
                    "period": f"\ucc29\uc218+{label_suffix}",
                    "offset_days": int(day_value),
                }
            )
        if period_columns:
            return {
                "worksheet": worksheet,
                "header_row": row_idx + 1,
                "data_start_row": row_idx + 2,
                "name_column": 1,
                "weight_column": 3,
                "memo_column": 13,
                "period_columns": period_columns,
            }
    return None


def _extract_value_pair(data_ws, formula_ws, row_idx, col_idx, data_merged_map, formula_merged_map):
    if not col_idx:
        return None, None
    return (
        _safe_cell_value(data_ws, row_idx, col_idx, data_merged_map),
        _safe_cell_value(formula_ws, row_idx, col_idx, formula_merged_map),
    )


def _cell_has_formula(raw_value) -> bool:
    return isinstance(raw_value, str) and raw_value.startswith("=")


def _append_formula_warning(row_warnings: list[str], raw_value, cached_value, label: str):
    if _cell_has_formula(raw_value) and cached_value in (None, ""):
        row_warnings.append(f"{label} \uc218\uc2dd\uc758 \uacc4\uc0b0\uac12\uc774 \uc5c6\uc5b4 \ud655\uc778\uc774 \ud544\uc694\ud569\ub2c8\ub2e4.")


def _infer_hierarchy(code: str, item_name: str) -> str:
    code_text = _normalize_text(code)
    if code_text:
        return str(max(len([token for token in re.split(r"[.\-]", code_text) if token]), 1))
    leading_spaces = len(str(item_name or "")) - len(str(item_name or "").lstrip())
    if leading_spaces > 0:
        return str(leading_spaces // 2 + 1)
    return ""


def _looks_like_budget_header_row(row) -> bool:
    code_key = _normalize_key(row.get("code"))
    item_key = _normalize_key(row.get("item_name"))
    unit_key = _normalize_key(row.get("unit"))
    return (
        code_key == _normalize_key("\uacf5\uc885")
        or item_key in {_normalize_key("\ud488\uba85"), _normalize_key("\ud56d\ubaa9\uba85")}
        or unit_key == _normalize_key("\ub2e8\uc704")
    )


def _is_summary_like_row(item_name: str, code: str) -> bool:
    text = f"{item_name} {code}".strip()
    return any(keyword in text for keyword in ["\ud569\uacc4", "\ucd1d\uacc4", "\uc18c\uacc4"])


def _find_summary_total(workbook) -> tuple[str | None, Decimal | None]:
    for preferred_name in BUDGET_SUMMARY_SHEET_PREFERENCES:
        for worksheet in workbook.worksheets:
            if preferred_name not in worksheet.title:
                continue
            merged_map = _build_merged_value_map(worksheet)
            for row_idx in range(1, min(worksheet.max_row or 0, 60) + 1):
                for col_idx in range(1, min(worksheet.max_column or 0, 20) + 1):
                    text = _normalize_text(_safe_cell_value(worksheet, row_idx, col_idx, merged_map))
                    if "\ud569\uacc4" not in text and "\ucd1d\uacc4" not in text:
                        continue
                    for scan_col in range(col_idx + 1, min(worksheet.max_column or 0, col_idx + 4) + 1):
                        amount = _coerce_decimal(_safe_cell_value(worksheet, row_idx, scan_col, merged_map))
                        if amount is not None:
                            return worksheet.title, amount
            return worksheet.title, None
    return None, None


def _extract_nearby_amount(worksheet, merged_map, row_idx, col_idx, *, max_scan_cols=4):
    for scan_col in range(col_idx + 1, min(worksheet.max_column or 0, col_idx + max_scan_cols) + 1):
        amount = _coerce_decimal(_safe_cell_value(worksheet, row_idx, scan_col, merged_map))
        if amount is not None:
            return amount
    return None


def _summary_aliases(*values: str) -> list[str]:
    return [_normalize_key(value) for value in values]


SUMMARY_HEADER_ALIASES = {
    "contract": _summary_aliases("도급금액", "도급 금액", "계약금액", "도급예정액", "금액"),
    "labor": _summary_aliases("노무비", "노 무 비", "노무비금액"),
    "material": _summary_aliases("재료비", "재 료 비", "재료비금액"),
    "expense": _summary_aliases("경비", "경 비", "경비금액"),
}

SUMMARY_ROW_LABEL_ALIASES = {
    "contract": _summary_aliases("도급예정액", "도급금액", "계약금액", "도급계약금액"),
    "owner_supplied": _summary_aliases("관급자재대", "관급자재", "관급"),
    "total_construction": _summary_aliases("총공사비"),
}

SUMMARY_PRIMARY_ROW_LABEL_ALIASES = {
    "contract": _summary_aliases("도급예정액"),
    "owner_supplied": _summary_aliases("관급자재대"),
    "total_construction": _summary_aliases("총공사비"),
}


def _summary_field_from_header_key(header_key: str) -> str | None:
    for field_name, aliases in SUMMARY_HEADER_ALIASES.items():
        for alias in aliases:
            if header_key == alias or header_key.endswith(alias) or alias in header_key:
                return field_name
    return None


def _find_summary_header_mapping(worksheet, merged_map, max_scan_rows: int = 30):
    max_row = min(worksheet.max_row or 0, max_scan_rows)
    best_mapping = {}
    best_row = None
    best_score = 0
    for row_idx in range(1, max_row + 1):
        current_keys = [_normalize_key(value) for value in _row_values(worksheet, row_idx, merged_map)]
        previous_keys = [_normalize_key(value) for value in _row_values(worksheet, row_idx - 1, merged_map)] if row_idx > 1 else []
        mapping = {}
        for col_idx, current_key in enumerate(current_keys, start=1):
            variants = {current_key}
            if row_idx > 1 and col_idx - 1 < len(previous_keys):
                previous_key = previous_keys[col_idx - 1]
                if previous_key or current_key:
                    variants.add(f"{previous_key}{current_key}")
            for variant in variants:
                field_name = _summary_field_from_header_key(variant)
                if field_name and field_name not in mapping:
                    mapping[field_name] = col_idx
        score = len(mapping) + (1 if "contract" in mapping else 0)
        if score > best_score and len(mapping) >= 3:
            best_mapping = mapping
            best_row = row_idx
            best_score = score
    return best_row, best_mapping


def _find_summary_row(
    worksheet,
    merged_map,
    primary_aliases: list[str],
    fallback_aliases: list[str] | None = None,
    max_scan_rows: int = 80,
) -> int | None:
    max_row = min(worksheet.max_row or 0, max_scan_rows)
    primary_set = set(primary_aliases)
    fallback_set = set(fallback_aliases or [])
    fallback_row_idx = None
    for row_idx in range(1, max_row + 1):
        row_values = _row_values(worksheet, row_idx, merged_map)
        label_keys = [_normalize_key(value) for value in row_values[:3]]
        if any(cell_key in primary_set for cell_key in label_keys if cell_key):
            return row_idx
        if fallback_row_idx is None and any(
            cell_key in fallback_set for cell_key in label_keys if cell_key
        ):
            fallback_row_idx = row_idx
    return fallback_row_idx


def _read_row_value_by_column(worksheet, merged_map, row_idx: int | None, col_idx: int | None) -> Decimal | None:
    if not row_idx or not col_idx:
        return None
    return _coerce_decimal(_safe_cell_value(worksheet, row_idx, col_idx, merged_map))


def _is_suspicious_bucket_amount(value: Decimal | None, contract_expected_amount: Decimal | None) -> bool:
    if value is None:
        return True
    if contract_expected_amount is None or contract_expected_amount <= Decimal("1000000"):
        return False
    return value in {Decimal("0"), Decimal("1"), Decimal("-1")}


def _extract_best_vertical_amount(worksheet, merged_map, row_idx: int, col_idx: int) -> Decimal | None:
    candidates: list[Decimal] = []
    for scan_col in range(col_idx + 1, min(worksheet.max_column or 0, col_idx + 6) + 1):
        amount = _coerce_decimal(_safe_cell_value(worksheet, row_idx, scan_col, merged_map))
        if amount is not None:
            candidates.append(amount)
    if not candidates:
        return None
    meaningful = [amount for amount in candidates if amount not in {Decimal("0"), Decimal("1"), Decimal("-1")}]
    if meaningful:
        return max(meaningful)
    return candidates[0]


def _fill_summary_from_vertical_labels(worksheet, merged_map, summary: dict):
    for row_idx in range(1, min(worksheet.max_row or 0, 80) + 1):
        row_values = _row_values(worksheet, row_idx, merged_map)
        for col_idx, cell_value in enumerate(row_values, start=1):
            cell_key = _normalize_key(cell_value)
            if not cell_key:
                continue
            if summary["contract_expected_amount"] is None and cell_key in SUMMARY_ROW_LABEL_ALIASES["contract"]:
                summary["contract_amount_label"] = _normalize_text(cell_value)
                summary["contract_expected_amount"] = _extract_best_vertical_amount(worksheet, merged_map, row_idx, col_idx)
            elif summary["contract_labor_amount"] is None and cell_key in SUMMARY_HEADER_ALIASES["labor"]:
                summary["contract_labor_amount"] = _extract_best_vertical_amount(worksheet, merged_map, row_idx, col_idx)
            elif summary["contract_material_amount"] is None and cell_key in SUMMARY_HEADER_ALIASES["material"]:
                summary["contract_material_amount"] = _extract_best_vertical_amount(worksheet, merged_map, row_idx, col_idx)
            elif summary["contract_expense_amount"] is None and cell_key in SUMMARY_HEADER_ALIASES["expense"]:
                summary["contract_expense_amount"] = _extract_best_vertical_amount(worksheet, merged_map, row_idx, col_idx)
            elif summary["owner_supplied_amount"] is None and cell_key in SUMMARY_ROW_LABEL_ALIASES["owner_supplied"]:
                summary["owner_supplied_amount"] = _extract_best_vertical_amount(worksheet, merged_map, row_idx, col_idx)
            elif summary["total_construction_amount"] is None and cell_key in SUMMARY_ROW_LABEL_ALIASES["total_construction"]:
                summary["total_construction_amount"] = _extract_best_vertical_amount(worksheet, merged_map, row_idx, col_idx)


def parse_budget_summary_sheet(workbook) -> dict:
    summary = {
        "summary_sheet_name": "",
        "contract_expected_amount": None,
        "contract_labor_amount": None,
        "contract_material_amount": None,
        "contract_expense_amount": None,
        "contract_amount_label": "",
        "owner_supplied_amount": None,
        "total_construction_amount": None,
        "warnings": [],
    }
    target_sheet = None
    for preferred_name in BUDGET_SUMMARY_SHEET_PREFERENCES:
        for worksheet in workbook.worksheets:
            if preferred_name == _normalize_text(worksheet.title) or preferred_name in worksheet.title:
                target_sheet = worksheet
                break
        if target_sheet is not None:
            break
    if target_sheet is None:
        return summary

    summary["summary_sheet_name"] = target_sheet.title
    merged_map = _build_merged_value_map(target_sheet)
    _header_row_idx, header_mapping = _find_summary_header_mapping(target_sheet, merged_map)
    contract_row_idx = _find_summary_row(
        target_sheet,
        merged_map,
        SUMMARY_PRIMARY_ROW_LABEL_ALIASES["contract"],
        SUMMARY_ROW_LABEL_ALIASES["contract"],
    )
    owner_row_idx = _find_summary_row(
        target_sheet,
        merged_map,
        SUMMARY_PRIMARY_ROW_LABEL_ALIASES["owner_supplied"],
        SUMMARY_ROW_LABEL_ALIASES["owner_supplied"],
    )
    total_row_idx = _find_summary_row(
        target_sheet,
        merged_map,
        SUMMARY_PRIMARY_ROW_LABEL_ALIASES["total_construction"],
        SUMMARY_ROW_LABEL_ALIASES["total_construction"],
    )

    if header_mapping and contract_row_idx:
        summary["contract_amount_label"] = "도급예정액"
        summary["contract_expected_amount"] = _read_row_value_by_column(target_sheet, merged_map, contract_row_idx, header_mapping.get("contract"))
        summary["contract_labor_amount"] = _read_row_value_by_column(target_sheet, merged_map, contract_row_idx, header_mapping.get("labor"))
        summary["contract_material_amount"] = _read_row_value_by_column(target_sheet, merged_map, contract_row_idx, header_mapping.get("material"))
        summary["contract_expense_amount"] = _read_row_value_by_column(target_sheet, merged_map, contract_row_idx, header_mapping.get("expense"))
        summary["owner_supplied_amount"] = _read_row_value_by_column(target_sheet, merged_map, owner_row_idx, header_mapping.get("contract"))
        summary["total_construction_amount"] = _read_row_value_by_column(target_sheet, merged_map, total_row_idx, header_mapping.get("contract"))

    if _is_suspicious_bucket_amount(summary["contract_labor_amount"], summary["contract_expected_amount"]):
        summary["contract_labor_amount"] = None
    if _is_suspicious_bucket_amount(summary["contract_material_amount"], summary["contract_expected_amount"]):
        summary["contract_material_amount"] = None
    if _is_suspicious_bucket_amount(summary["contract_expense_amount"], summary["contract_expected_amount"]):
        summary["contract_expense_amount"] = None

    if (
        summary["contract_expected_amount"] is None
        or summary["contract_labor_amount"] is None
        or summary["contract_material_amount"] is None
        or summary["contract_expense_amount"] is None
        or summary["owner_supplied_amount"] is None
        or summary["total_construction_amount"] is None
    ):
        _fill_summary_from_vertical_labels(target_sheet, merged_map, summary)

    if summary["summary_sheet_name"] and summary["contract_expected_amount"] is None:
        summary["warnings"].append(
            f"{summary['summary_sheet_name']} 시트에서 도급예정액을 찾지 못했습니다."
        )

    if any(
        value is None
        for value in [
            summary["contract_labor_amount"],
            summary["contract_material_amount"],
            summary["contract_expense_amount"],
        ]
    ):
        summary["warnings"].append(
            "총괄표 3분류 금액을 안정적으로 읽지 못해 상세행 합계를 참고값으로 표시합니다."
        )

    if (
        summary["contract_expected_amount"] is not None
        and summary["contract_labor_amount"] is not None
        and summary["contract_material_amount"] is not None
        and summary["contract_expense_amount"] is not None
    ):
        bucket_total = (
            summary["contract_labor_amount"]
            + summary["contract_material_amount"]
            + summary["contract_expense_amount"]
        )
        if abs(bucket_total - summary["contract_expected_amount"]) > Decimal("1"):
            summary["warnings"].append(
                "총괄표 도급예정액이 노무비+재료비+경비 합계와 일치하지 않습니다."
            )
    return summary


def _parse_budget_rows(data_ws, formula_ws, mapping, data_start_row):
    data_merged_map = _build_merged_value_map(data_ws)
    formula_merged_map = _build_merged_value_map(formula_ws)
    rows = []
    blank_streak = 0
    for row_idx in range(data_start_row, (data_ws.max_row or 0) + 1):
        values = _row_values(data_ws, row_idx, data_merged_map)
        if not _row_has_values(values):
            blank_streak += 1
            if blank_streak >= 10:
                break
            continue
        blank_streak = 0
        row_warnings = []
        code_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("code"), data_merged_map, formula_merged_map)
        item_name_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("item_name"), data_merged_map, formula_merged_map)
        spec_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("spec"), data_merged_map, formula_merged_map)
        quantity_value, quantity_raw = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("quantity"), data_merged_map, formula_merged_map)
        unit_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("unit"), data_merged_map, formula_merged_map)
        unit_price_value, unit_price_raw = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("unit_price"), data_merged_map, formula_merged_map)
        amount_value, amount_raw = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("amount"), data_merged_map, formula_merged_map)
        labor_value, labor_raw = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("labor_amount"), data_merged_map, formula_merged_map)
        material_value, material_raw = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("material_amount"), data_merged_map, formula_merged_map)
        expense_value, expense_raw = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("expense_amount"), data_merged_map, formula_merged_map)
        memo_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("memo"), data_merged_map, formula_merged_map)

        row = {
            "sheet_name": data_ws.title,
            "row_no": row_idx,
            "row_number": row_idx,
            "code": _normalize_text(code_value),
            "hierarchy": "",
            "item_name": _normalize_text(item_name_value),
            "spec": _normalize_text(spec_value),
            "quantity": _coerce_decimal(quantity_value),
            "unit": _normalize_text(unit_value),
            "unit_price": _coerce_decimal(unit_price_value),
            "amount": _coerce_decimal(amount_value),
            "labor_amount": _coerce_decimal(labor_value),
            "material_amount": _coerce_decimal(material_value),
            "expense_amount": _coerce_decimal(expense_value),
            "memo": _normalize_text(memo_value),
            "warnings": row_warnings,
        }
        row["hierarchy"] = _infer_hierarchy(row["code"], row["item_name"])
        if _looks_like_budget_header_row(row):
            continue
        if not " ".join([row["code"], row["item_name"], row["spec"]]).strip():
            continue
        if row["amount"] is None and row["quantity"] is not None and row["unit_price"] is not None:
            row["amount"] = row["quantity"] * row["unit_price"]
        if _is_summary_like_row(row["item_name"], row["code"]) and row["amount"] is None:
            continue

        _append_formula_warning(row_warnings, quantity_raw, quantity_value, "\uc218\ub7c9")
        _append_formula_warning(row_warnings, unit_price_raw, unit_price_value, "\ub2e8\uac00")
        _append_formula_warning(row_warnings, amount_raw, amount_value, "\uae08\uc561")
        _append_formula_warning(row_warnings, labor_raw, labor_value, "\ub178\ubb34\ube44")
        _append_formula_warning(row_warnings, material_raw, material_value, "\uc7ac\ub8cc\ube44")
        _append_formula_warning(row_warnings, expense_raw, expense_value, "\uacbd\ube44")
        if row["amount"] is None:
            row_warnings.append("\uae08\uc561\uc744 \ud655\uc778\ud574 \uc8fc\uc138\uc694.")
        rows.append(row)
    return rows


def _extract_commencement_info(workbook, scheduled_ws=None) -> dict:
    info = {"project_name": "", "start_date": None, "end_date": None}
    if scheduled_ws is not None:
        merged_map = _build_merged_value_map(scheduled_ws)
        for row_idx in range(1, min(scheduled_ws.max_row or 0, 5) + 1):
            row_text = " ".join(
                _normalize_text(value)
                for value in _row_values(scheduled_ws, row_idx, merged_map)
                if _normalize_text(value)
            )
            match = re.search(r"\uacf5\uc0ac\uba85\s*:\s*(.+)", row_text)
            if match:
                info["project_name"] = match.group(1).strip()
                break
    for worksheet in workbook.worksheets:
        if not any(name in worksheet.title for name in COMMENCEMENT_INFO_SHEET_PREFERENCES):
            continue
        merged_map = _build_merged_value_map(worksheet)
        for row_idx in range(1, min(worksheet.max_row or 0, 40) + 1):
            for col_idx in range(1, min(worksheet.max_column or 0, 12) + 1):
                cell_value = _safe_cell_value(worksheet, row_idx, col_idx, merged_map)
                normalized = _normalize_text(cell_value)
                if not normalized:
                    continue
                for field_name, labels in COMMENCEMENT_INFO_LABELS.items():
                    if normalized not in labels:
                        continue
                    for scan_col in range(col_idx + 1, min(worksheet.max_column or 0, col_idx + 3) + 1):
                        neighbor = _safe_cell_value(worksheet, row_idx, scan_col, merged_map)
                        if field_name == "project_name":
                            project_name = _normalize_text(neighbor)
                            if project_name:
                                info["project_name"] = project_name
                        else:
                            parsed_date = _coerce_date(neighbor)
                            if parsed_date:
                                info[field_name] = parsed_date
    return info


def _format_period_value(value):
    if value in (None, ""):
        return ""
    return _normalize_text(value)


def parse_contract_budget_excel(file_or_path) -> dict:
    data_book, formula_book = _load_workbook_pair(file_or_path)

    layout = None
    for worksheet in data_book.worksheets:
        layout = _detect_budget_multilevel_layout(worksheet)
        if layout:
            break

    if layout is not None:
        data_ws = layout["worksheet"]
        formula_ws = formula_book[data_ws.title]
        header_row = layout["header_row"]
        rows = _parse_budget_rows(data_ws, formula_ws, layout["mapping"], layout["data_start_row"])
    else:
        data_ws, header_row, mapping = _find_best_sheet(
            data_book,
            BUDGET_HEADER_ALIASES,
            BUDGET_SHEET_PREFERENCES,
        )
        if data_ws is None:
            raise ValidationError("\uacc4\uc57d \uc608\uc0b0 \uc5d1\uc140\uc5d0\uc11c \ub0b4\uc5ed\uc11c \uc2dc\ud2b8\ub97c \ucc3e\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.")
        formula_ws = formula_book[data_ws.title]
        rows = _parse_budget_rows(data_ws, formula_ws, mapping, header_row + 1)

    if not rows:
        raise ValidationError("\uacc4\uc57d \uc608\uc0b0 \uc5d1\uc140\uc5d0\uc11c \uc608\uc0b0 \ub77c\uc778\uc744 \ucc3e\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.")

    summary_sheet = parse_budget_summary_sheet(data_book)
    summary_sheet_name = summary_sheet.get("summary_sheet_name")
    summary_total = summary_sheet.get("contract_expected_amount")
    calculated_total = sum((row["amount"] or Decimal("0")) for row in rows)
    warnings = list(summary_sheet.get("warnings") or [])
    if summary_total is not None and calculated_total != summary_total:
        warnings.append(
            f"{summary_sheet_name} \ub3c4\uae09\uc608\uc815\uc561({summary_total})\uacfc \ucd94\ucd9c \ud569\uacc4({calculated_total})\uac00 \ub2e4\ub985\ub2c8\ub2e4."
        )

    warning_count = len(warnings) + sum(len(row["warnings"]) for row in rows)
    return {
        "sheet_name": data_ws.title,
        "header_row": header_row,
        "rows": rows,
        "warnings": warnings,
        "warning_count": warning_count,
        "summary_sheet": summary_sheet,
        "summary_validation": {
            "sheet_name": summary_sheet_name,
            "summary_total": summary_total,
            "calculated_total": calculated_total,
        },
    }


def parse_commencement_wbs_excel(file_or_path) -> dict:
    data_book, formula_book = _load_workbook_pair(file_or_path)

    layout = None
    for worksheet in data_book.worksheets:
        layout = _detect_wbs_schedule_layout(worksheet)
        if layout:
            break

    if layout is not None:
        data_ws = layout["worksheet"]
        formula_ws = formula_book[data_ws.title]
        data_merged_map = _build_merged_value_map(data_ws)
        formula_merged_map = _build_merged_value_map(formula_ws)
        commencement_info = _extract_commencement_info(data_book, scheduled_ws=data_ws)
        rows = []
        seen_keys = set()
        blank_streak = 0
        for row_idx in range(layout["data_start_row"], (data_ws.max_row or 0) + 1):
            physical_name_raw = data_ws.cell(row_idx, layout["name_column"]).value
            physical_weight_raw = data_ws.cell(row_idx, layout["weight_column"]).value
            physical_row_values = [
                data_ws.cell(row_idx, col_idx).value
                for col_idx in range(1, (data_ws.max_column or 0) + 1)
            ]

            if not _row_has_values(physical_row_values):
                blank_streak += 1
                if blank_streak >= 10:
                    break
                continue
            blank_streak = 0

            physical_name_text = _normalize_text(physical_name_raw)
            physical_name_key = _normalize_key(physical_name_text)
            physical_weight_decimal = _coerce_decimal(physical_weight_raw)

            if not physical_name_text:
                continue
            if physical_name_key in {
                _normalize_key("\uacf5\uc885"),
                _normalize_key("\uacf5\uc885\uba85"),
                _normalize_key("\uc18c\uacc4"),
                _normalize_key("\ub204\uacc4"),
                _normalize_key("\ud569\uacc4"),
                _normalize_key("0"),
            }:
                continue
            if physical_weight_decimal is None or physical_weight_decimal == 0:
                continue

            _, weight_raw = _extract_value_pair(
                data_ws, formula_ws, row_idx, layout["weight_column"], data_merged_map, formula_merged_map
            )
            memo_value, _ = _extract_value_pair(
                data_ws, formula_ws, row_idx, layout["memo_column"], data_merged_map, formula_merged_map
            )

            row_warnings = []
            _append_formula_warning(row_warnings, weight_raw, physical_weight_raw, "\uac00\uc911\uce58")
            weight_decimal = physical_weight_decimal
            if Decimal("0") < weight_decimal <= Decimal("1"):
                weight_decimal *= Decimal("100")
            seen_key = (
                _normalize_key(physical_name_text),
                str(weight_decimal.quantize(Decimal("0.0000000001"))),
            )
            if seen_key in seen_keys:
                continue
            seen_keys.add(seen_key)

            row = {
                "sheet_name": data_ws.title,
                "row_no": row_idx,
                "row_number": row_idx,
                "code": physical_name_text,
                "name": physical_name_text,
                "wbs_name": physical_name_text,
                "weight": weight_decimal,
                "plan_start_date": None,
                "plan_end_date": None,
                "memo": _normalize_text(memo_value),
                "period_progress_values": [],
                "warnings": row_warnings,
            }
            for period in layout["period_columns"]:
                period_value, _ = _extract_value_pair(
                    data_ws,
                    formula_ws,
                    row_idx,
                    period["column"],
                    data_merged_map,
                    formula_merged_map,
                )
                row["period_progress_values"].append(
                    {
                        "period": period["period"],
                        "offset_days": period["offset_days"],
                        "value": _format_period_value(period_value),
                    }
                )
            rows.append(row)

        if not rows:
            raise ValidationError("\ucc29\uacf5\uacc4 \uc5d1\uc140\uc5d0\uc11c WBS \ub77c\uc778\uc744 \ucc3e\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.")

        warning_count = sum(len(row["warnings"]) for row in rows)
        return {
            "sheet_name": data_ws.title,
            "header_row": layout["header_row"],
            "rows": rows,
            "warnings": [],
            "warning_count": warning_count,
            "project_name": commencement_info.get("project_name") or "",
            "start_date": commencement_info.get("start_date"),
            "end_date": commencement_info.get("end_date"),
        }

    data_ws, header_row, mapping = _find_best_sheet(
        data_book,
        WBS_HEADER_ALIASES,
        WBS_SHEET_PREFERENCES,
        max_scan_rows=40,
    )
    if data_ws is None:
        raise ValidationError("\ucc29\uacf5\uacc4 \uc5d1\uc140\uc5d0\uc11c \uc608\uc815\uacf5\uc815\ud45c \uc2dc\ud2b8\ub97c \ucc3e\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.")

    formula_ws = formula_book[data_ws.title]
    data_merged_map = _build_merged_value_map(data_ws)
    formula_merged_map = _build_merged_value_map(formula_ws)
    commencement_info = _extract_commencement_info(data_book, scheduled_ws=data_ws)

    core_columns = set(mapping.values())
    period_columns = []
    for col_idx in range(1, (data_ws.max_column or 0) + 1):
        if col_idx in core_columns:
            continue
        header_text = _normalize_text(_safe_cell_value(data_ws, header_row, col_idx, data_merged_map))
        if header_text:
            period_columns.append((col_idx, header_text))

    rows = []
    blank_streak = 0
    for row_idx in range(header_row + 1, (data_ws.max_row or 0) + 1):
        values = _row_values(data_ws, row_idx, data_merged_map)
        if not _row_has_values(values):
            blank_streak += 1
            if blank_streak >= 10:
                break
            continue
        blank_streak = 0

        code_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("code"), data_merged_map, formula_merged_map)
        name_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("name"), data_merged_map, formula_merged_map)
        weight_value, weight_raw = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("weight"), data_merged_map, formula_merged_map)
        plan_start_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("plan_start_date"), data_merged_map, formula_merged_map)
        plan_end_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("plan_end_date"), data_merged_map, formula_merged_map)
        memo_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, mapping.get("memo"), data_merged_map, formula_merged_map)

        name_text = _normalize_text(name_value)
        code_text = _normalize_text(code_value)
        if not name_text and not code_text:
            continue

        row_warnings = []
        _append_formula_warning(row_warnings, weight_raw, weight_value, "\uac00\uc911\uce58")
        weight_decimal = _coerce_decimal(weight_value)
        if weight_decimal is None:
            weight_decimal = Decimal("0")
            row_warnings.append("\uac00\uc911\uce58\ub97c \ud655\uc778\ud574 \uc8fc\uc138\uc694.")

        row = {
            "sheet_name": data_ws.title,
            "row_no": row_idx,
            "row_number": row_idx,
            "code": code_text,
            "name": name_text,
            "wbs_name": name_text,
            "weight": weight_decimal,
            "plan_start_date": _coerce_date(plan_start_value),
            "plan_end_date": _coerce_date(plan_end_value),
            "memo": _normalize_text(memo_value),
            "period_progress_values": [],
            "warnings": row_warnings,
        }
        for col_idx, header_text in period_columns:
            period_value, _ = _extract_value_pair(data_ws, formula_ws, row_idx, col_idx, data_merged_map, formula_merged_map)
            if period_value in (None, ""):
                continue
            row["period_progress_values"].append(
                {"period": header_text, "offset_days": None, "value": _format_period_value(period_value)}
            )
        if not row["plan_start_date"] and not row["plan_end_date"] and row["period_progress_values"]:
            row_warnings.append("\uae30\uac04\uac12\uc744 \ucd94\uc815\ud558\uc9c0 \ubabb\ud574 \uc218\ub3d9 \ud655\uc778\uc774 \ud544\uc694\ud569\ub2c8\ub2e4.")
        rows.append(row)

    if not rows:
        raise ValidationError("\ucc29\uacf5\uacc4 \uc5d1\uc140\uc5d0\uc11c WBS \ub77c\uc778\uc744 \ucc3e\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.")

    warning_count = sum(len(row["warnings"]) for row in rows)
    return {
        "sheet_name": data_ws.title,
        "header_row": header_row,
        "rows": rows,
        "warnings": [],
        "warning_count": warning_count,
        "project_name": commencement_info.get("project_name") or "",
        "start_date": commencement_info.get("start_date"),
        "end_date": commencement_info.get("end_date"),
    }


def parse_budget_workbook(source) -> dict:
    return parse_contract_budget_excel(source)


def parse_wbs_workbook(source) -> dict:
    return parse_commencement_wbs_excel(source)
