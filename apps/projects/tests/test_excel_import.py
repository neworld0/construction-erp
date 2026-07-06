from io import BytesIO

import pytest
from django.core.exceptions import ValidationError
from openpyxl import Workbook

import apps.projects.excel_import as ei


def _workbook_bytes(workbook) -> BytesIO:
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def test_excel_import_namespace_exports_expected_functions():
    assert hasattr(ei, "_detect_budget_multilevel_layout")
    assert hasattr(ei, "_detect_wbs_schedule_layout")
    assert hasattr(ei, "parse_contract_budget_excel")
    assert hasattr(ei, "parse_commencement_wbs_excel")
    assert hasattr(ei, "parse_budget_workbook")
    assert hasattr(ei, "parse_wbs_workbook")
    assert ei.parse_contract_budget_excel.__name__ == "parse_contract_budget_excel"
    assert ei.parse_commencement_wbs_excel.__name__ == "parse_commencement_wbs_excel"


def test_parse_contract_budget_excel_handles_generic_sheet_and_formula_fallback():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "\ub3c4\uae09\uacc4\uc57d\ub0b4\uc5ed\uc11c"
    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "\uacc4\uc57d\ub0b4\uc5ed"
    worksheet.append(["", "", "", "", "", "", "", "", "", "", ""])
    worksheet.append(
        [
            "\uacf5\uc885",
            "\ud488\uba85",
            "\uaddc\uaca9",
            "\uc218\ub7c9",
            "\ub2e8\uc704",
            "\ub2e8\uac00",
            "\uae08\uc561",
            "\ub178\ubb34\ube44",
            "\uc7ac\ub8cc\ube44",
            "\uacbd\ube44",
            "\ube44\uace0",
        ]
    )
    worksheet.append(["A-100", "\uc544\uc2a4\ucf58 \ud3ec\uc7a5", "t=5cm", 10, "m2", 15000, 150000, 30000, 100000, 20000, ""])
    worksheet.append(["A-200", "\ucc28\uc120 \ub3c4\uc0c9", "", 5, "\uc2dd", 5000, "=D5*F5", "", "", "", ""])

    result = ei.parse_contract_budget_excel(_workbook_bytes(workbook))

    assert result["sheet_name"] == "\ub3c4\uae09\uacc4\uc57d\ub0b4\uc5ed\uc11c"
    assert len(result["rows"]) == 2
    assert result["rows"][1]["amount"] == 25000
    assert any("\uc218\uc2dd" in warning for warning in result["rows"][1]["warnings"])


def test_parse_contract_budget_excel_real_layout_maps_amount_columns_and_skips_headers():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "\ub3c4\uae09\uacc4\uc57d\ub0b4\uc5ed\uc11c"
    worksheet["A1"] = "\uad6d\ub3c446\ud638\uc120 \ud638\ud3c9IC\uad50(\uc0c1)\ub4f1 \ud3ec\uc7a5\ubc0f\uc2dc\uc124\ubb3c \ubcf4\uc218\uacf5\uc0ac"
    worksheet["A3"] = "\uacf5\uc885"
    worksheet["B3"] = "\ud488\uba85"
    worksheet["C3"] = "\uaddc\uaca9"
    worksheet["D3"] = "\uc218\ub7c9"
    worksheet["E3"] = "\ub2e8\uc704"
    worksheet["F3"] = "\ub3c4\uae09\uae08\uc561"
    worksheet["H3"] = "\ub178\ubb34\ube44"
    worksheet["J3"] = "\uc7ac\ub8cc\ube44"
    worksheet["L3"] = "\uacbd\ube44"
    worksheet["N3"] = "\ube44\uace0"
    worksheet["O3"] = "\ud22c\ucc30\uc728"
    worksheet["F4"] = "\ub2e8\uac00"
    worksheet["G4"] = "\uae08\uc561"
    worksheet["H4"] = "\ub2e8\uac00"
    worksheet["I4"] = "\uae08\uc561"
    worksheet["J4"] = "\ub2e8\uac00"
    worksheet["K4"] = "\uae08\uc561"
    worksheet["L4"] = "\ub2e8\uac00"
    worksheet["M4"] = "\uae08\uc561"
    worksheet["A5"] = "1"
    worksheet["B5"] = "\uc808\uc0ad \ud3ec\uc7a5"
    worksheet["C5"] = "A-Type"
    worksheet["D5"] = 100
    worksheet["E5"] = "\u33a1"
    worksheet["F5"] = 1000
    worksheet["G5"] = 100000
    worksheet["H5"] = 300
    worksheet["I5"] = 30000
    worksheet["J5"] = 500
    worksheet["K5"] = 50000
    worksheet["L5"] = 200
    worksheet["M5"] = 20000
    worksheet["B8"] = "\uc544\uc2a4\ud314\ud2b8 \ud3ec\uc7a5 \uc808\uc0ad\ud6c4 \uc544\uc2a4\ud314\ud2b8 \ub367\uc50c\uc6b0\uae30"
    worksheet["C8"] = "A-Type (1\ud68c\uc808\uc0ad,2\ud68c\ud3ec\uc7a5)-\uc57c\uac04\uc791\uc5c5"
    worksheet["D8"] = 1969
    worksheet["E8"] = "\u33a1"
    worksheet["F8"] = 6368
    worksheet["G8"] = 12538592
    worksheet["H8"] = 3674
    worksheet["I8"] = 7234106
    worksheet["J8"] = 1208
    worksheet["K8"] = 2378552
    worksheet["L8"] = 1486
    worksheet["M8"] = 2925934

    result = ei.parse_contract_budget_excel(_workbook_bytes(workbook))
    target_row = next(row for row in result["rows"] if row["row_no"] == 8)

    assert all(row["code"] != "\uacf5\uc885" for row in result["rows"])
    assert all(row["item_name"] != "\ud488\uba85" for row in result["rows"])
    assert target_row["amount"] == 12538592
    assert target_row["unit_price"] == 6368
    assert target_row["labor_amount"] == 7234106
    assert target_row["material_amount"] == 2378552
    assert target_row["expense_amount"] == 2925934


def test_parse_commencement_wbs_excel_real_layout_maps_periods_and_weight_percent():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "\uc608\uc815\uacf5\uc815\ud45c"
    worksheet["A1"] = "\uc608\uc815\uacf5\uc815\ud45c"
    worksheet["A2"] = "\uacf5\uc0ac\uba85 : \uad6d\ub3c446\ud638\uc120 \ud638\ud3c9IC\uad50(\uc0c1)\ub4f1 \ud3ec\uc7a5 \ubc0f \uc2dc\uc124\ubb3c\ubcf4\uc218\uacf5\uc0ac"
    worksheet["A3"] = "\uacf5\uc885"
    worksheet["C3"] = "\ubcf4\ud560"
    worksheet["D3"] = "\ucc29\uc218"
    worksheet["M3"] = "\ube44\uace0"
    worksheet["D4"] = 20
    worksheet["E4"] = 40
    worksheet["F4"] = 60
    worksheet["G4"] = 80
    worksheet["H4"] = 100
    worksheet["I4"] = 120
    worksheet["A5"] = "1. \uad6d\ub3c446\ud638\uc120 \ud638\ud3c9IC\uad50(\uc0c1)"
    worksheet["C5"] = 0.6336947475660303
    worksheet["F5"] = 0.15842368689150757
    worksheet["G5"] = 0.15842368689150757
    worksheet["H5"] = 0.15842368689150757
    worksheet["I5"] = 0.15842368689150757
    worksheet["D7"] = "\uc790\uc7ac \ubc0f \ud604\uc7a5\uc900\ube44"
    worksheet["A8"] = "2. \uad6d\ub3c446\ud638\uc120 \uc2e0\uad6c\uc6b4\uad50\ub4f1"
    worksheet["C8"] = 0.36630525243396966
    worksheet["F8"] = 0.09157631310849242
    worksheet["G8"] = 0.09157631310849242
    worksheet["H8"] = 0.09157631310849242
    worksheet["I8"] = 0.09157631310849242
    worksheet["A68"] = "\uc18c \uacc4"
    worksheet["C68"] = 1.0
    worksheet.merge_cells("A5:A7")
    worksheet.merge_cells("C5:C7")
    worksheet.merge_cells("A8:A10")
    worksheet.merge_cells("C8:C10")
    worksheet.merge_cells("A68:A70")
    worksheet.merge_cells("C68:C70")

    result = ei.parse_commencement_wbs_excel(_workbook_bytes(workbook))

    assert result["project_name"] == "\uad6d\ub3c446\ud638\uc120 \ud638\ud3c9IC\uad50(\uc0c1)\ub4f1 \ud3ec\uc7a5 \ubc0f \uc2dc\uc124\ubb3c\ubcf4\uc218\uacf5\uc0ac"
    assert [row["name"] for row in result["rows"]] == [
        "1. \uad6d\ub3c446\ud638\uc120 \ud638\ud3c9IC\uad50(\uc0c1)",
        "2. \uad6d\ub3c446\ud638\uc120 \uc2e0\uad6c\uc6b4\uad50\ub4f1",
    ]
    assert len(result["rows"]) == 2
    assert [row["row_no"] for row in result["rows"]] == [5, 8]
    assert result["rows"][0]["name"] == "1. \uad6d\ub3c446\ud638\uc120 \ud638\ud3c9IC\uad50(\uc0c1)"
    assert round(float(result["rows"][0]["weight"]), 4) == 63.3695
    assert round(sum(float(row["weight"]) for row in result["rows"]), 4) == 100.0
    assert all(float(row["weight"]) > 0 for row in result["rows"])
    assert [entry["period"] for entry in result["rows"][0]["period_progress_values"]] == [
        "\ucc29\uc218+20",
        "\ucc29\uc218+40",
        "\ucc29\uc218+60",
        "\ucc29\uc218+80",
        "\ucc29\uc218+100",
        "\ucc29\uc218+120",
    ]
    assert all(entry["period"] not in {"\uacf5\uc885", "\ubcf4\ud560"} for entry in result["rows"][0]["period_progress_values"])
    assert all(row["name"] != "\uc790\uc7ac \ubc0f \ud604\uc7a5\uc900\ube44" for row in result["rows"])
    assert all(row["name"] not in {"\uc18c \uacc4", "\ub204 \uacc4", "\ud569 \uacc4", "0", "\uc7a5 \uace1 \uad50"} for row in result["rows"])


def test_parse_contract_budget_excel_raises_clear_validation_error_for_invalid_workbook():
    with pytest.raises(ValidationError) as exc_info:
        ei.parse_contract_budget_excel(BytesIO(b"not-an-excel-file"))

    assert "\uc5d1\uc140 \ud30c\uc77c\uc744 \uc77d\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4." in str(exc_info.value)


def test_parse_budget_summary_sheet_matrix_row_reads_three_bucket_amounts():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "내역서총괄표(도급)"
    sheet.append(["구분", "도급금액", "노무비", "재료비", "경비"])
    sheet.append(["도급예정액", 571022700, 234054163, 107322412, 229646125])
    sheet.append(["관급자재대", 37700000, 0, 37700000, 0])
    sheet.append(["총공사비", 608722700, 234054163, 145022412, 229646125])

    summary = ei.parse_budget_summary_sheet(workbook)

    assert summary["contract_expected_amount"] == 571022700
    assert summary["contract_labor_amount"] == 234054163
    assert summary["contract_material_amount"] == 107322412
    assert summary["contract_expense_amount"] == 229646125
    assert summary["owner_supplied_amount"] == 37700000
    assert summary["total_construction_amount"] == 608722700


def test_parse_budget_summary_sheet_does_not_read_one_as_labor_or_expense():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "내역서총괄표(도급)"
    sheet.append(["구분", "도급금액", "노무비", "재료비", "경비", "비고"])
    sheet.append(["도급예정액", 571022700, 234054163, 107322412, 229646125, 1])
    sheet.append(["노무비", 1, 1, "", "", ""])
    sheet.append(["경비", 1, "", "", 1, ""])

    summary = ei.parse_budget_summary_sheet(workbook)

    assert summary["contract_labor_amount"] == 234054163
    assert summary["contract_expense_amount"] == 229646125
    assert summary["contract_labor_amount"] != 1
    assert summary["contract_expense_amount"] != 1


def test_parse_budget_summary_sheet_vertical_label_value_layout():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "내역서총괄표(도급)"
    sheet["A1"] = "도급예정액"
    sheet["B1"] = 571022700
    sheet["A2"] = "노무비"
    sheet["B2"] = 234054163
    sheet["A3"] = "재료비"
    sheet["B3"] = 107322412
    sheet["A4"] = "경비"
    sheet["B4"] = 229646125
    sheet["A5"] = "관급자재대"
    sheet["B5"] = 37700000
    sheet["A6"] = "총공사비"
    sheet["B6"] = 608722700

    summary = ei.parse_budget_summary_sheet(workbook)

    assert summary["contract_expected_amount"] == 571022700
    assert summary["contract_labor_amount"] == 234054163
    assert summary["contract_material_amount"] == 107322412
    assert summary["contract_expense_amount"] == 229646125
    assert summary["owner_supplied_amount"] == 37700000
    assert summary["total_construction_amount"] == 608722700
