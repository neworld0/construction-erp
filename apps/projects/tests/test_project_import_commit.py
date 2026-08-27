from io import BytesIO, StringIO

import pytest
from django.contrib.auth import get_user_model
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.sessions.middleware import SessionMiddleware
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import RequestFactory
from openpyxl import Workbook

from apps.audit.models import AuditLog
from apps.core.rbac.models import UserProfile
from apps.cost.models import CostItem, CostItemAlias
from apps.cost.seed_civil_road_cbs import seed_civil_road_cbs
from apps.evidence.models import Evidence
from apps.projects.hq_views import (
    _build_budget_initial_from_import_rows,
    _get_project_budget_totals,
    _make_budget_import_key,
    hq_project_detail,
    hq_project_new,
)
from apps.projects.models import BudgetCategory, BudgetItem, Project, ProjectContract, WBSItem


BUDGET_SHEET = "도급계약내역서"
WBS_SHEET = "예정공정표"
ITEM_ASCON = "아스콘 포장"
ITEM_BRIDGE = "교면난간"
ITEM_UNMATCHED = "UNMATCHED-XYZ-987"
PROJECT_NAME = "엑셀 등록 공사"
PARTIAL_NAME = "부분등록 공사"
CLIENT_A = "발주처 A"
CLIENT_C = "발주처 C"
SITE_SEOUL = "서울"
SITE_CHUNCHEON = "춘천"
WBS_ROW_1 = "1. 국도46호선 호평IC교(상)"
WBS_ROW_2 = "2. 국도46호선 신구로(춘천)"
BLOCK_MSG = "CBS 미매칭 예산 행이 남아 있어 등록할 수 없습니다."


def _build_request(user, data):
    request = RequestFactory().post("/app/hq/projects/new/", data)
    SessionMiddleware(lambda req: None).process_request(request)
    request.session.save()
    setattr(request, "_messages", FallbackStorage(request))
    request.user = user
    return request


def _xlsx_upload(name, workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _budget_workbook_matched_only():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["A-100", ITEM_ASCON, "t=5cm", 10, "m2", 15000, 150000, "", "", "", ""])
    return workbook


def _budget_workbook_with_bucket_split():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1-1-1", "아스팔트 포장 절삭후 아스팔트 덧씌우기", "A-Type (1회절삭,2회포장)-야간작업", 1969, "㎡", 6368, 12538592, 7234106, 2378552, 2925934, ""])
    return workbook


def _budget_workbook_with_duplicate_labor_rows():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1-1", "안전관리책임자", "", 1, "식", 1000, 1000, 1000, 0, 0, ""])
    ws.append(["1-2", "신호수", "", 1, "식", 2000, 2000, 2000, 0, 0, ""])
    return workbook


def _budget_workbook_with_unmatched():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1", "국도46호선 호평IC교(상)", "", "-", "-", "", 181183936, "", "", "", ""])
    ws.append(["1-1", ITEM_ASCON, "t=5cm", 10, "m2", 15000, 150000, "", "", "", ""])
    ws.append(["1-2", ITEM_UNMATCHED, "특수규격", 1, "식", 200000, 200000, 0, 0, 0, ""])
    return workbook


def _budget_workbook_bridge_unmatched():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["B-100", ITEM_BRIDGE, "보수", 1, "식", 100000, 100000, 0, 0, 0, ""])
    return workbook


def _wbs_workbook():
    workbook = Workbook()
    ws = workbook.active
    ws.title = WBS_SHEET
    ws["A1"] = WBS_SHEET
    ws["A2"] = "공사명: 국도46호선 포장 보수공사"
    ws["A3"] = "공종"
    ws["C3"] = "보할"
    ws["D3"] = "착수"
    ws["M3"] = "비고"
    ws["D4"] = 20
    ws["E4"] = 40
    ws["F4"] = 60
    ws["G4"] = 80
    ws["H4"] = 100
    ws["I4"] = 120
    ws["A5"] = WBS_ROW_1
    ws["C5"] = 0.6336947475660303
    ws["F5"] = 0.15842368689150757
    ws["G5"] = 0.15842368689150757
    ws["H5"] = 0.15842368689150757
    ws["I5"] = 0.15842368689150757
    ws["A8"] = WBS_ROW_2
    ws["C8"] = 0.36630525243396966
    ws["F8"] = 0.09157631310849242
    ws["G8"] = 0.09157631310849242
    ws["H8"] = 0.09157631310849242
    ws["I8"] = 0.09157631310849242
    return workbook


def _base_form_data():
    return {
        "name": PROJECT_NAME,
        "project_type": "civil",
        "status": "draft",
        "client_name": CLIENT_A,
        "site_address": SITE_SEOUL,
        "start_date": "2026-05-01",
        "end_date": "2026-05-31",
        "contract_amount": "1000000",
        "contract_start_date": "2026-05-01",
        "contract_end_date": "2026-05-31",
        "contract_file": SimpleUploadedFile("contract.pdf", b"%PDF-1.4\n", content_type="application/pdf"),
        "wbs-TOTAL_FORMS": "2",
        "wbs-INITIAL_FORMS": "0",
        "wbs-MIN_NUM_FORMS": "0",
        "wbs-MAX_NUM_FORMS": "1000",
        "wbs-0-name": WBS_ROW_1,
        "wbs-0-parent": "",
        "wbs-0-weight": "63.37",
        "wbs-0-sort_order": "1",
        "wbs-0-plan_start_date": "",
        "wbs-0-plan_end_date": "",
        "wbs-1-name": WBS_ROW_2,
        "wbs-1-parent": "",
        "wbs-1-weight": "36.63",
        "wbs-1-sort_order": "2",
        "wbs-1-plan_start_date": "",
        "wbs-1-plan_end_date": "",
        "confirm_import_warnings": "1",
    }


def _build_detail_request(user, project_id, method=None, data=None):
    factory = RequestFactory()
    if method is None:
        method = "post" if data is not None else "get"
    request = getattr(factory, method.lower())(f"/app/hq/projects/{project_id}/", data or {})
    SessionMiddleware(lambda req: None).process_request(request)
    request.session.save()
    setattr(request, "_messages", FallbackStorage(request))
    request.user = user
    return request


def _labor_cost_item():
    return CostItem.objects.filter(code="CIVIL-LABOR").first() or CostItem.objects.get(code="LABOR-GENERAL")


@pytest.mark.django_db
def test_hq_project_new_commits_imported_budget_and_wbs_rows():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="commit-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    cost_item = CostItem.objects.get(code="CIVIL-ASCON-PAVING")

    data = _base_form_data()
    data.update(
        {
            "action": "save_project",
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_matched_only()),
            "commencement_excel": _xlsx_upload("wbs.xlsx", _wbs_workbook()),
            "budget-TOTAL_FORMS": "1",
            "budget-INITIAL_FORMS": "0",
            "budget-MIN_NUM_FORMS": "0",
            "budget-MAX_NUM_FORMS": "1000",
            "budget-0-category": "MATERIAL",
            "budget-0-cost_item": str(cost_item.id),
            "budget-0-name": ITEM_ASCON,
            "budget-0-planned_amount": "150000",
            "budget-0-note": "엑셀 합산 1행 / 원본:도급계약내역서:4",
        }
    )
    request = _build_request(user, data)

    response = hq_project_new(request)

    assert response.status_code == 302
    project = Project.objects.latest("id")
    assert BudgetItem.objects.filter(project=project).count() == 1
    assert WBSItem.objects.filter(project=project).count() == 2
    assert Evidence.objects.filter(object_type="PROJECT", object_id=project.id).count() == 2
    actions = list(AuditLog.objects.filter(project=project).values_list("action", flat=True))
    assert "PROJECT_BUDGET_IMPORT_COMMIT" in actions
    assert "PROJECT_WBS_IMPORT_COMMIT" in actions
    assert "PROJECT_IMPORT_COMMIT" in actions


@pytest.mark.django_db
def test_source_row_bucket_save_allows_same_cost_item_multiple_budget_rows():
    seed_civil_road_cbs()
    project = Project.objects.create(
        code="PRJ-DUP-CBS-001",
        name="중복 CBS 예산 공사",
        project_type="civil",
        status="draft",
    )
    labor_item = _labor_cost_item()

    BudgetItem.objects.create(
        project=project,
        cost_item=labor_item,
        category=BudgetCategory.LABOR,
        name="아스팔트 포장 절삭후 아스팔트 덧씌우기 - 노무비",
        planned_amount=1000,
    )
    BudgetItem.objects.create(
        project=project,
        cost_item=labor_item,
        category=BudgetCategory.LABOR,
        name="교면방수 - 노무비",
        planned_amount=2000,
    )

    items = list(
        BudgetItem.objects.filter(project=project, cost_item=labor_item).order_by("id")
    )
    assert len(items) == 2
    assert sum(item.planned_amount for item in items) == 3000


@pytest.mark.django_db
def test_hq_project_new_commit_source_row_bucket_with_duplicate_cbs_does_not_500():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(
        username="duplicate-cbs-hq", password="pass"
    )
    UserProfile.objects.create(user=user, role="hq")

    data = _base_form_data()
    data.update(
        {
            "action": "save_project",
            "name": "중복 CBS 등록 공사",
            "budget_excel": _xlsx_upload(
                "budget-duplicate-labor.xlsx",
                _budget_workbook_with_duplicate_labor_rows(),
            ),
            "commencement_excel": _xlsx_upload("wbs.xlsx", _wbs_workbook()),
            "confirm_budget_reconciliation_gap": "1",
        }
    )
    request = _build_request(user, data)

    response = hq_project_new(request)

    assert response.status_code == 302
    project = Project.objects.get(name="중복 CBS 등록 공사")
    labor_item = _labor_cost_item()
    items = list(
        BudgetItem.objects.filter(project=project, cost_item=labor_item).order_by("id")
    )
    assert len(items) >= 2
    assert sum(item.planned_amount for item in items) == 3000
    assert any("안전관리책임자" in (item.name or "") for item in items)
    assert any("신호수" in (item.name or "") for item in items)


@pytest.mark.django_db
def test_save_blocks_when_budget_excel_has_unmatched_rows_without_confirm():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="block-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    cost_item = CostItem.objects.get(code="CIVIL-ASCON-PAVING")

    data = _base_form_data()
    data.update(
        {
            "action": "save_project",
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_unmatched()),
            "commencement_excel": _xlsx_upload("wbs.xlsx", _wbs_workbook()),
            "budget-TOTAL_FORMS": "1",
            "budget-INITIAL_FORMS": "0",
            "budget-MIN_NUM_FORMS": "0",
            "budget-MAX_NUM_FORMS": "1000",
            "budget-0-category": "MATERIAL",
            "budget-0-cost_item": str(cost_item.id),
            "budget-0-name": ITEM_ASCON,
            "budget-0-planned_amount": "150000",
            "budget-0-note": "엑셀 합산 1행 / 원본:도급계약내역서:4",
        }
    )
    request = _build_request(user, data)

    response = hq_project_new(request)
    content = response.content.decode("utf-8")

    assert response.status_code == 200
    assert not Project.objects.filter(name=PROJECT_NAME).exists()
    assert "CBS \ubbf8\ub9e4\uce6d \ud589\uc774 \uc788\uc5b4 \uc608\uc0b0 \uae30\uc900\uc120\uc744 \uc800\uc7a5\ud560 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4." in content


@pytest.mark.django_db
def test_save_blocks_partial_import_even_when_confirmed():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="partial-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    cost_item = CostItem.objects.get(code="CIVIL-ASCON-PAVING")

    data = _base_form_data()
    data.update(
        {
            "action": "save_project",
            "name": PARTIAL_NAME,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_unmatched()),
            "commencement_excel": _xlsx_upload("wbs.xlsx", _wbs_workbook()),
            "budget-TOTAL_FORMS": "1",
            "budget-INITIAL_FORMS": "0",
            "budget-MIN_NUM_FORMS": "0",
            "budget-MAX_NUM_FORMS": "1000",
            "budget-0-category": "MATERIAL",
            "budget-0-cost_item": str(cost_item.id),
            "budget-0-name": ITEM_ASCON,
            "budget-0-planned_amount": "150000",
            "budget-0-note": "엑셀 합산 1행 / 원본:도급계약내역서:4",
            "confirm_partial_budget_import": "1",
        }
    )
    request = _build_request(user, data)

    response = hq_project_new(request)

    content = response.content.decode("utf-8")

    assert response.status_code == 200
    assert not Project.objects.filter(name=PARTIAL_NAME).exists()
    assert "CBS \ubbf8\ub9e4\uce6d \ud589\uc774 \uc788\uc5b4 \uc608\uc0b0 \uae30\uc900\uc120\uc744 \uc800\uc7a5\ud560 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4." in content


@pytest.mark.django_db
def test_alias_saved_from_manual_mapping_is_used_next_preview():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="alias-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    bridge_item = CostItem.objects.get(code="CIVIL-BRIDGE-REPAIR")

    preview_row = {
        "code": "B-100",
        "row_no": 4,
        "item_name": ITEM_BRIDGE,
        "spec": "보수",
        "amount": 100000,
        "quantity": 1,
        "unit": "식",
        "warnings": [],
    }
    import_key = _make_budget_import_key(preview_row, 0)

    request = _build_request(
        user,
        {
            "action": "apply_budget_cbs_mapping",
            "name": "수동 매핑 공사",
            "project_type": "civil",
            "client_name": CLIENT_C,
            "site_address": SITE_CHUNCHEON,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_bridge_unmatched()),
            f"budget_match__{import_key}": str(bridge_item.id),
            f"budget_alias_save__{import_key}": "1",
        },
    )
    response = hq_project_new(request)

    assert response.status_code == 200
    assert CostItemAlias.objects.filter(cost_item=bridge_item, alias=ITEM_BRIDGE).exists()

    rows = [
        {
            "code": "B-100",
            "row_no": 4,
            "item_name": ITEM_BRIDGE,
            "spec": "보수",
            "amount": 100000,
            "quantity": 1,
            "unit": "식",
            "warnings": [],
        }
    ]
    initial, warnings, _stats = _build_budget_initial_from_import_rows(rows)
    assert warnings == []
    assert len(initial) == 1
    assert CostItem.objects.get(id=initial[0]["cost_item"]).code == "CIVIL-BRIDGE-REPAIR"


@pytest.mark.django_db
def test_project_detail_shows_budget_totals():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="detail-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    project = Project.objects.create(code="PRJ-BUDGET-001", name="예산 합계 확인 공사", project_type="civil", status="draft")
    ProjectContract.objects.create(
        project=project,
        contract_amount=10000,
        contract_start_date="2026-05-01",
        contract_end_date="2026-05-31",
        status="approved",
    )
    expense_item = CostItem.objects.get(code="CIVIL-EXPENSE")
    profit_item = CostItem.objects.get(code="CIVIL-PROFIT")
    labor_item = _labor_cost_item()
    BudgetItem.objects.create(project=project, cost_item=expense_item, category=BudgetCategory.OTHER, name="경비", planned_amount=1000)
    BudgetItem.objects.create(project=project, cost_item=profit_item, category=BudgetCategory.OTHER, name="이윤", planned_amount=2000)
    BudgetItem.objects.create(project=project, cost_item=labor_item, category=BudgetCategory.LABOR, name="노무비", planned_amount=3000)

    request = _build_detail_request(user, project.id)
    response = hq_project_detail(request, project.id)
    content = response.content.decode("utf-8")

    assert response.status_code == 200
    assert "노무비 합계" in content
    assert "재료비 합계" in content
    assert "경비 합계" in content
    assert "예산 기준선 총합계" in content
    assert "3,000" in content
    assert "6,000" in content


@pytest.mark.django_db
def test_repair_project_budget_cbs_command_moves_safety_manager_out_of_profit():
    seed_civil_road_cbs()
    project = Project.objects.create(code="PRJ-REPAIR-001", name="CBS 보정 공사", project_type="civil", status="draft")
    profit_item = CostItem.objects.get(code="CIVIL-PROFIT")
    labor_item = _labor_cost_item()
    BudgetItem.objects.create(
        project=project,
        cost_item=profit_item,
        category=BudgetCategory.OTHER,
        name="이윤",
        planned_amount=14127540,
        note="엑셀 합산 1행 / 주요 항목:안전관리책임자 / 원본:도급계약내역서:204",
    )

    out = StringIO()
    call_command("repair_project_budget_cbs", project_id=project.id, stdout=out)

    assert "After totals:" in out.getvalue()
    assert not BudgetItem.objects.filter(project=project, cost_item=profit_item, note__contains="안전관리책임자").exists()
    labor_budget = BudgetItem.objects.get(project=project, cost_item=labor_item)
    assert labor_budget.planned_amount == 14127540
    assert "안전관리책임자" in labor_budget.note


@pytest.mark.django_db
def test_profit_is_general_budget_even_if_saved_as_labor():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="profit-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    project = Project.objects.create(code="PRJ-PROFIT-001", name="이윤 분류 공사", project_type="civil", status="draft")
    ProjectContract.objects.create(project=project, contract_amount=1000, contract_start_date="2026-05-01", contract_end_date="2026-05-31", status="approved")
    profit_item = CostItem.objects.get(code="CIVIL-PROFIT")
    BudgetItem.objects.create(project=project, cost_item=profit_item, category=BudgetCategory.LABOR, name="이윤", planned_amount=1000)

    request = _build_detail_request(user, project.id)
    response = hq_project_detail(request, project.id)
    content = response.content.decode("utf-8")

    assert response.status_code == 200
    assert "노무비 합계" in content
    assert "경비 합계" in content
    budget = BudgetItem.objects.get(project=project, cost_item=profit_item)
    assert budget.category == BudgetCategory.OTHER


@pytest.mark.django_db
def test_commit_work_item_bucket_import_saves_three_budget_items():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="bucket-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")

    data = _base_form_data()
    data.update(
        {
            "action": "save_project",
            "name": "버킷 분할 등록 공사",
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_bucket_split()),
            "commencement_excel": _xlsx_upload("wbs.xlsx", _wbs_workbook()),
            "confirm_budget_reconciliation_gap": "1",
        }
    )
    request = _build_request(user, data)

    response = hq_project_new(request)

    assert response.status_code == 302
    project = Project.objects.get(name="버킷 분할 등록 공사")
    items = list(BudgetItem.objects.filter(project=project).order_by("id"))
    assert len(items) == 3
    assert sum(item.planned_amount for item in items) == 12538592
    assert {item.category for item in items} == {BudgetCategory.LABOR, BudgetCategory.MATERIAL, BudgetCategory.OTHER}
    notes = " ".join(item.note or "" for item in items)
    assert "버킷:노무비" in notes
    assert "버킷:재료비" in notes
    assert "버킷:경비" in notes


@pytest.mark.django_db
def test_project_detail_bucket_totals_treat_non_labor_material_as_expense():
    seed_civil_road_cbs()
    user = get_user_model().objects.create_user(username="practical-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    project = Project.objects.create(code="PRJ-BUCKET-001", name="버킷 합계 공사", project_type="civil", status="draft")
    ProjectContract.objects.create(project=project, contract_amount=2100, contract_start_date="2026-05-01", contract_end_date="2026-05-31", status="approved")
    labor_item = _labor_cost_item()
    material_item = CostItem.objects.get(code="CIVIL-MATERIAL")
    subcon_item = CostItem.objects.filter(category="subcon").first()
    if subcon_item is None:
        subcon_item = CostItem.objects.create(code="TEST-SUBCON-BUCKET", name="하도급 테스트", category="subcon", is_active=True)
    expense_item = CostItem.objects.get(code="CIVIL-EXPENSE")
    equipment_item = CostItem.objects.get(code="CIVIL-EQUIPMENT")
    profit_item = CostItem.objects.get(code="CIVIL-PROFIT")
    BudgetItem.objects.create(project=project, cost_item=labor_item, category=BudgetCategory.LABOR, name="노무비", planned_amount=100)
    BudgetItem.objects.create(project=project, cost_item=material_item, category=BudgetCategory.MATERIAL, name="재료비", planned_amount=200)
    BudgetItem.objects.create(project=project, cost_item=subcon_item, category=BudgetCategory.SUBCON, name="하도급", planned_amount=300)
    BudgetItem.objects.create(project=project, cost_item=expense_item, category=BudgetCategory.OTHER, name="경비", planned_amount=400)
    BudgetItem.objects.create(project=project, cost_item=equipment_item, category=BudgetCategory.EQUIP, name="장비비", planned_amount=500)
    BudgetItem.objects.create(project=project, cost_item=profit_item, category=BudgetCategory.OTHER, name="이윤", planned_amount=600)

    request = _build_detail_request(user, project.id)
    response = hq_project_detail(request, project.id)
    content = response.content.decode("utf-8")
    totals = _get_project_budget_totals(project)

    assert response.status_code == 200
    assert "계약금액 - 예산 기준선 차이" in content
    assert totals["material_budget_total_amount"] == 200
    assert totals["subcontract_budget_total_amount"] == 300
    assert totals["labor_budget_total_amount"] == 100
    assert totals["expense_budget_total_amount"] == 1500
    assert totals["budget_grand_total_amount"] == 2100


@pytest.mark.django_db
def test_repair_project_budget_categories_command():
    seed_civil_road_cbs()
    project = Project.objects.create(code="PRJ-REPAIR-002", name="카테고리 보정 공사", project_type="civil", status="draft")
    ProjectContract.objects.create(project=project, contract_amount=5000, contract_start_date="2026-05-01", contract_end_date="2026-05-31", status="approved")
    profit_item = CostItem.objects.get(code="CIVIL-PROFIT")
    BudgetItem.objects.create(project=project, cost_item=profit_item, category=BudgetCategory.LABOR, name="", planned_amount=1000)

    out = StringIO()
    call_command("repair_project_budget_categories", project_id=project.id, stdout=out)
    call_command("repair_project_budget_categories", project_id=project.id, stdout=out)

    repaired = BudgetItem.objects.get(project=project, cost_item=profit_item)
    assert repaired.category == BudgetCategory.OTHER
    assert repaired.name == "이윤"
    output = out.getvalue()
    assert "재료비 합계" in output
    assert "경비 합계" in output


@pytest.mark.django_db
def test_repair_project_budget_categories_maps_equip_to_other():
    seed_civil_road_cbs()
    project = Project.objects.create(code="PRJ-REPAIR-EQUIP-001", name="장비 카테고리 보정 공사", project_type="civil", status="draft")
    equipment_item = CostItem.objects.get(code="CIVIL-EQUIPMENT")
    budget_item = BudgetItem.objects.create(project=project, cost_item=equipment_item, category=BudgetCategory.EQUIP, name="장비비", planned_amount=500)

    call_command("repair_project_budget_categories", project_id=project.id)

    budget_item.refresh_from_db()
    assert budget_item.category == BudgetCategory.OTHER
