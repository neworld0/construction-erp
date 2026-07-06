from io import BytesIO
import re

import pytest
from django.contrib.auth import get_user_model
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.sessions.middleware import SessionMiddleware
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory
from openpyxl import Workbook

from apps.core.rbac.models import UserProfile
from apps.cost.models import CostItem, CostItemAlias
from apps.cost.seed_civil_road_cbs import seed_civil_road_cbs
from apps.projects.hq_views import hq_project_new


BUDGET_SHEET = "도급계약내역서"
WBS_SHEET = "예정공정표"
PROJECT_NAME = "미리보기 공사"
CLIENT_A = "발주처 A"
CLIENT_B = "발주처 B"
SITE_SEOUL = "서울"
SITE_GENERIC = "현장주소"
ITEM_PAVING = "아스팔트 포장 절삭후 아스팔트 덧씌우기"
ITEM_UNMATCHED = "UNMATCHED-XYZ-987"
ROW_TITLE = "국도46호선 호평IC교(상)"
WBS_ROW_1 = "1. 국도46호선 호평IC교(상)"
WBS_ROW_2 = "2. 국도46호선 신구로(춘천)"
MSG_BUDGET_NOTICE = "도급계약내역서 시트에서 예산 행을 불러왔습니다."
MSG_WBS_NOTICE = "예정공정표 시트에서 예정공정표를 불러왔습니다."
MSG_INVALID_XLSX = "계약내역서 Excel은 .xlsx 파일만 업로드할 수 있습니다."
MSG_TEMPLATE_FALLBACK = "현재 선택한 공종의 템플릿이 없어 조경 기본 템플릿이 적용되었습니다. 필요한 경우 수정해 주세요."
MSG_TEMPLATE_EMPTY = "선택 가능한 템플릿이 없어 기본 입력 행으로 표시합니다."


def _xlsx_upload(name, workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _build_request(user, data):
    request = RequestFactory().post("/app/hq/projects/new/", data)
    SessionMiddleware(lambda req: None).process_request(request)
    request.session.save()
    setattr(request, "_messages", FallbackStorage(request))
    request.user = user
    return request


def _build_get_request(user, query_string=""):
    path = "/app/hq/projects/new/"
    if query_string:
        path = f"{path}?{query_string}"
    request = RequestFactory().get(path)
    SessionMiddleware(lambda req: None).process_request(request)
    request.session.save()
    setattr(request, "_messages", FallbackStorage(request))
    request.user = user
    return request


def _budget_workbook_with_match_and_unmatch():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1", ROW_TITLE, "", "-", "-", "", 181183936, "", "", "", ""])
    ws.append(["1-1", ITEM_PAVING, "A-Type", 1969, "㎡", 6368, 12538592, 7234106, 2378552, 2925934, ""])
    ws.append(["1-2", ITEM_UNMATCHED, "특수규격", 1, "식", 100000, 100000, 0, 0, 0, ""])
    return workbook


def _budget_workbook_with_summary_sheet():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1", "도급예정액", "", "-", "-", "", 572446289, "", "", "", ""])
    ws.append(["1-1", ITEM_PAVING, "A-Type", 1969, "㎡", 6368, 12538592, 7234106, 2378552, 2925934, ""])

    summary_ws = workbook.create_sheet("내역서총괄표(도급)")
    summary_ws["A1"] = "도급예정액"
    summary_ws["B1"] = 571022700
    summary_ws["A2"] = "총공사비"
    summary_ws["B2"] = 609884852
    return workbook


def _budget_workbook_with_owner_supplied():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1", ROW_TITLE, "", "-", "-", "", 181183936, "", "", "", ""])
    ws.append(["1-1", ITEM_PAVING, "A-Type", 1969, "㎡", 6368, 12538592, 7234106, 2378552, 2925934, ""])
    ws.append(["5.", "관급자재대", "", "-", "-", "-", 37700000, "", "", "", ""])
    ws.append(["5-1", "호평IC교(상)", "", "-", "-", "-", 37645120, "", "", "", ""])
    ws.append(["", "아스콘(관급)-서울,인천,경기(강북,노원,도봉구 제외)", "조합공통품목, WC-6, 가열, 1등급", 376, "톤", 100120, 37645120, "", "", "", ""])
    ws.append(["5-2", "단수조정", "", 1, "식", 54880, 54880, 0, 0, 0, ""])
    return workbook


def _budget_workbook_with_bucket_gap_and_manual_profit():
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1-1", ITEM_PAVING, "A-Type", 1, "식", 310437127, 310437127, 203114715, 107322412, 0, ""])
    ws.append(["1-2", "이윤", "", 1, "식", 30664306, 30664306, 0, 0, 30664306, ""])

    summary_ws = workbook.create_sheet("내역서총괄표(도급)")
    summary_ws["A1"] = "도급예정액"
    summary_ws["B1"] = 571022700
    summary_ws["A2"] = "노무비"
    summary_ws["B2"] = 218435187
    summary_ws["A3"] = "재료비"
    summary_ws["B3"] = 107322412
    summary_ws["A4"] = "경비"
    summary_ws["B4"] = 245265101
    return workbook


def _wbs_workbook():
    workbook = Workbook()
    ws = workbook.active
    ws.title = WBS_SHEET
    ws["A1"] = WBS_SHEET
    ws["A2"] = "공사명: 국도46호선 포장 보수공사"
    ws["A3"] = "공종"
    ws["C3"] = "보할"
    ws["D3"] = "착수"
    ws["M3"] = "비고"
    ws["D4"] = 20
    ws["E4"] = 40
    ws["F4"] = 60
    ws["G4"] = 80
    ws["H4"] = 100
    ws["I4"] = 120
    ws["A5"] = WBS_ROW_1
    ws["C5"] = 0.6336947475660303
    ws["F5"] = 0.15842368689150757
    ws["G5"] = 0.15842368689150757
    ws["H5"] = 0.15842368689150757
    ws["I5"] = 0.15842368689150757
    ws["A8"] = WBS_ROW_2
    ws["C8"] = 0.36630525243396966
    ws["F8"] = 0.09157631310849242
    ws["G8"] = 0.09157631310849242
    ws["H8"] = 0.09157631310849242
    ws["I8"] = 0.09157631310849242
    return workbook


@pytest.fixture
def hq_user(db):
    user = get_user_model().objects.create_user(username="preview-hq", password="pass")
    UserProfile.objects.create(user=user, role="hq")
    return user


@pytest.mark.django_db
def test_hq_project_new_preview_renders_budget_and_wbs_preview(hq_user):
    seed_civil_road_cbs()
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": PROJECT_NAME,
            "project_type": "civil",
            "client_name": CLIENT_A,
            "site_address": SITE_SEOUL,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
            "commencement_excel": _xlsx_upload("wbs.xlsx", _wbs_workbook()),
        },
    )
    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert PROJECT_NAME in content
    assert MSG_BUDGET_NOTICE in content
    assert MSG_WBS_NOTICE in content
    assert ITEM_PAVING in content
    assert WBS_ROW_1 in content


@pytest.mark.django_db
def test_hq_project_new_preview_shows_manual_cbs_mapping_for_unmatched_row(hq_user):
    seed_civil_road_cbs()
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": "CBS 수동 매핑 공사",
            "project_type": "civil",
            "client_name": CLIENT_B,
            "site_address": SITE_GENERIC,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
        },
    )
    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert ITEM_UNMATCHED in content
    assert "CBS 매핑 적용" in content
    assert "budget_match__" in content
    assert "budget_alias_save__" in content
    assert "예산 자동 매칭 상세 조정" in content
    assert "budget_review_category__" in content
    assert "budget_review_cost_item__" in content
    assert "budget_review_amount__" in content


@pytest.mark.django_db
def test_apply_budget_cbs_mapping_saves_alias_and_shows_message(hq_user):
    seed_civil_road_cbs()
    bridge_item = CostItem.objects.get(code="CIVIL-BRIDGE-REPAIR")
    preview_request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": "CBS alias 생성 공사",
            "project_type": "civil",
            "client_name": CLIENT_B,
            "site_address": SITE_GENERIC,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
        },
    )
    preview_response = hq_project_new(preview_request)
    preview_content = preview_response.content.decode("utf-8")
    match = re.search(r'name="budget_match__([^"]+)"', preview_content)
    assert match is not None
    import_key = match.group(1)

    request = _build_request(
        hq_user,
        {
            "action": "apply_budget_cbs_mapping",
            "name": "CBS alias 생성 공사",
            "project_type": "civil",
            "client_name": CLIENT_B,
            "site_address": SITE_GENERIC,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
            f"budget_match__{import_key}": str(bridge_item.id),
            f"budget_alias_save__{import_key}": "1",
        },
    )

    response = hq_project_new(request)

    assert response.status_code == 200
    assert CostItemAlias.objects.filter(
        cost_item=bridge_item,
        alias=ITEM_UNMATCHED,
        note="project_import_manual_mapping",
    ).exists()
    content = response.content.decode("utf-8")
    assert "CBS alias 2건을 저장했습니다." in content
    assert ITEM_UNMATCHED in content
    assert f"{ITEM_UNMATCHED} → {bridge_item.name}" in content


@pytest.mark.django_db
def test_saved_alias_is_used_on_next_preview(hq_user):
    seed_civil_road_cbs()
    bridge_item = CostItem.objects.get(code="CIVIL-BRIDGE-REPAIR")
    CostItemAlias.objects.get_or_create(
        cost_item=bridge_item,
        alias=ITEM_UNMATCHED,
        defaults={"is_primary": False, "note": "project_import_manual_mapping"},
    )
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": "alias 재확인 공사",
            "project_type": "civil",
            "client_name": CLIENT_A,
            "site_address": SITE_SEOUL,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
        },
    )

    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert ITEM_UNMATCHED in content
    assert "CBS 미매칭: 0건" in content
    assert "CBS 자동 매칭: 2건" in content
    assert "CBS 자동 매칭이 필요합니다." not in content


@pytest.mark.django_db
def test_alias_not_saved_when_cbs_not_selected(hq_user):
    seed_civil_road_cbs()
    preview_request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": "CBS alias 미선택 공사",
            "project_type": "civil",
            "client_name": CLIENT_B,
            "site_address": SITE_GENERIC,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
        },
    )
    preview_response = hq_project_new(preview_request)
    preview_content = preview_response.content.decode("utf-8")
    match = re.search(r'name="budget_match__([^"]+)"', preview_content)
    assert match is not None
    import_key = match.group(1)

    request = _build_request(
        hq_user,
        {
            "action": "apply_budget_cbs_mapping",
            "name": "CBS alias 미선택 공사",
            "project_type": "civil",
            "client_name": CLIENT_B,
            "site_address": SITE_GENERIC,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
            f"budget_alias_save__{import_key}": "1",
        },
    )

    response = hq_project_new(request)

    assert response.status_code == 200
    assert not CostItemAlias.objects.filter(alias=ITEM_UNMATCHED).exists()
    content = response.content.decode("utf-8")
    assert "CBS alias 1건은 저장되지 않았습니다. CBS 선택과 품명을 확인해 주세요." in content
    assert "CBS가 선택되지 않아 alias를 저장하지 않았습니다." in content


@pytest.mark.django_db
def test_hq_project_new_preview_rejects_invalid_excel_file_type(hq_user):
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": "잘못된 파일 공사",
            "project_type": "landscape",
            "budget_excel": SimpleUploadedFile("budget.txt", b"invalid", content_type="text/plain"),
        },
    )
    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert MSG_INVALID_XLSX in content


@pytest.mark.django_db
def test_hq_project_new_template_fallback_notices_are_valid_korean(hq_user):
    request = _build_get_request(hq_user, "project_type=civil")
    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert MSG_TEMPLATE_EMPTY in content or MSG_TEMPLATE_FALLBACK in content


@pytest.mark.django_db
def test_summary_sheet_contract_amount_is_authoritative(hq_user):
    seed_civil_road_cbs()
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": PROJECT_NAME,
            "project_type": "civil",
            "client_name": CLIENT_A,
            "site_address": SITE_SEOUL,
            "contract_amount": "571022700",
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_summary_sheet()),
        },
    )

    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert "총괄표 도급예정액:" in content
    assert "571,022,700" in content
    assert "계약금액 대비 차이:" in content
    assert "572,446,289" not in content


@pytest.mark.django_db
def test_preview_shows_owner_supplied_as_reference_not_cbs_mapping(hq_user):
    seed_civil_road_cbs()
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": "관급자재 미리보기 공사",
            "project_type": "civil",
            "client_name": CLIENT_A,
            "site_address": SITE_SEOUL,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_owner_supplied()),
        },
    )
    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert "관급자재 참고금액" in content
    assert "관급자재 참고항목" in content
    assert "관급자재 제외 행" in content
    assert "아스콘(관급)-서울,인천,경기(강북,노원,도봉구 제외)" in content
    assert "budget_match__" not in content
    assert "budget_review_category__" in content  # normal editable rows still exist
    assert "원본:도급계약내역서:7" not in content


@pytest.mark.django_db
def test_preview_reports_bucket_gap_when_generated_totals_do_not_match_summary(hq_user):
    seed_civil_road_cbs()
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": PROJECT_NAME,
            "project_type": "civil",
            "client_name": CLIENT_A,
            "site_address": SITE_SEOUL,
            "budget_excel": _xlsx_upload("budget-gap.xlsx", _budget_workbook_with_bucket_gap_and_manual_profit()),
        },
    )
    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert "노무비 차이" in content
    assert "15,320,472" in content
    assert "경비 차이" in content
    assert "214,600,795" in content
    assert "비목 대사 오류 후보" in content
    assert "수동 매핑 필요" in content
    assert "예산 자동 매칭 상세 조정" in content
    assert "budget_review_category__" in content


@pytest.mark.django_db
def test_budget_review_table_uses_wide_note_column(hq_user):
    seed_civil_road_cbs()
    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": PROJECT_NAME,
            "project_type": "civil",
            "client_name": CLIENT_A,
            "site_address": SITE_SEOUL,
            "budget_excel": _xlsx_upload("budget.xlsx", _budget_workbook_with_match_and_unmatch()),
        },
    )

    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert "예산 자동 매칭 상세 조정" in content
    assert "budget-review-table-wrap" in content
    assert "budget-review-table" in content
    assert "col-note" in content
    assert "budget-review-note" in content
    assert "min-width: 1500px" in content or "min-width:1500px" in content


@pytest.mark.django_db
def test_project_new_preview_shows_correct_summary_three_bucket_amounts(hq_user):
    seed_civil_road_cbs()
    workbook = Workbook()
    ws = workbook.active
    ws.title = BUDGET_SHEET
    ws.append(["계약내역"])
    ws.append(["", "", "", "", "", "", "", "", "", "", ""])
    ws.append(["공종", "품명", "규격", "수량", "단위", "단가", "금액", "노무비", "재료비", "경비", "비고"])
    ws.append(["1", "노무비 직접행", "", 1, "식", 234054163, 234054163, 234054163, 0, 0, ""])
    ws.append(["2", "재료비 직접행", "", 1, "식", 107322412, 107322412, 0, 107322412, 0, ""])
    ws.append(["3", "경비 직접행", "", 1, "식", 229646125, 229646125, 0, 0, 229646125, ""])
    summary_ws = workbook.create_sheet("내역서총괄표(도급)")
    summary_ws.append(["구분", "도급금액", "노무비", "재료비", "경비"])
    summary_ws.append(["도급예정액", 571022700, 234054163, 107322412, 229646125])
    summary_ws.append(["관급자재대", 37700000, 0, 37700000, 0])
    summary_ws.append(["총공사비", 608722700, 234054163, 145022412, 229646125])

    request = _build_request(
        hq_user,
        {
            "action": "parse_all",
            "name": PROJECT_NAME,
            "project_type": "civil",
            "client_name": CLIENT_A,
            "site_address": SITE_SEOUL,
            "budget_excel": _xlsx_upload("budget-summary.xlsx", workbook),
        },
    )

    response = hq_project_new(request)

    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert "총괄표 도급예정액" in content
    assert "571,022,700" in content
    assert "총괄표 노무비" in content
    assert "234,054,163" in content
    assert "총괄표 재료비" in content
    assert "107,322,412" in content
    assert "총괄표 경비" in content
    assert "229,646,125" in content
    assert "총괄표 노무비: 1" not in content
    assert "총괄표 경비: 1" not in content
    assert "노무비 차이: 0" in content
    assert "재료비 차이: 0" in content
    assert "경비 차이: 0" in content
    assert "총액 차이: 0" in content
